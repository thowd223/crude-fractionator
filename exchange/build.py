"""Discipline task and information exchange register.

    python exchange/build.py            # check data/*.json, write out/ (workbook, page, item catalog)
    python exchange/build.py --check PIPE   # check only (errors for one discipline), writes nothing

Reads one JSON file per discipline (see SPEC.md), checks every reference, derives the exchanges (one row per
item sent from its owner task to a task that uses it) and writes:
  out/Task-and-Interface-Register.xlsx   tasks, items, exchanges, discipline and system matrices, manual hand-offs
  out/interfaces.html                    interactive register
  out/item-catalog.md                    every item by discipline (reference for authors)
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
DATA, OUT = HERE / "data", HERE / "out"

DISC = {
    "OWNR": "Owner and operations", "PROJ": "Project and controls", "PROC": "Process",
    "SAFE": "Process safety and environmental", "FIRE": "Fire protection", "MECH": "Mechanical", "PIPE": "Piping", "CIVL": "Civil and structural",
    "INST": "Instrumentation and control", "ELEC": "Electrical", "SCM": "Procurement and materials",
    "VEND": "Vendors and fabricators", "CONS": "Construction", "COMM": "Completions and commissioning",
}
SYSTEMS = ["HYSYS", "HTRI", "Flare / relief tool", "SPID", "SI", "SPEL", "S3D", "SPF", "CAESAR II", "Tekla",
           "ETAP", "Mechanical calc tool", "PHA tool", "Document Locator", "Purchasing DB", "Jovix", "P6",
           "iConstruct", "Procore", "Smart Completions", "DCS / SIS config", "Excel / Word"]
PHASES = ["FEL1", "FEL2", "FEL3", "DD", "CON", "COM"]
FORMS = ["data", "document", "drawing", "model", "decision"]
METHODS = ["integrated", "file", "manual", "review", "meeting"]
METHOD_NOTE = {
    "integrated": "system to system (live or publish/retrieve)",
    "file": "export / import of a file",
    "manual": "read and re-entered by hand",
    "review": "read and commented or approved; no data moves",
    "meeting": "workshop or coordination meeting",
}


GAPS = []   # needs an author could not match to any item: {"task", "from", "info", ...}


def load():
    tasks, items, errs = {}, {}, []
    GAPS.clear()
    for p in sorted(DATA.glob("*.json")):
        try:
            d = json.loads(p.read_text())
        except json.JSONDecodeError as e:
            errs.append(f"{p.name}: {e}")
            continue
        code = d.get("discipline")
        if code not in DISC:
            errs.append(f"{p.name}: unknown discipline {code!r}")
            continue
        for t in d.get("tasks", []):
            t["disc"] = code
            if t["id"] in tasks:
                errs.append(f"duplicate task {t['id']}")
            tasks[t["id"]] = t
        for g in d.get("gaps", []):
            GAPS.append(dict(g, by=code))
        for i in d.get("items", []):
            i["disc"] = code
            if i["id"] in items:
                errs.append(f"duplicate item {i['id']}")
            items[i["id"]] = i
    return tasks, items, errs


def activities():
    sys.path.insert(0, str(HERE.parent))
    try:
        from cfu.wrapup.phases import Plan
        p = Plan()
        return {i: p.A[i]["name"] for i in p.A}
    except Exception:
        return None


def check(tasks, items):
    errs, warns = [], []
    acts = activities()
    for t in tasks.values():
        c = t["disc"]
        if not re.fullmatch(rf"{c}-T\d{{3}}", t["id"]):
            errs.append(f"{t['id']}: task id should be {c}-T###")
        if t.get("system") not in SYSTEMS:
            errs.append(f"{t['id']}: unknown system {t.get('system')!r}")
        bad = [p for p in t.get("phases", []) if p not in PHASES]
        if bad or not t.get("phases"):
            errs.append(f"{t['id']}: phases {t.get('phases')}")
        if acts is not None and t.get("activity") not in acts:
            errs.append(f"{t['id']}: unknown activity {t.get('activity')!r}")
        for i in t.get("produces", []):
            if i not in items:
                errs.append(f"{t['id']} produces unknown item {i}")
            elif items[i]["owner_task"] != t["id"]:
                errs.append(f"{t['id']} produces {i}, but its owner is {items[i]['owner_task']}")
        for k, c_ in enumerate(t.get("consumes", [])):
            if c_.get("item") not in items:
                errs.append(f"{t['id']} consumes unknown item {c_.get('item')}")
            if c_.get("method") not in METHODS:
                errs.append(f"{t['id']} input {c_.get('item')}: method {c_.get('method')!r}")
            if c_.get("level") not in range(5):
                errs.append(f"{t['id']} input {c_.get('item')}: level {c_.get('level')!r}")
            if c_.get("item") in t.get("produces", []):
                errs.append(f"{t['id']} consumes its own output {c_.get('item')}")
    for i in items.values():
        c = i["disc"]
        if not re.fullmatch(rf"{c}-I\d{{3}}", i["id"]):
            errs.append(f"{i['id']}: item id should be {c}-I###")
        if i.get("system") not in SYSTEMS:
            errs.append(f"{i['id']}: unknown system {i.get('system')!r}")
        if i.get("form") not in FORMS:
            errs.append(f"{i['id']}: form {i.get('form')!r}")
        o = tasks.get(i.get("owner_task"))
        if not o:
            errs.append(f"{i['id']}: owner task {i.get('owner_task')} not found")
        elif o["disc"] != c:
            errs.append(f"{i['id']}: owner task {o['id']} is in another discipline")
        elif i["id"] not in o.get("produces", []):
            errs.append(f"{i['id']}: owner {o['id']} does not list it in produces")
    for g in GAPS:
        if g.get("task") not in tasks:
            errs.append(f"gap on unknown task {g.get('task')}")
        warns.append(f"{g.get('task')} needs from {g.get('from')}: {g.get('info')} (no item yet)")
    for t in tasks.values():
        if t.get("needs"):
            warns.append(f"{t['id']}: {len(t['needs'])} needs not yet linked to items")
    used = {c["item"] for t in tasks.values() for c in t.get("consumes", [])}
    for i in items.values():
        if i["id"] not in used:
            warns.append(f"{i['id']} {i['name']}: produced but no task uses it")
    for t in tasks.values():
        if not t.get("consumes") and t["disc"] != "OWNR":
            warns.append(f"{t['id']} {t['name']}: no inputs")
        if not t.get("produces"):
            warns.append(f"{t['id']} {t['name']}: produces nothing")
    return errs, warns


def exchanges(tasks, items):
    out = []
    for t in tasks.values():
        for c in t.get("consumes", []):
            i = items.get(c["item"])
            if not i:
                continue
            src = tasks[i["owner_task"]]
            out.append(dict(item=i["id"], item_name=i["name"], form=i["form"], content=i.get("content", ""),
                            from_disc=i["disc"], from_task=src["id"], from_sys=i["system"],
                            to_disc=t["disc"], to_task=t["id"], to_sys=t["system"], use=c.get("use", ""),
                            method=c["method"], level=c["level"], phases=t["phases"]))
    return out


def order_key(x):
    return (list(DISC).index(x["disc"]), x["id"])


# ---------------------------------------------------------------------------------------------- outputs
def catalog(tasks, items):
    md = ["# Item catalog\n\nEvery information item, by owning discipline. Use these IDs in `consumes`.\n"]
    for c in DISC:
        its = sorted((i for i in items.values() if i["disc"] == c), key=lambda i: i["id"])
        if not its:
            continue
        md.append(f"\n## {c} {DISC[c]}\n\n| ID | Item | System | Form | Content |\n|---|---|---|---|---|\n")
        for i in its:
            md.append(f"| {i['id']} | {i['name']} | {i['system']} | {i['form']} | "
                      f"{i.get('content', '').replace('|', '/')} |\n")
    (OUT / "item-catalog.md").write_text("".join(md))


def workbook(tasks, items, ex, errs, warns):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    hf, hfont = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")

    def sheet(title, hdr, widths, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(hdr)
        for c in ws[1]:
            c.font, c.fill = hfont, hf
        for k, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(k + 1)].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}1"
        return ws

    ws = sheet("Exchanges", ["Item", "Information", "Form", "From discipline", "From task", "From system",
                             "To discipline", "To task", "To system", "What is used", "Method today", "Level needed",
                             "Phases"], (11, 34, 9, 12, 34, 16, 12, 34, 16, 50, 11, 8, 18), True)
    for e in sorted(ex, key=lambda e: (list(DISC).index(e["from_disc"]), e["item"], e["to_task"])):
        ws.append([e["item"], e["item_name"], e["form"], e["from_disc"],
                   f"{e['from_task']} {tasks[e['from_task']]['name']}", e["from_sys"], e["to_disc"],
                   f"{e['to_task']} {tasks[e['to_task']]['name']}", e["to_sys"], e["use"], e["method"], e["level"],
                   ", ".join(e["phases"])])
    ws = sheet("Tasks", ["ID", "Discipline", "Task", "System", "Phases", "Activity", "Description", "Inputs",
                         "Outputs"], (11, 8, 40, 16, 18, 9, 60, 7, 7))
    for t in sorted(tasks.values(), key=order_key):
        ws.append([t["id"], t["disc"], t["name"], t["system"], ", ".join(t["phases"]), t.get("activity", ""),
                   t.get("description", ""), len(t.get("consumes", [])), len(t.get("produces", []))])
    ws = sheet("Items", ["ID", "Discipline", "Information", "System", "Form", "Owner task", "Content", "Used by"],
               (11, 8, 40, 16, 9, 11, 70, 7))
    use = Counter(e["item"] for e in ex)
    for i in sorted(items.values(), key=order_key):
        ws.append([i["id"], i["disc"], i["name"], i["system"], i["form"], i["owner_task"], i.get("content", ""),
                   use[i["id"]]])
    codes = [c for c in DISC if any(t["disc"] == c for t in tasks.values())]
    ws = sheet("Discipline matrix", ["From \\ To"] + codes, [14] + [7] * len(codes))
    m = Counter((e["from_disc"], e["to_disc"]) for e in ex)
    for a in codes:
        ws.append([a] + [m[(a, b)] or "" for b in codes])
    ws.append([])
    ws.append(["Exchanges from the row discipline to the column discipline."])
    syss = [s for s in SYSTEMS if any(e["from_sys"] == s or e["to_sys"] == s for e in ex)]
    ws = sheet("System matrix", ["From \\ To"] + syss, [20] + [9] * len(syss))
    ms = Counter((e["from_sys"], e["to_sys"]) for e in ex)
    mm = Counter((e["from_sys"], e["to_sys"]) for e in ex if e["method"] == "manual")
    for a in syss:
        ws.append([a] + [(f"{ms[(a, b)]} ({mm[(a, b)]} manual)" if mm[(a, b)] else ms[(a, b)]) or ""
                         for b in syss])
    ws = sheet("Manual hand-offs", ["From system", "To system", "Item", "Information", "From task", "To task",
                                    "What is re-entered"], (16, 16, 11, 34, 34, 34, 60))
    for e in sorted((e for e in ex if e["method"] == "manual" and e["from_sys"] != e["to_sys"]),
                    key=lambda e: (e["from_sys"], e["to_sys"], e["item"])):
        ws.append([e["from_sys"], e["to_sys"], e["item"], e["item_name"],
                   f"{e['from_task']} {tasks[e['from_task']]['name']}",
                   f"{e['to_task']} {tasks[e['to_task']]['name']}", e["use"]])
    ws = sheet("Checks", ["Severity", "Message"], (10, 120))
    for s in errs:
        ws.append(["error", s])
    for s in warns:
        ws.append(["warning", s])
    for w in wb.worksheets:
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT / "Task-and-Interface-Register.xlsx")


def page(tasks, items, ex, errs, warns):
    T = {t["id"]: dict(d=t["disc"], n=t["name"], s=t["system"], p=t["phases"], a=t.get("activity", ""),
                       x=t.get("description", "")) for t in tasks.values()}
    I = {i["id"]: dict(d=i["disc"], n=i["name"], s=i["system"], f=i["form"], c=i.get("content", ""),
                       o=i["owner_task"]) for i in items.values()}
    E = [[e["item"], e["to_task"], e["use"], METHODS.index(e["method"]), e["level"]] for e in ex]
    data = dict(disc=DISC, systems=SYSTEMS, phases=PHASES, methods=METHODS, mnote=METHOD_NOTE, T=T, I=I, E=E,
                order=[t["id"] for t in sorted(tasks.values(), key=order_key)], warns=len(warns), errs=errs)
    tpl = (HERE / "page.html").read_text()
    js = json.dumps(data, separators=(",", ":")).replace("</", "<\\/")
    (OUT / "interfaces.html").write_text(tpl.replace("/*DATA*/null", js))


def main():
    OUT.mkdir(exist_ok=True)
    tasks, items, errs = load()
    e2, warns = check(tasks, items)
    errs += e2
    if "--check" in sys.argv:       # validation only, no outputs (safe to run concurrently)
        only = next((a for a in sys.argv[1:] if a in DISC), None)
        mine = [s for s in errs if not only or s.startswith(only) or f" {only}-" in s]
        print(f"{len(mine)} errors" + (f" for {only}" if only else ""))
        for s in mine:
            print("  ERROR", s)
        return mine
    ex = exchanges(tasks, items)
    catalog(tasks, items)
    workbook(tasks, items, ex, errs, warns)
    if (HERE / "page.html").exists():
        page(tasks, items, ex, errs, warns)
    by = Counter(t["disc"] for t in tasks.values())
    print(f"{len(tasks)} tasks, {len(items)} items, {len(ex)} exchanges; " +
          ", ".join(f"{c} {by[c]}" for c in DISC if by[c]))
    print(f"methods: {dict(Counter(e['method'] for e in ex))}")
    print(f"{len(errs)} errors, {len(warns)} warnings")
    for s in errs[:40]:
        print("  ERROR", s)
    return errs


if __name__ == "__main__":
    bad = [a for a in sys.argv[1:] if a != "--check" and a not in DISC]
    if bad:
        print(__doc__)
        sys.exit(2)
    sys.exit(1 if main() else 0)
