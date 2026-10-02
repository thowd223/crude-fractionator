"""Generic datasheet page renderer (reportlab canvas) and xlsx writer.

A datasheet item = dict(tag, title, sections=[...]); a section is
    ("kv", "TITLE", [(label, value), ...])            -> two-column numbered lines (API/TEMA style)
    ("tbl", "TITLE", [hdr...], [[row...], ...], [col widths])
    ("note", "TITLE", [lines])
"""
from __future__ import annotations

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from .. import basis

NAVY = colors.HexColor("#1F3864")
SHADE = colors.HexColor("#E4EAF4")
PW, PH = A4
M = 10 * mm
RH = 4.15 * mm


def _fit(c, text, x, y, w, font="Helvetica", size=6.8, align="left"):
    text = "" if text is None else str(text)
    s = size
    while s > 4.6 and stringWidth(text, font, s) > w:
        s -= 0.2
    while stringWidth(text, font, s) > w and len(text) > 3:
        text = text[:-2] + "~"
    c.setFont(font, s)
    if align == "left":
        c.drawString(x, y, text)
    elif align == "center":
        c.drawCentredString(x + w / 2, y, text)
    else:
        c.drawRightString(x + w, y, text)


class DSWriter:
    def __init__(self, path, doc_no, cls_title, standard):
        self.c = canvas.Canvas(str(path), pagesize=A4)
        self.c.setTitle(f"{doc_no} {cls_title} datasheets")
        self.c.setAuthor("CLAUDE CODE")
        self.doc_no, self.cls, self.std = doc_no, cls_title, standard
        self.page = 0
        self.items = []

    def _header(self, item, cont):
        c = self.c
        P = basis.PROJECT
        x0, y1 = M, PH - M
        w = PW - 2 * M
        c.setLineWidth(0.8)
        c.rect(x0, M, w, PH - 2 * M)
        h = 22 * mm
        c.setLineWidth(0.5)
        c.line(x0, y1 - h, x0 + w, y1 - h)
        c.line(x0 + 55 * mm, y1, x0 + 55 * mm, y1 - h)
        c.line(x0 + w - 55 * mm, y1, x0 + w - 55 * mm, y1 - h)
        c.setFillColor(colors.black)
        _fit(c, P["client"].upper(), x0 + 2 * mm, y1 - 5 * mm, 51 * mm, "Helvetica", 6.5)
        _fit(c, P["name"], x0 + 2 * mm, y1 - 9.5 * mm, 51 * mm, "Helvetica-Bold", 7)
        _fit(c, f"Stage: {P['stage']}  -  Discipline: MECHANICAL", x0 + 2 * mm, y1 - 14 * mm, 51 * mm, "Helvetica", 6.2)
        _fit(c, f"Unit area: {item.get('area', '-')}", x0 + 2 * mm, y1 - 18.5 * mm, 51 * mm, "Helvetica", 6.2)
        mx = x0 + 55 * mm
        mw = w - 110 * mm
        _fit(c, f"{self.cls.upper()} DATASHEET", mx, y1 - 7 * mm, mw, "Helvetica-Bold", 10, "center")
        _fit(c, self.std, mx, y1 - 11.5 * mm, mw, "Helvetica", 6.8, "center")
        _fit(c, f"{item['tag']}  -  {item['title']}" + ("  (continued)" if cont else ""), mx, y1 - 17.5 * mm, mw,
             "Helvetica-Bold", 8.5, "center")
        rx = x0 + w - 55 * mm
        rows = [("DOC. No.", self.doc_no), ("REV / DATE", f"{P['rev']} / 2026-10-02"), ("STATUS", P["rev_desc"]),
                ("PREPARED", "CLAUDE CODE   CHK: -   APP: -"), ("SHEET", f"{self.page}")]
        for i, (a, b) in enumerate(rows):
            yy = y1 - (i + 1) * 4.2 * mm + 1.1 * mm
            _fit(c, a, rx + 1.5 * mm, yy, 15 * mm, "Helvetica", 5.6)
            _fit(c, b, rx + 17 * mm, yy, 37 * mm, "Helvetica-Bold" if i == 0 else "Helvetica", 6.4)
            if i:
                c.setLineWidth(0.2)
                c.line(rx, y1 - i * 4.2 * mm, x0 + w, y1 - i * 4.2 * mm)
        self.sheet_ref = rows
        return y1 - h

    def _new_page(self, item, cont=False):
        if self.page:
            self.c.showPage()
        self.page += 1
        self.y = self._header(item, cont) - 1.5 * mm
        self.line_no = getattr(self, "line_no", 0) if cont else 0

    def _bar(self, title):
        c = self.c
        c.setFillColor(NAVY)
        c.rect(M, self.y - RH, PW - 2 * M, RH, stroke=0, fill=1)
        c.setFillColor(colors.white)
        _fit(c, title.upper(), M + 2 * mm, self.y - RH + 1.25 * mm, PW - 2 * M - 4 * mm, "Helvetica-Bold", 7)
        c.setFillColor(colors.black)
        self.y -= RH

    def add(self, item):
        self.items.append(item)
        self._new_page(item)
        for sec in item["sections"]:
            kind, title = sec[0], sec[1]
            if kind == "kv":
                rows = sec[2]
                nrow = (len(rows) + 1) // 2
                if self.y - RH * (nrow + 1) < M + 12 * mm and nrow < 40:
                    self._new_page(item, True)
                self._bar(title)
                half = (PW - 2 * M) / 2
                i = 0
                while i < len(rows):
                    if self.y - RH < M + 10 * mm:
                        self._new_page(item, True)
                        self._bar(title + " (cont.)")
                    for col in range(2):
                        if i >= len(rows):
                            break
                        lab, val = rows[i]
                        x = M + col * half
                        self.line_no += 1
                        c = self.c
                        if (self.line_no // 2) % 2 == 0:
                            c.setFillColor(SHADE)
                            c.rect(x, self.y - RH, half, RH, stroke=0, fill=1)
                            c.setFillColor(colors.black)
                        _fit(c, f"{self.line_no}", x + 0.6 * mm, self.y - RH + 1.3 * mm, 5 * mm, "Helvetica", 5.2)
                        _fit(c, lab, x + 6 * mm, self.y - RH + 1.3 * mm, half * 0.47, "Helvetica", 6.4)
                        vb = str(val) if val is not None else "-"
                        font = "Helvetica-Oblique" if vb.upper().startswith(("TBD", "VENDOR", "BY VENDOR")) else "Helvetica-Bold"
                        _fit(c, vb, x + 6 * mm + half * 0.48, self.y - RH + 1.3 * mm, half * 0.5 - 7 * mm, font, 6.6)
                        i += 1
                    c.setStrokeColor(colors.HexColor("#9AA5B8"))
                    c.setLineWidth(0.15)
                    c.line(M, self.y - RH, PW - M, self.y - RH)
                    c.line(M + half, self.y, M + half, self.y - RH)
                    c.setStrokeColor(colors.black)
                    self.y -= RH
                self.y -= 1.2 * mm
            elif kind == "tbl":
                hdr, rows, widths = sec[2], sec[3], sec[4]
                tot = sum(widths)
                W = PW - 2 * M
                cw = [W * w_ / tot for w_ in widths]
                if self.y - RH * min(len(rows) + 2, 8) < M + 12 * mm:
                    self._new_page(item, True)
                self._bar(title)

                def hdr_row():
                    c = self.c
                    c.setFillColor(SHADE)
                    c.rect(M, self.y - RH, W, RH, stroke=0, fill=1)
                    c.setFillColor(colors.black)
                    x = M
                    for h_, w_ in zip(hdr, cw):
                        _fit(c, h_, x + 0.8 * mm, self.y - RH + 1.3 * mm, w_ - 1.6 * mm, "Helvetica-Bold", 6.2)
                        x += w_
                    self.y -= RH

                hdr_row()
                for r in rows:
                    if self.y - RH < M + 10 * mm:
                        self._new_page(item, True)
                        self._bar(title + " (cont.)")
                        hdr_row()
                    x = M
                    for v, w_ in zip(r, cw):
                        _fit(self.c, v, x + 0.8 * mm, self.y - RH + 1.3 * mm, w_ - 1.6 * mm, "Helvetica", 6.3)
                        self.c.setLineWidth(0.15)
                        self.c.line(x, self.y, x, self.y - RH)
                        x += w_
                    self.c.line(M, self.y - RH, PW - M, self.y - RH)
                    self.y -= RH
                self.y -= 1.2 * mm
            elif kind == "note":
                lines = sec[2]
                if self.y - RH * (len(lines) + 1) < M + 10 * mm:
                    self._new_page(item, True)
                self._bar(title)
                for ln in lines:
                    if self.y - RH < M + 10 * mm:
                        self._new_page(item, True)
                    _fit(self.c, ln, M + 2 * mm, self.y - RH + 1.3 * mm, PW - 2 * M - 4 * mm, "Helvetica", 6.3)
                    self.y -= RH * 0.95
                self.y -= 1.2 * mm

    def save(self):
        self.c.showPage()
        self.c.save()
        return self.page


def write_xlsx(path, classes):
    """classes: list of (sheet_name, doc_no, items).  One sheet per class; parameters x tags."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="A0A8B8")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    sec_fill = PatternFill("solid", fgColor="E4EAF4")
    P = basis.PROJECT
    idx = wb.create_sheet("Index")
    idx.append([f"{P['name']} - Mechanical equipment datasheets (combined)", "", ""])
    idx.append([f"Rev {P['rev']} {P['rev_desc']} 2026-10-02, prepared CLAUDE CODE", "", ""])
    idx.append([])
    idx.append(["Sheet", "Datasheet document", "Items"])
    for c in idx[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
    idx["A1"].font = Font(bold=True, size=13)
    for name, doc, items in classes:
        idx.append([name, doc, ", ".join(i["tag"] for i in items)])
        ws = wb.create_sheet(name[:31])
        tags = [i["tag"] for i in items]
        ws.append([f"{doc} - {name}", "", *[""] * len(tags)])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append(["Section", "Parameter", *tags])
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        # union of (section, label) keeping order
        order = []
        vals = {}
        for it in items:
            for sec in it["sections"]:
                if sec[0] != "kv":
                    continue
                for lab, v in sec[2]:
                    k = (sec[1], lab)
                    if k not in vals:
                        order.append(k)
                        vals[k] = {}
                    vals[k][it["tag"]] = v
        last = None
        for k in order:
            sec, lab = k
            if sec != last:
                ws.append([sec.upper()])
                r = ws.max_row
                for col in range(1, len(tags) + 3):
                    ws.cell(r, col).fill = sec_fill
                ws.cell(r, 1).font = Font(bold=True)
                last = sec
            row = ["", lab]
            for t in tags:
                v = vals[k].get(t, "")
                row.append(v if isinstance(v, (int, float)) else ("" if v is None else str(v)))
            ws.append(row)
        for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(tags) + 2):
            for c in r:
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if c.column > 2:
                    c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 40
        for i in range(len(tags)):
            ws.column_dimensions[get_column_letter(i + 3)].width = 22
        ws.freeze_panes = "C3"
        # tables (e.g. nozzle schedules) on their own block below
        for it in items:
            for sec in it["sections"]:
                if sec[0] != "tbl":
                    continue
                ws.append([])
                ws.append([f"{it['tag']} - {sec[1]}"])
                ws.cell(ws.max_row, 1).font = Font(bold=True)
                ws.append(list(sec[2]))
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
                    c.fill = sec_fill
                for r in sec[3]:
                    ws.append([str(x) for x in r])
    idx.column_dimensions["A"].width = 22
    idx.column_dimensions["B"].width = 34
    idx.column_dimensions["C"].width = 110
    wb.save(path)
