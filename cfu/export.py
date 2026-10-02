"""Export master data (JSON) and Phase-0 deliverables (HMB workbook, equipment list, reports)."""
from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import assay, basis
from .hmb import Model
from .thermo import STEAM, enthalpy, flash_TP

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
DLV = ROOT / "deliverables"


def _clean(o):
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items() if not isinstance(v, (np.ndarray,)) and not callable(v)
                and k not in ("model",)}
    if isinstance(o, (list, tuple)):
        return [_clean(x) for x in o]
    if isinstance(o, (np.floating,)):
        return round(float(o), 4)
    if isinstance(o, float):
        return round(o, 4)
    if isinstance(o, (np.integer,)):
        return int(o)
    return o


def stream_props(m: Model, s) -> dict:
    sl = m.sl
    hc = float(s.m.sum())
    d = dict(no=s.no, name=s.name, frm=s.frm, to=s.to, T_C=round(s.T, 1), P_barg=round(s.P - 1.013, 2),
             P_bara=round(s.P, 3), phase=s.phase, total_kg_h=round(s.total(), 0), hc_kg_h=round(hc, 0),
             water_kg_h=round(s.water + s.steam, 0))
    if hc > 0:
        f = flash_TP(sl, s.m, s.T, s.P, s.steam)
        vol15 = m.vol(s.m)
        d.update(std_m3h=round(vol15, 1), bpsd=round(assay.bpsd(sl, s.m), 0), api=round(assay.api(sl, s.m), 1),
                 mw=round(m.mw(s.m), 1), vf_mass=round(f.vf_mass, 3))
        if f.liq.sum() > 0:
            rho_l = float(f.liq.sum() / (f.liq / sl.rho_l(s.T)).sum())
            d.update(liq_kg_h=round(float(f.liq.sum()), 0), rho_liq=round(rho_l, 1),
                     liq_act_m3h=round(float(f.liq.sum()) / rho_l, 1))
        if f.vap.sum() + s.steam > 0:
            mwv = (f.vap.sum() + s.steam) / ((f.vap / sl.MW).sum() + s.steam / STEAM["MW"])
            rho_v = s.P * 1e5 * mwv / (8314.46 * (s.T + 273.15))
            d.update(vap_kg_h=round(float(f.vap.sum() + s.steam), 0), mw_vap=round(mwv, 1), rho_vap=round(rho_v, 3),
                     vap_act_m3h=round(float(f.vap.sum() + s.steam) / rho_v, 0))
        d["enthalpy_kW"] = round(enthalpy(sl, f.liq, f.vap, s.T, s.steam, s.water), 0)
        tbp = assay.tbp_points(sl, s.m, (5, 50, 95))
        d.update(tbp5=round(tbp[5]), tbp50=round(tbp[50]), tbp95=round(tbp[95]),
                 sulfur_wt=round(float((s.m * sl.S).sum() / hc * 100), 3))
    else:
        d.update(std_m3h=round((s.water + s.steam) / 999, 1) if not s.gas else "", liq_kg_h=round(s.water, 0),
                 vap_kg_h=round(s.steam + s.gas, 0))
        if s.gas or s.steam:
            n = s.steam / STEAM["MW"] + s.gas / s.gas_mw
            mwv = (s.steam + s.gas) / n
            rho_v = s.P * 1e5 * mwv / (8314.46 * (s.T + 273.15))
            d.update(mw_vap=round(mwv, 1), rho_vap=round(rho_v, 4), vap_act_m3h=round((s.steam + s.gas) / rho_v, 0),
                     vf_mass=1.0)
    d["gas_kg_h"] = round(s.gas, 1)
    return d


def write_json(name, obj):
    DATA.mkdir(exist_ok=True)
    (DATA / name).write_text(json.dumps(_clean(obj), indent=1, default=str))


# ---------------------------------------------------------------------------
HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def _sheet(wb, title, rows, header, widths=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, size=9)
        r0 = 3
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), HDR, Alignment(wrap_text=True,
                                                                                               vertical="center"), BOX
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.border = BOX
            if isinstance(v, float):
                c.number_format = "#,##0.0" if abs(v) < 1000 else "#,##0"
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths[j - 1] if widths else 14)
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    return ws


def hmb_workbook(m: Model, streams: list[dict], path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    # stream table transposed (streams as columns) - classic HMB layout
    ws = wb.create_sheet("Stream Table")
    keys = [("no", "Stream No."), ("name", "Description"), ("frm", "From"), ("to", "To"), ("phase", "Phase"),
            ("T_C", "Temperature, C"), ("P_barg", "Pressure, barg"), ("total_kg_h", "Total mass, kg/h"),
            ("hc_kg_h", "Hydrocarbon, kg/h"), ("water_kg_h", "Water/steam, kg/h"), ("vf_mass", "Vapour frac. (mass)"),
            ("std_m3h", "Std liq. vol (15C), m3/h"), ("bpsd", "Std liq. vol, BPSD"), ("api", "API gravity"),
            ("mw", "Molecular weight"), ("liq_kg_h", "Liquid, kg/h"), ("rho_liq", "Liq. density @T, kg/m3"),
            ("liq_act_m3h", "Liq. actual, m3/h"), ("vap_kg_h", "Vapour, kg/h"), ("mw_vap", "Vapour MW"),
            ("rho_vap", "Vap. density, kg/m3"), ("vap_act_m3h", "Vap. actual, m3/h"), ("enthalpy_kW", "Enthalpy, kW"),
            ("tbp5", "TBP 5%, C"), ("tbp50", "TBP 50%, C"), ("tbp95", "TBP 95%, C"), ("sulfur_wt", "Sulfur, wt%")]
    ws.cell(1, 1, f"{basis.PROJECT['name']} - Heat & Material Balance, Design Case ({basis.CRUDE['name']}), "
                  f"Rev {basis.PROJECT['rev']}").font = Font(bold=True, size=12)
    for i, (k, lbl) in enumerate(keys, 3):
        c = ws.cell(i, 1, lbl)
        c.font, c.fill, c.border = Font(bold=True), SUB, BOX
        for j, s in enumerate(streams, 2):
            v = s.get(k, "")
            c = ws.cell(i, j, v)
            c.border = BOX
            c.alignment = Alignment(horizontal="center", wrap_text=k in ("name", "frm", "to"))
            if k == "no":
                c.font, c.fill = Font(bold=True, color="FFFFFF"), HDR
    ws.column_dimensions["A"].width = 26
    for j in range(2, len(streams) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.row_dimensions[4].height = 45
    ws.freeze_panes = "B4"

    # component flows
    names = m.sl.names()
    rows = []
    for i, n in enumerate(names):
        c = m.sl.c[i]
        rows.append([n, round(c.Tb_C, 1), round(c.SG, 4), round(c.MW, 1), round(c.Tc - 273.15, 1), round(c.Pc, 2),
                     round(c.omega, 3), round(c.K, 2), round(c.lv_frac * 100, 3), round(c.sulfur * 100, 3)]
                    + [round(float(m.S[s["no"]].m[i]), 1) for s in streams])
    _sheet(wb, "Component Flows", rows,
           ["Component", "NBP C", "SG", "MW", "Tc C", "Pc bar", "omega", "Watson K", "LV% crude", "S wt%"]
           + [f"S{s['no']} kg/h" for s in streams],
           note="Pseudo-components from TBP; Riazi-Daubert 1987 critical properties; Lee-Kesler omega")

    # products summary
    prods = [("Off-gas (atm drum)", "9"), ("Stabiliser off-gas", "29"), ("LPG", "30"), ("Light naphtha", "32"),
             ("Heavy naphtha", "33"), ("Kerosene", "16"), ("Diesel", "17"), ("AGO", "18"), ("LVGO", "23"),
             ("HVGO", "24"), ("Slop wax", "25"), ("Vacuum residue", "26")]
    crude = m.crude.sum()
    rows = []
    for n, no in prods:
        s = next(x for x in streams if x["no"] == no)
        rows.append([n, no, s["hc_kg_h"], round(s["hc_kg_h"] / crude * 100, 2), s.get("bpsd", 0),
                     round(s.get("bpsd", 0) / basis.CAPACITY_BPSD * 100, 2), s.get("api", ""), s.get("tbp5", ""),
                     s.get("tbp50", ""), s.get("tbp95", ""), s.get("sulfur_wt", ""), s["to"]])
    rows.append(["TOTAL", "", round(sum(r[2] for r in rows), 0), round(sum(r[3] for r in rows), 2),
                 round(sum(r[4] for r in rows), 0), round(sum(r[5] for r in rows), 2), "", "", "", "", "", ""])
    _sheet(wb, "Product Yields", rows, ["Product", "Stream", "kg/h", "wt% crude", "BPSD", "LV% crude", "API",
                                        "TBP5 C", "TBP50 C", "TBP95 C", "S wt%", "Destination"],
           widths=[22, 8, 12, 10, 10, 10, 8, 8, 8, 8, 8, 24])

    # energy summary
    a, v = m.res["atm"], m.res["vac"]
    rows = [["H-101 absorbed (process + SS coil)", m.res["H-101"]["Q_abs_kw"] / 1000],
            ["H-101 fired", m.res["H-101"]["Q_fired_kw"] / 1000],
            ["H-201 absorbed", m.res["H-201"]["Q_abs_kw"] / 1000], ["H-201 fired", m.res["H-201"]["Q_fired_kw"] / 1000],
            ["C-101 total heat removal", a["Q_tot"] / 1000], ["C-101 overhead condensing", a["Q_cond"] / 1000],
            ["TPA duty", a["pa"]["TPA"]["duty_kw"] / 1000], ["MPA duty", a["pa"]["MPA"]["duty_kw"] / 1000],
            ["BPA duty", a["pa"]["BPA"]["duty_kw"] / 1000],
            ["C-201 total heat removal", v["Q_tot"] / 1000], ["LVGO PA duty", v["pa"]["LVGO"]["duty_kw"] / 1000],
            ["HVGO PA duty", v["pa"]["HVGO"]["duty_kw"] / 1000],
            ["Preheat recovered to crude", sum(e["Q_kw"] for e in m.res["preheat"]["exch"]) / 1000],
            ["Crude inlet temperature (CIT), C", m.res["preheat"]["CIT"]],
            ["Stabiliser reboiler", m.res["stab"]["Q_reb"] / 1000], ["Splitter reboiler", m.res["split"]["Q_reb"] / 1000]]
    rows = [[r[0], round(float(r[1]), 2)] for r in rows]
    _sheet(wb, "Energy Summary", rows, ["Item", "MW (or C)"], widths=[40, 14])

    # exchanger duties
    rows = [[e["tag"], e["hot"], round(e["Q_kw"] / 1000, 2), round(e["Th_in"], 1), round(e["Th_out"], 1),
             round(e["Tc_in"], 1), round(e["Tc_out"], 1), e["train"]] for e in m.res["preheat"]["exch"]]
    rows += [[t["tag"], t["hot"], round(t["Q_kw"] / 1000, 2), round(t["T_in"], 1), round(t["T_out"], 1), "", "",
              t["type"]] for t in m.res["preheat"]["trims"]]
    _sheet(wb, "Preheat Train", rows, ["Tag", "Hot stream", "Duty MW", "Hot in C", "Hot out C", "Crude in C",
                                       "Crude out C", "Train / type"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def equipment_workbook(sz: dict, path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    rows = []
    for e in sz["equipment"]:
        rows.append([e["tag"], e["type"], e["service"], e.get("area", ""), e.get("size", ""),
                     str(e.get("op_P", "")), str(e.get("op_T", "")), str(e.get("des_P", "")), str(e.get("des_T", "")),
                     e.get("moc", ""), e.get("ca_mm", ""),
                     round(e["duty_kw"] / 1000, 2) if e.get("duty_kw") else "", e.get("motor_kw") or ""])
    _sheet(wb, "Equipment List", rows, ["Tag", "Type", "Service", "Area", "Size / rating", "Op. P barg",
                                        "Op. T C", "Des. P barg", "Des. T C", "Material", "CA mm", "Duty MW",
                                        "Motor kW (each)"],
           widths=[11, 13, 38, 7, 52, 14, 14, 11, 9, 40, 7, 9, 10])
    rows = [[p["tag"], p["service"], round(p["flow_m3h"], 1), round(p["head_m"], 0), round(p["dP_bar"], 1),
             p["T"] if isinstance(p["T"], (int, float)) else p["T"], round(p["rho"], 0), round(p["eta"], 2),
             round(p["absorbed_kw"], 1), p["motor_kw"], p["api610"], p["seal"], p["moc"]] for p in sz["pumps"]]
    _sheet(wb, "Pump Schedule", rows, ["Tag", "Service", "Rated m3/h", "Head m", "dP bar", "T C", "rho kg/m3",
                                       "Eff.", "Abs. kW", "Motor kW", "API 610", "Seal plan", "API 610 matl."],
           widths=[10, 34, 10, 8, 7, 7, 9, 6, 8, 9, 8, 22, 12])
    rows = [[p["tag"], p["protects"], p["case"], round(p["load_kg_h"], 0), p["set_barg"], round(p["T"], 0),
             round(p["MW"], 1), round(p["area_mm2"], 0), f"{p['count']} x {p['orifice']}", p["dest"]]
            for p in sz["psv"]]
    _sheet(wb, "PSV Schedule", rows, ["Tag", "Protects", "Governing case", "Relief kg/h", "Set barg", "Relief T C",
                                      "MW", "Req. area mm2", "Orifice (API 526)", "Discharge to"],
           widths=[10, 14, 44, 11, 9, 9, 7, 12, 14, 12])
    rows = [[h["tag"], h["service"], h.get("tema", ""), round(h["duty_kw"] / 1000, 2), round(h.get("area_m2", 0), 0),
             h.get("n_shells", ""), h.get("U", ""), round(h.get("F", 0.9), 2) if h.get("F") else "", h["op_T"],
             h.get("shellside", ""), h.get("tubeside", ""), h["moc"]] for h in sz["hx"]]
    _sheet(wb, "Exchanger Schedule", rows, ["Tag", "Service", "TEMA", "Duty MW", "Area m2", "Shells", "U W/m2K",
                                            "F", "Temps C", "Shell side", "Tube side", "Material"],
           widths=[8, 34, 18, 8, 8, 7, 8, 6, 30, 22, 22, 30])
    rows = [[a["tag"], a["service"], round(a["duty_kw"] / 1000, 2), round(a["area_m2"], 0), a["bays"], a["fans"],
             a["motor_kw"], a["op_T"]] for a in sz["aircoolers"]]
    _sheet(wb, "Air Cooler Schedule", rows, ["Tag", "Service", "Duty MW", "Bare area m2", "Bays", "Fans",
                                             "Fan motor kW", "Temps C"], widths=[8, 34, 9, 12, 6, 6, 12, 34])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def export_all(m: Model, sz: dict):
    streams = [stream_props(m, s) for s in sorted(m.S.values(), key=lambda s: int(s.no))]
    write_json("streams.json", streams)
    write_json("equipment.json", sz["equipment"])
    write_json("psv.json", sz["psv"])
    a, v = m.res["atm"], m.res["vac"]
    key = dict(
        crude=dict(bpsd=basis.CAPACITY_BPSD, kg_h=float(m.crude.sum()), m3h=a["crude_m3h"], api=basis.CRUDE["api"],
                   K=m.sl.K0),
        atm=dict(T_fz=a["T_fz"], cot=a["cot"], P_fz_barg=a["P_fz"] - 1.013, T_top=a["T_top"], T_bot=a["T_bot"],
                 P_top_barg=basis.DESIGN["atm_top_P"] - 1.013, drum_P_barg=basis.DESIGN["atm_drum_P"] - 1.013,
                 T_draw=a["T_draw"], T_strip_out=a["T_out"], reflux_kg_h=a["R"], Q_tot_kw=a["Q_tot"],
                 Q_cond_kw=a["Q_cond"], T_water_dew=a["T_wdew"], steam=a["steam"], tray=a["tray"],
                 pa={k: {kk: vv for kk, vv in x.items() if kk != "comp"} for k, x in a["pa"].items()},
                 L_int=a["L_int"], sections=a["sections"]),
        vac=dict(T_fz=v["T_fz"], cot=v["cot"], P_fz_mbar=v["P_fz"] * 1000, P_top_mbar=v["P_top"] * 1000,
                 T_top=v["T_top"], T_bot=v["T_bot"], T_lvgo=v["T_lvgo"], T_hvgo=v["T_hvgo"], T_slop=v["T_slop"],
                 steam=v["steam"], Q_tot_kw=v["Q_tot"],
                 pa={k: {kk: vv for kk, vv in x.items() if kk != "comp"} for k, x in v["pa"].items()},
                 sections=v["sections"]),
        ejector=m.res["ejector"],
        stab={k: v_ for k, v_ in m.res["stab"].items() if not isinstance(v_, (np.ndarray,)) and k not in ("drum",)},
        split={k: v_ for k, v_ in m.res["split"].items() if not isinstance(v_, (np.ndarray,)) and k not in ("drum",)},
        preheat=dict(CIT=m.res["preheat"]["CIT"], T_desalter=m.res["preheat"]["T_desalter"],
                     wash_water=m.res["preheat"]["wash_water"], exch=m.res["preheat"]["exch"],
                     trims=m.res["preheat"]["trims"]),
        heaters={k: m.res[k] for k in ("H-101", "H-201")},
        warnings=m.warn,
        units="T degC; P_out, P_top, P_bot, P_fz (atm) in bar(a) unless key says barg/mbar; duties kW; flows kg/h",
    )
    write_json("process_results.json", key)
    hmb_workbook(m, streams, DLV / "00-basis" / "CFU-000-PR-HMB-001_Heat-Material-Balance.xlsx")
    equipment_workbook(sz, DLV / "02-equipment" / "CFU-000-ME-LST-001_Equipment-List.xlsx")
    return streams
