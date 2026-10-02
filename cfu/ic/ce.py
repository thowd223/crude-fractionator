"""Cause & Effect matrix CFU-000-IC-CE-001 (xlsx + A1 grid drawing) and SIF list / SIL determination
(LOPA-style) CFU-000-IC-RPT-002 (md + pdf).

Initiators: the SIFs of cfu.control_loops (expanded per transmitter group), ESD push-buttons and confirmed
F&G events per fire zone. Set points are derived from process_results / psv / equipment data with the margins
stated in SETPOINT_BASIS. LOPA frequencies are FEED judgements (CCPS 'LOPA' typical values) to be confirmed
in the SIL assessment workshop.
"""
from __future__ import annotations

import math

from ..drawing.sheet import Sheet
from .common import GREY, OUT, SIS_RED, Pen, load, pr, sifs

TURNDOWN = 0.5
# ------------------------------------------------------------------------- actions (columns)
ACTIONS = [
    # key, tag, description, group
    ("A1", "XV-1021/1022", "H-101 main FG SSOVs close", "H-101"),
    ("A2", "XV-1230", "H-101 FG vent opens (DB&B)", "H-101"),
    ("A3", "XV-1026", "H-101 pilot gas SSOV close", "H-101"),
    ("A4", "PV-1021 / TIC-1020", "H-101 FG control to min / track", "H-101"),
    ("A5", "K-101/K-102", "H-101 fans to purge (keep running), damper open", "H-101"),
    ("A6", "XV-2006A/B", "H-201 FG SSOVs close", "H-201"),
    ("A7", "XV-2105", "H-201 FG vent opens (DB&B)", "H-201"),
    ("A8", "FV-2007", "H-201 coil steam to max (sweep coils)", "H-201"),
    ("A9", "XY-1007", "D-101A transformer trip", "DESALTER"),
    ("A10", "XY-1008", "D-101B transformer trip", "DESALTER"),
    ("A11", "XV-1001", "Crude charge XV close", "CHARGE"),
    ("A12", "P-101A/B", "Crude charge pumps trip", "CHARGE"),
    ("A13", "P-102A/B", "Desalted crude booster pumps trip", "CHARGE"),
    ("A14", "XV-1083", "AR to H-201 XV close", "VDU FEED"),
    ("A15", "P-112A/B", "AR / H-201 charge pumps trip", "VDU FEED"),
    ("A16", "XV-1093", "LPG product XV close", "LIGHT ENDS"),
    ("A17", "XV-1096", "C-105 reboiler steam XV close", "LIGHT ENDS"),
    ("A18", "EIV-1121", "P-112 suction ROSOV close", "ROSOV"),
    ("A19", "EIV-2041", "P-204 suction ROSOV close", "ROSOV"),
    ("A20", "P-204A/B", "VR pumps trip", "ROSOV"),
    ("A21", "HV-1290/2190", "Heater snuffing steam open (F&G / manual)", "FIRE"),
    ("A22", "F&G", "Zone beacons / sounders, PAGA alarm", "ALARM"),
    ("A23", "CCR", "CCR alarm (SIS first-out) / DCS", "ALARM"),
    ("A24", "OSBL", "Signal to OSBL fire-water pumps / deluge", "FIRE"),
]
AKEYS = [a[0] for a in ACTIONS]


def setpoints():
    R = pr()
    H1, H2 = R["heaters"]["H-101"], R["heaters"]["H-201"]
    psv = {p["protects"]: p for p in load("psv.json")}
    eq = {e["tag"]: e for e in load("equipment.json")}
    p1 = H1["flow"] / H1["passes"] / 1000
    p2 = H2["flow"] / H2["passes"] / 1000
    c105_op = float(str(eq["C-105"]["op_P"]).split()[0])
    c105_set = psv["C-105"]["set_barg"]
    return dict(
        pass1=(p1, 0.4 * p1), pass2=(p2, 0.4 * p2),
        cot1=(H1["T_out"], H1["T_out"] + 15), cot2=(H2["T_out"], H2["T_out"] + 12),
        c105=(c105_op, round(min(0.94 * c105_set, c105_set - 0.5), 1), c105_set),
    )


SETPOINT_BASIS = [
    "Pass flow LL = 40 % of design pass flow (below the 50 % turndown minimum with 10 % margin).",
    "COT HH = design COT + 15 °C (H-101) / + 12 °C (H-201); high alarm at + 8 °C (+ 6 °C).",
    "Burner FG pressure LL / HH per burner vendor stability curve (FEED: 0.15 / 2.2 barg with 2.0 barg max "
    "normal at design firing).",
    "Arch pressure HH = +2.5 mmH2O (normal -2.5 mmH2O draft).",
    "Vessel pressure HH = min(94 % of PSV set, set - 0.5 bar) - keeps >= 6 % margin to PSV lift.",
    "Level HH = 85 % of transmitter span (above HLA at 75 %); LL = 15 % (below LLA at 25 %).",
]


def initiators():
    S = sifs()
    sp = setpoints()
    p1n, p1ll = sp["pass1"]
    p2n, p2ll = sp["pass2"]
    c1, c1hh = sp["cot1"]
    c2, c2hh = sp["cot2"]
    c105op, c105hh, c105set = sp["c105"]
    H101 = ["A1", "A2", "A3", "A4", "A5", "A23"]
    H201 = ["A6", "A7", "A8", "A23"]
    rows = []
    for k in range(1, 9):
        rows.append(dict(sif="SIF-101", tag=f"FZLL-10{10 + k}", init=f"FT-10{10 + k}A/B/C",
                         desc=f"H-101 pass {k} flow low-low", vote="2oo3", sp=f"{p1ll:.1f} t/h",
                         normal=f"{p1n:.1f} t/h", rt="2 s", act=H101))
    rows += [
        dict(sif="SIF-102", tag="PZLL-1027", init="PT-1027A/B/C", desc="H-101 burner FG pressure low-low",
             vote="2oo3", sp="0.15 barg", normal="1.0-2.0 barg", rt="2 s", act=H101),
        dict(sif="SIF-103", tag="PZHH-1027", init="PT-1027A/B/C", desc="H-101 burner FG pressure high-high",
             vote="2oo3", sp="2.2 barg", normal="1.0-2.0 barg", rt="2 s", act=H101),
        dict(sif="SIF-104", tag="BZLL-1028", init="BS-1028 (16)", desc="H-101 loss of flame (all burners)",
             vote="per burner / all", sp="flame off", normal="flame on", rt="4 s (FFRT)", act=H101),
        dict(sif="SIF-105", tag="PZHH-1029", init="PT-1029A/B/C", desc="H-101 arch pressure high-high / ID fan loss",
             vote="2oo3", sp="+2.5 mmH2O", normal="-2.5 mmH2O", rt="3 s", act=H101),
        dict(sif="SIF-106", tag="TZHH-1020", init="TT-1020A/B/C", desc="H-101 COT high-high",
             vote="2oo3", sp=f"{c1hh:.0f} °C", normal=f"{c1:.0f} °C", rt="5 s", act=H101),
        dict(sif="SIF-107", tag="LZLL-1007", init="LT-1007B", desc="D-101A interface low-low (grid short)",
             vote="1oo1", sp="15 %", normal="50 %", rt="5 s", act=["A9", "A23"]),
        dict(sif="SIF-107", tag="LZLL-1008", init="LT-1008B", desc="D-101B interface low-low (grid short)",
             vote="1oo1", sp="15 %", normal="50 %", rt="5 s", act=["A10", "A23"]),
        dict(sif="SIF-108", tag="LZHH-1082", init="LT-1082B", desc="C-101 bottom level high-high",
             vote="1oo1", sp="85 %", normal="50 %", rt="10 s", act=["A11", "A12", "A23"]),
        dict(sif="SIF-109", tag="LZLL-1092", init="LT-1092B", desc="D-105 level low-low (gas blow-by to LPG)",
             vote="1oo1", sp="15 %", normal="50 %", rt="5 s", act=["A16", "A23"]),
        dict(sif="SIF-110", tag="PZHH-1091", init="PT-1091B", desc="C-105 pressure high-high",
             vote="1oo1", sp=f"{c105hh:.1f} barg", normal=f"{c105op:.1f} barg", rt="5 s", act=["A17", "A23"]),
    ]
    for k in range(1, 5):
        rows.append(dict(sif="SIF-201", tag=f"FZLL-200{k}", init=f"FT-200{k}A/B/C", desc=f"H-201 pass {k} flow low-low",
                         vote="2oo3", sp=f"{p2ll:.1f} t/h", normal=f"{p2n:.1f} t/h", rt="2 s", act=H201))
    rows += [
        dict(sif="SIF-202", tag="PZLL-2009", init="PT-2009A/B/C", desc="H-201 burner FG pressure low-low",
             vote="2oo3", sp="0.15 barg", normal="1.0-2.0 barg", rt="2 s", act=H201),
        dict(sif="SIF-202", tag="BZLL-2008", init="BS-2008 (6)", desc="H-201 loss of flame (all burners)",
             vote="per burner / all", sp="flame off", normal="flame on", rt="4 s (FFRT)", act=H201),
        dict(sif="SIF-203", tag="LZHH-2024", init="LT-2024B", desc="C-201 bottom level high-high",
             vote="1oo1", sp="85 %", normal="50 %", rt="10 s", act=["A14", "A15", "A23"]),
        dict(sif="SIF-204", tag="HS-1121", init="HS-1121 / BY-1121", desc="P-112 area fire - manual / F&G confirmed",
             vote="1oo2", sp="-", normal="-", rt="30 s (valve)", act=["A18", "A15", "A23"]),
        dict(sif="SIF-204", tag="HS-2041", init="HS-2041 / BY-2041", desc="P-204 area fire - manual / F&G confirmed",
             vote="1oo2", sp="-", normal="-", rt="30 s (valve)", act=["A19", "A20", "A23"]),
        dict(sif="SIF-901", tag="HS-9000", init="HS-9000 (CCR + field)", desc="Unit ESD-1 push-button",
             vote="1oo2", sp="-", normal="-", rt="2 s",
             act=["A1", "A2", "A3", "A4", "A5", "A6", "A7", "A8", "A9", "A10", "A11", "A12", "A13", "A14",
                  "A15", "A16", "A17", "A22", "A23"]),
    ]
    from .iolist import FG_ZONES
    for z, d, *_ in FG_ZONES:
        heat = z == "FZ-05"
        rows.append(dict(sif="F&G", tag=f"GZHH-{z[3:]}", init=f"GD-{z[3:]}xx", desc=f"Confirmed gas - {d}",
                         vote="2ooN", sp="20 / 50 % LEL", normal="0", rt="10 s",
                         act=(["A1", "A3", "A6"] if heat else []) + ["A22", "A23"]))
        fa = ["A21"] if heat else []
        if z == "FZ-03":
            fa += ["A18"]
        if z == "FZ-06":
            fa += ["A19", "A20"]
        rows.append(dict(sif="F&G", tag=f"BZHH-{z[3:]}", init=f"BD-{z[3:]}xx / HS-{z[3:]}xx",
                         desc=f"Confirmed fire - {d}", vote="2ooN / MCP", sp="flame", normal="-", rt="10 s",
                         act=fa + ["A22", "A23", "A24"] + (["A1", "A3", "A6"] if heat else [])))
    for r in rows:
        r["sil"] = S.get(r["sif"], {}).get("sil", "F&G" if r["sif"] == "F&G" else "")
    return rows


def motor_trips():
    """(motor, SIF, reason) for SIS hard-wired motor trip outputs (used by the I/O list)."""
    out = []
    for m in ("P-101A", "P-101B"):
        out.append((m, "SIF-108", "C-101 LZHH / ESD"))
    for m in ("P-102A", "P-102B"):
        out.append((m, "SIF-901", "ESD"))
    for m in ("P-112A", "P-112B"):
        out.append((m, "SIF-203", "C-201 LZHH / fire / ESD"))
    for m in ("P-204A", "P-204B"):
        out.append((m, "SIF-204", "fire"))
    return out


# ------------------------------------------------------------------------- LOPA
LOPA = [
    # sif, scenario, IE, IEF/yr, consequence cat, TMEL, IPLs [(name, PFD)], CMs [(name, p)]
    ("SIF-101", "Loss of pass flow (FV fails closed, coking, charge loss) -> tube overheating / rupture, firebox fire",
     "BPCS loop failure / charge pump trip", 0.1, "C4 single fatality", 1e-5,
     [("Low-flow alarm FAL + operator (>= 10 min)", 0.1)], [("Occupancy (heater area)", 0.25)]),
    ("SIF-102", "Flame-out on low FG pressure -> fuel accumulation, re-ignition explosion",
     "PV-1021 fails closed / FG supply loss", 0.1, "C4 single fatality", 1e-5,
     [("Low-pressure alarm + operator", 0.1)], [("Probability of delayed ignition", 0.1), ("Occupancy", 0.25)]),
    ("SIF-103", "High FG pressure -> flame lift-off / unstable flame -> flame-out, explosion",
     "PV-1021 fails open", 0.1, "C4 single fatality", 1e-5,
     [("High-pressure alarm + operator", 0.1)], [("Probability of flame-out given HH", 0.1), ("Occupancy", 0.25)]),
    ("SIF-104", "Loss of flame with fuel flowing -> firebox explosion",
     "Burner instability / air-fuel upset", 0.1, "C4 single fatality", 1e-5,
     [("Operator observation (no credit - too fast)", 1.0)],
     [("Probability of explosive accumulation", 0.1), ("Probability of ignition", 0.5), ("Occupancy", 0.25)]),
    ("SIF-105", "Positive firebox pressure (ID fan trip / damper closed) -> flue gas / flame release",
     "ID fan trip or damper failure", 0.2, "C3 serious injury", 1e-4,
     [("Draft alarm + operator", 0.1)], [("Occupancy (platform)", 0.1)]),
    ("SIF-106", "High COT -> coking / tube overheating, transfer-line overpressure",
     "TIC-1020 / PIC-1021 failure", 0.1, "C3 serious injury / major asset", 1e-4,
     [("High-temperature alarm + operator", 0.1)], [("Probability of tube rupture", 0.5)]),
    ("SIF-107", "Low interface -> water on electrodes, grid short / arcing, vapour generation, fire",
     "LIC-1007/1008 failure", 0.1, "C3 serious injury", 1e-4,
     [("Interface alarm + operator", 0.1)], [("Probability of ignition", 0.3)]),
    ("SIF-108", "C-101 overfill -> liquid into overhead / PSV liquid relief, flare carry-over",
     "LIC-1082 / FV-1083 failure, P-112 trip", 0.1, "C3 major environmental", 1e-4,
     [("High-level alarm + operator", 0.1)], [("Probability PSV relieves liquid", 0.5)]),
    ("SIF-109", "D-105 loss of level -> gas blow-by into LPG treating / rundown, overpressure",
     "LIC-1092 / FV-1093 fails open", 0.1, "C3 serious injury", 1e-4,
     [("Low-level alarm + operator", 0.1)], [("Probability of downstream LOPC", 0.3)]),
    ("SIF-110", "C-105 overpressure on reboiler upset / loss of cooling -> PSV-1005 lift, LPG release",
     "FIC-1096 failure / A-106 fan loss", 0.2, "C3 serious injury", 1e-4,
     [("PSV-1005 (sized for case)", 0.01)], [("Probability of ignition of flare release", 1.0)]),
    ("SIF-201", "H-201 loss of pass flow -> coking, tube rupture, fire",
     "FIC-2001..2004 failure / P-112 trip", 0.1, "C4 single fatality", 1e-5,
     [("Low-flow alarm + operator", 0.1)], [("Occupancy", 0.25)]),
    ("SIF-202", "H-201 low FG pressure / flame failure -> fuel accumulation, explosion",
     "PV-2006 failure / burner instability", 0.1, "C4 single fatality", 1e-5,
     [("Alarm + operator", 0.1)], [("Probability of delayed ignition", 0.1), ("Occupancy", 0.25)]),
    ("SIF-203", "C-201 overfill -> liquid into flash zone / wash bed, loss of vacuum, overpressure",
     "LIC-2024 / FV-2025 failure, P-204 trip", 0.1, "C3 major asset", 1e-4,
     [("High-level alarm + operator", 0.1)], [("Probability of damage", 0.5)]),
    ("SIF-204", "Fire at hot pump (> AIT) -> escalation; isolate inventory of C-101 / C-201 bottoms",
     "Pump seal failure with ignition", 0.01, "C4 single fatality", 1e-5,
     [("Fire-fighting (OSBL response)", 1.0)], [("Probability of escalation within 15 min", 0.1), ("Occupancy", 0.5)]),
    ("SIF-901", "Major emergency - unit-wide isolation (manual ESD)", "Operator-initiated (escalation)", None,
     "C4 single fatality", 1e-5, [], []),
]


def sil_from_rrf(rrf):
    if rrf <= 1:
        return "none required"
    if rrf <= 10:
        return "SIL a (BPCS / alarm)"
    if rrf <= 100:
        return "SIL 1"
    if rrf <= 1000:
        return "SIL 2"
    if rrf <= 10000:
        return "SIL 3"
    return "SIL 4 (redesign)"


def lopa():
    S = sifs()
    out = []
    for sif, scen, ie, ief, cat, tmel, ipls, cms in LOPA:
        srs = S.get(sif, {}).get("sil", "")
        if ief is None:
            out.append(dict(sif=sif, scenario=scen, ie=ie, ief=None, cat=cat, tmel=tmel, ipls=ipls, cms=cms,
                            mel=None, rrf=None, sil=f"{srs} (company standard, manual ESD - LOPA n/a)", srs=srs,
                            match=True, function=S.get(sif, {}).get("function", "")))
            continue
        pfd = math.prod(p for _, p in ipls) * math.prod(p for _, p in cms)
        mel = ief * pfd
        rrf = mel / tmel
        sil = sil_from_rrf(rrf)
        out.append(dict(sif=sif, scenario=scen, ie=ie, ief=ief, cat=cat, tmel=tmel, ipls=ipls, cms=cms,
                        mel=mel, rrf=rrf, sil=sil, srs=srs,
                        match=(sil.startswith(srs) or (srs == "SIL 1" and rrf <= 100 and rrf > 1)),
                        function=S.get(sif, {}).get("function", "")))
    return out


# ------------------------------------------------------------------------- outputs
def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Cause & Effect"
    thin = Side(style="thin", color="808080")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr = PatternFill("solid", fgColor="1F3864")
    red = PatternFill("solid", fgColor="F4CCCC")
    grn = PatternFill("solid", fgColor="D9EAD3")
    ws["A1"] = "CFU-000-IC-CE-001  CAUSE & EFFECT MATRIX - SIS / BMS / ESD / F&G (Rev A, FEED)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("X = trip / close, O = open, A = alarm, P = purge / safe state. Initiators de-energise to trip; "
                "manual reset in CCR after cause cleared. Set points per CFU-000-IC-RPT-002.")
    fixed = ["Item", "SIF", "Initiator tag", "Sensor(s)", "Description", "Voting", "Normal", "Trip set point",
             "SIL", "Resp. time"]
    hrow = 5
    for j, h in enumerate(fixed, 1):
        c = ws.cell(hrow, j, h)
        c.fill, c.font, c.border = hdr, Font(color="FFFFFF", bold=True), bd
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    for k, (key, tag, desc, grp) in enumerate(ACTIONS):
        j = len(fixed) + 1 + k
        ws.cell(hrow - 1, j, grp).alignment = Alignment(text_rotation=90, horizontal="center")
        c = ws.cell(hrow, j, f"{tag} - {desc}")
        c.alignment = Alignment(text_rotation=90, wrap_text=True, horizontal="center", vertical="bottom")
        c.fill, c.font, c.border = hdr, Font(color="FFFFFF", bold=True, size=8), bd
        ws.column_dimensions[get_column_letter(j)].width = 4.2
    ws.row_dimensions[hrow].height = 200
    for i, r in enumerate(rows, 1):
        rr = hrow + i
        vals = [i, r["sif"], r["tag"], r["init"], r["desc"], r["vote"], r["normal"], r["sp"], r["sil"], r["rt"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(rr, j, v)
            c.border = bd
        for k, key in enumerate(AKEYS):
            j = len(fixed) + 1 + k
            v = ""
            if key in r["act"]:
                v = "A" if key in ("A22", "A23") else "O" if key in ("A2", "A7", "A21", "A8") else \
                    "P" if key in ("A4", "A5") else "X"
            c = ws.cell(rr, j, v)
            c.border = bd
            c.alignment = Alignment(horizontal="center")
            if v == "X":
                c.fill = red
            elif v:
                c.fill = grn
    for j, w in enumerate([5, 8, 12, 18, 40, 10, 11, 12, 7, 10], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(hrow + 1, len(fixed) + 1)
    ws2 = wb.create_sheet("Set-point basis")
    for i, s in enumerate(SETPOINT_BASIS, 1):
        ws2.cell(i, 1, s)
    ws2.column_dimensions["A"].width = 120
    ws3 = wb.create_sheet("SIF LOPA")
    heads = ["SIF", "Function", "Scenario", "Initiating event", "IEF /yr", "Consequence", "TMEL /yr", "IPLs (PFD)",
             "Conditional modifiers", "MEL w/o SIF /yr", "Required RRF", "SIL (LOPA)", "SIL (SRS)", "Check"]
    for j, h in enumerate(heads, 1):
        c = ws3.cell(1, j, h)
        c.fill, c.font, c.border = hdr, Font(color="FFFFFF", bold=True), bd
    for i, L in enumerate(lopa(), 2):
        vals = [L["sif"], L["function"], L["scenario"], L["ie"], L["ief"], L["cat"], L["tmel"],
                "; ".join(f"{n} ({p:g})" for n, p in L["ipls"]), "; ".join(f"{n} ({p:g})" for n, p in L["cms"]),
                L["mel"], round(L["rrf"], 1) if L["rrf"] else "-", L["sil"], L["srs"], "OK" if L["match"] else "REVIEW"]
        for j, v in enumerate(vals, 1):
            c = ws3.cell(i, j, v)
            c.border = bd
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate([8, 34, 40, 26, 7, 18, 8, 32, 30, 10, 9, 14, 9, 8], 1):
        ws3.column_dimensions[get_column_letter(j)].width = w
    wb.save(path)


def draw_grid(rows, stem):
    sh = Sheet("A1", "CAUSE & EFFECT MATRIX", "SIS / BMS / ESD / F&G - CDU / VDU", "CFU-000-IC-CE-001",
               discipline="INSTRUMENTATION")
    p = Pen(sh)
    HH = 78
    x0, y0 = 16, 17 + HH
    cw = [10, 19, 26, 36, 92, 24, 22, 24, 13, 20]
    heads = ["ITEM", "SIF", "INITIATOR", "SENSOR(S)", "DESCRIPTION", "VOTING", "NORMAL", "TRIP SP", "SIL", "RESP."]
    ac = 15.0
    rh = 9.6
    xa = x0 + sum(cw)
    W = sum(cw) + ac * len(ACTIONS)
    H = rh * len(rows)
    p.box(x0, y0 - HH, W, HH, None, fill="#1F3864", rx=0)
    xx = x0
    for c, h in zip(cw, heads):
        p.text(h, xx + 1.5, y0 - 3, 3.0, bold=True, color="white")
        xx += c
    p.text("CAUSES (INITIATORS)", x0 + 3, y0 - HH + 10, 4.5, bold=True, color="white")
    p.text("EFFECTS (FINAL ELEMENTS / ACTIONS) -->", x0 + 3, y0 - HH + 18, 3.2, color="#DDE6F5")
    # groups
    groups = []
    for k, (key, tag, desc, grp) in enumerate(ACTIONS):
        if groups and groups[-1][0] == grp:
            groups[-1][2] = k
        else:
            groups.append([grp, k, k])
    for grp, k0, k1 in groups:
        xg0, xg1 = xa + ac * k0, xa + ac * (k1 + 1)
        p.line([(xg0, y0 - HH), (xg0, y0 + H)], w=0.5, color="black")
        p.text(grp, (xg0 + xg1) / 2, y0 - HH + 5.5, 2.6 if (k1 - k0) else 2.0, "middle", bold=True,
               color="#FFD966")
    p.line([(xa, y0 - HH + 8), (xa + ac * len(ACTIONS), y0 - HH + 8)], w=0.3, color="white")
    for k, (key, tag, desc, grp) in enumerate(ACTIONS):
        x = xa + ac * k
        if k:
            p.line([(x, y0 - HH + 8), (x, y0)], w=0.25, color="white")
        p.text(tag, x + 4.6, y0 - 2, 3.0, bold=True, color="white", rotate=-90)
        lines = p.wrap(desc, HH - 12, 2.3)[:2]
        for i, ln in enumerate(lines):
            p.text(ln, x + 8.4 + 3.0 * i, y0 - 2, 2.3, color="#DDE6F5", rotate=-90)
    for i, r in enumerate(rows):
        y = y0 + rh * i
        if r["sif"] == "F&G":
            p.gs.add(p.d.rect((x0, y), (W, rh), fill="#FFF7E6", stroke="none"))
        elif i % 2:
            p.gs.add(p.d.rect((x0, y), (W, rh), fill="#F2F5FA", stroke="none"))
        vals = [str(i + 1), r["sif"], r["tag"], r["init"], r["desc"], r["vote"], r["normal"], r["sp"], r["sil"],
                r["rt"]]
        xx = x0
        for c, v in zip(cw, vals):
            sz = 2.9
            while len(v) and sz > 2.0 and len(v) * sz * 0.56 > c - 2:
                sz -= 0.1
            p.text(v, xx + 1.5, y + rh * 0.66, sz)
            xx += c
        for k, key in enumerate(AKEYS):
            if key in r["act"]:
                v = "A" if key in ("A22", "A23") else "O" if key in ("A2", "A7", "A21", "A8") else \
                    "P" if key in ("A4", "A5") else "X"
                cx = xa + ac * k + ac / 2
                col = SIS_RED if v == "X" else "#2E7D32"
                p.gs.add(p.d.circle((cx, y + rh / 2), 3.3, fill="white" if v != "X" else "#F4CCCC", stroke=col,
                                    stroke_width=0.4))
                p.text(v, cx, y + rh / 2 + 1.3, 3.4, "middle", bold=True, color=col)
        p.line([(x0, y + rh), (x0 + W, y + rh)], w=0.2, color=GREY)
    p.box(x0, y0, W, H, None, fill="none", sw=0.5, rx=0)
    xx = x0
    for c in cw:
        xx += c
        p.line([(xx, y0), (xx, y0 + H)], w=0.25)
    for k in range(len(ACTIONS)):
        p.line([(xa + ac * k, y0), (xa + ac * k, y0 + H)], w=0.2, color=GREY)
    yN = y0 + H + 6
    p.text("LEGEND:  X = trip / close (de-energise)   O = open   P = purge / safe position, controller to manual-"
           "track   A = alarm.   Shaded rows = F&G confirmed events.", x0, yN, 2.8, bold=True)
    notes = ["SIS / BMS hosted in the SIL 3-capable logic solver (separate from DCS); BMS per NFPA 85/86 & API 556. "
             "All trips latch; reset from CCR after initiator healthy and purge / permissives complete.",
             "Set points derived from process_results / psv data (set-point basis in RPT-002); vendor confirmation "
             "required. F&G: 2ooN confirmed gas (20 % LEL alarm, 50 % LEL action) or fire; layout per mapping study.",
             "SIF-204 ROSOVs fire-safe, close <= 30 s. SIF-109 = D-105 LL -> XV-1093 (SIF list revision)."]
    for i, n in enumerate(notes):
        p.text(f"{i + 1}. {n}", x0, yN + 4.5 + 3.8 * i, 2.5)
    sh.save(stem)


def build():
    rows = initiators()
    OUT.mkdir(parents=True, exist_ok=True)
    write_xlsx(rows, OUT / "CFU-000-IC-CE-001_Cause-Effect-Matrix.xlsx")
    draw_grid(rows, OUT / "CFU-000-IC-CE-001_Cause-Effect-Matrix")
    return rows
