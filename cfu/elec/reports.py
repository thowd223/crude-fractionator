"""Electrical deliverable writers: load list (xlsx + pdf), sizing calculation (md + pdf), cable schedule (xlsx),
data/electrical.json."""
from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import basis
from . import common as C
from .loads import B4, B13, BLV, BUPS, demand

P = basis.PROJECT
LDL = "CFU-000-EL-LDL-001"
CAL = "CFU-000-EL-CAL-001"
CBL = "CFU-000-EL-CBL-001"
HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
TOT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
DUTY = {"C": "Continuous", "I": "Intermittent", "S": "Standby"}


def _r(x, n=1):
    return None if x is None else round(float(x), n)


def _title(ws, docno, title, ncol):
    ws.cell(1, 1, f"{P['name']} - {title}").font = Font(bold=True, size=13)
    ws.cell(2, 1, f"{docno}   Rev {P['rev']} ({P['rev_desc']}), 2026-10-02   |   {P['client']}").font = \
        Font(italic=True, size=9)
    return 4


def _table(ws, r0, header, rows, widths, fmt=None, wrap_cols=()):
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF", size=9), HDR, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.border = BOX
            c.font = Font(size=9)
            if isinstance(v, float):
                c.number_format = (fmt or {}).get(j, "#,##0.0" if abs(v) < 100 else "#,##0")
            if j in wrap_cols:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 3)
    return r0 + len(rows) + 1


# ============================================================================================ load list
LDL_HDR = ["Item", "Tag", "Description", "Equip.", "Category", "Volt (V)", "Starter / feeder", "Bus",
           "Rated kW", "Rated kVA", "Absorbed kW", "Load factor", "Eff.", "PF", "Input kW", "kvar", "kVA",
           "Duty", "C kW", "I kW", "S kW", "C kvar", "I kvar", "S kvar", "Notes"]


def ldl_rows(rows):
    out = []
    for r in rows:
        st = r["starter"] if r["kind"] == "motor" else ("Feeder (via LTR)" if r["volt"] == 208 else "Feeder")
        out.append([r["item"], r["tag"], r["desc"], r["eq"], r["group"], int(r["volt"]), st, r["bus"],
                    _r(r["rated_kw"], 2), _r(r.get("kva_rated")), _r(r["absorbed_kw"], 2), _r(r["lf"], 2),
                    _r(r["eff"], 3), _r(r["pf"], 2), _r(r["kw"], 2), _r(r["kvar"], 2), _r(r["kva"], 2),
                    DUTY[r["duty"]], _r(r["c_kw"], 2), _r(r["i_kw"], 2), _r(r["s_kw"], 2), _r(r["c_kvar"], 2),
                    _r(r["i_kvar"], 2), _r(r["s_kvar"], 2), r["note"]])
    return out


def bus_summary(ctx):
    rows, R = ctx["rows"], ctx["R"]
    out = []
    for b in BLV + B4:
        d = R["md"][b]
        out.append([b, "0.48" if b in BLV else "4.16", d["n"], d["connected_kW"], d["C_kW"], d["I_kW"], d["S_kW"],
                    d["kW"], d["kvar"], d["kVA"], d["pf"]])
    for b in B4:
        bb = R["bus4"][b]
        out.append([b + " incl. LV through-load", "4.16", None, None, None, None, None, bb["kW"], bb["kvar"],
                    bb["kVA"], bb["kW"] / bb["kVA"]])
    for b in B13:
        bb = R["bus13"][b]
        out.append([b, "13.8", None, None, None, None, None, bb["kW"], bb["kvar"], bb["kVA"], bb["kW"] / bb["kVA"]])
    t = R["total"]
    lt, mt = R["md_lv_total"], R["md_mv_total"]
    out.append(["LV total (MCC-101A+B)", "0.48", lt["n"], lt["connected_kW"], lt["C_kW"], lt["I_kW"], lt["S_kW"],
                lt["kW"], lt["kvar"], lt["kVA"], lt["pf"]])
    out.append(["MV motors total (SWG-102A+B)", "4.16", mt["n"], mt["connected_kW"], mt["C_kW"], mt["I_kW"],
                mt["S_kW"], mt["kW"], mt["kvar"], mt["kVA"], mt["pf"]])
    out.append(["UNIT TOTAL at 13.8 kV (incl. TR losses)", "13.8", len(rows), None, None, None, None, t["kW"],
                t["kvar"], t["kVA"], t["pf"]])
    return out


def write_ldl_xlsx(ctx, path: Path):
    rows, R = ctx["rows"], ctx["R"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Load List"
    r0 = _title(ws, LDL, "Electrical Load List", len(LDL_HDR))
    end = _table(ws, r0, LDL_HDR, ldl_rows(rows),
                 [5, 11, 46, 9, 18, 7, 12, 10, 8, 8, 9, 7, 6, 6, 8, 8, 8, 11, 8, 8, 8, 8, 8, 8, 40],
                 fmt={12: "0.00", 13: "0.000", 14: "0.00"}, wrap_cols=(3,))
    # totals with formulas
    n0, n1 = r0 + 1, end - 1
    ws.cell(end + 1, 3, "TOTAL (all buses)").font = Font(bold=True)
    for j in (15, 16, 17, 19, 20, 21, 22, 23, 24):
        col = get_column_letter(j)
        c = ws.cell(end + 1, j, f"=SUM({col}{n0}:{col}{n1})")
        c.font, c.fill, c.number_format = Font(bold=True), TOT, "#,##0"
    rr = end + 3
    ws.cell(rr, 3, "Maximum demand per bus = C + 0.3 x I + 0.1 x S (kW and kvar separately), kVA = vector sum").font = \
        Font(italic=True)
    rr += 1
    for j, h in enumerate(["Bus", "", "MD kW (formula)", "MD kvar (formula)", "MD kVA (formula)"], 1):
        if h:
            c = ws.cell(rr, j + 1 if j > 1 else 2, h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HDR
    for b in BLV + B4:
        rr += 1
        ws.cell(rr, 2, b)
        rng = f"$H${n0}:$H${n1}"
        kw = "+".join(f"{k}*SUMIF({rng},B{rr},{get_column_letter(j)}${n0}:{get_column_letter(j)}${n1})"
                      for k, j in ((1, 19), (0.3, 20), (0.1, 21)))
        kv = "+".join(f"{k}*SUMIF({rng},B{rr},{get_column_letter(j)}${n0}:{get_column_letter(j)}${n1})"
                      for k, j in ((1, 22), (0.3, 23), (0.1, 24)))
        ws.cell(rr, 4, "=" + kw).number_format = "#,##0"
        ws.cell(rr, 5, "=" + kv).number_format = "#,##0"
        ws.cell(rr, 6, f"=SQRT(D{rr}^2+E{rr}^2)").number_format = "#,##0"

    ws2 = wb.create_sheet("Bus Summary")
    r0 = _title(ws2, LDL, "Maximum Demand per Bus", 11)
    end = _table(ws2, r0, ["Bus", "kV", "No. loads", "Connected kW", "C kW", "I kW", "S kW", "MD kW", "MD kvar",
                           "MD kVA", "PF"], bus_summary(ctx), [40, 6, 9, 12, 10, 10, 10, 10, 10, 10, 7],
                 fmt={11: "0.00"})
    trm, trl = R["tr_mv"], R["tr_lv"]
    rows2 = [[t["tag"], f"{t['pri_kV']}/{t['sec_kV']} kV", t["onan_kVA"], t["onaf_kVA"], t["z_pct"], t["md_kVA"],
              t["req_onaf_kVA"], t["load_onan_pct"], t["cooling"], t["vector"], t["grounding"]] for t in (trm, trl)]
    rows2 += [[f"LTR-101{ab}", "0.48/0.208 kV", R["ltr"][ab]["kva"], None, 4.0, R["ltr"][ab]["kva_load"],
               R["ltr"][ab]["kva_load"] * 1.25, 100 * R["ltr"][ab]["kva_load"] / R["ltr"][ab]["kva"], "Dry type (AN)",
               "Dyn1", "Solidly grounded"] for ab in "AB"]
    ws2.cell(end + 2, 1, "Transformers").font = Font(bold=True, size=11)
    _table(ws2, end + 3, ["Tag", "Ratio", "ONAN kVA", "ONAF kVA", "Z %", "MD kVA (one TR, tie closed)",
                          "MD x 1.25 kVA", "Loading at MD % ONAN", "Cooling", "Vector", "Neutral"], rows2,
           [40, 14, 10, 10, 6, 14, 12, 12, 14, 8, 40])

    ws3 = wb.create_sheet("UPS Loads")
    r0 = _title(ws3, LDL, "UPS (120 V) Consumers", 8)
    ur = [[r["tag"], r["desc"], r["bus"], r["kva_rated"], _r(r["kva"], 2), _r(r["kw"], 2), r["pf"]]
          for r in ctx["ups_rows"]]
    _table(ws3, r0, ["Tag", "Description", "UPS bus", "Consumer total kVA", "This feed kVA (50 %)", "This feed kW",
                     "PF"], ur, [14, 70, 11, 12, 12, 10, 6], wrap_cols=(2,))

    ws4 = wb.create_sheet("Basis")
    r0 = _title(ws4, LDL, "Load List Basis", 2)
    for i, t in enumerate(basis_lines(ctx), r0):
        ws4.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.column_dimensions["A"].width = 150
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def basis_lines(ctx):
    return [
        "1. Sources: motor ratings, absorbed power, fan counts and package data from data/equipment.json "
        "(process/mechanical sizing). Absorbed power used where available; 'Input kW' = absorbed / efficiency "
        "(/ VFD efficiency 0.97 where applicable).",
        "2. Voltage levels: motors >= 200 kW at 4.16 kV (VCB, DOL unless VFD); motors < 200 kW at 480 V (MCC, DOL). "
        "VFDs: A-101 overhead condenser fans (4 x 90 kW, process temperature control and energy) and H-101 "
        "ID fans K-102A/B (draft control, 4.16 kV MV drive).",
        "3. Efficiency / PF by motor size (NEMA Premium / IEEE 841 typical): <=1.5 kW 0.80/0.76; <=7.5 kW 0.875/0.82; "
        "<=30 kW 0.91/0.85; <=90 kW 0.935/0.86; <=160 kW 0.95/0.87; LV >160 kW 0.955/0.87; MV <500 kW 0.955/0.87; "
        "MV >=500 kW 0.965/0.89. VFD supply PF 0.95.",
        "4. Duty: C = continuous, I = intermittent, S = standby. Pump pairs A/B (2 x 100 %): A always on bus A, "
        "B on bus B; the normally running unit of each pair (C) is chosen to balance the two bus sections, the "
        "other is standby (S). Air-cooler fans all continuous, split alternately across the buses.",
        "5. Maximum demand (IEEE 141 / IEC style diversity): MD = 1.0 C + 0.3 I + 0.1 S, applied to kW and kvar "
        "separately; kVA by vector sum.",
        "6. Allowances (no process data available - FEED assumptions, to be replaced by vendor data): MOV/actuators "
        "36 x 1.5 kW; heat tracing 70 kW (slop/HN/slop-wax lines); lighting 2 x 29 kW LED + building lighting; small "
        "power/receptacles; welding receptacles 6 x 63 A; SS-100 HVAC 2 x 60 kW, FAR-100 HVAC 2 x 25 kW; cathodic "
        "protection 8 kW; analyser house 15 kW; UPS loads 60 kVA; 125 V DC chargers 2 x 6 kW.",
        "7. Desalter transformers: 2 x 150 kVA per stage (D-101A, D-101B), 480 V primary, load factor 0.35 (vendor "
        "to confirm), PF 0.85; one transformer of each stage on each bus.",
        "8. Chemical injection packages X-101..X-104: 2 x 100 % metering pumps 0.75 kW each, absorbed 60 %.",
        "9. 208 V loads are supplied via lighting transformers LTR-101A/B (480-208Y/120 V) from MCC-101A/B.",
        "10. UPS consumers (120 V) are listed separately; they are supplied through UPS-101A/B whose rectifier input "
        "(incl. battery recharge) appears on MCC-101A/B. Not double counted.",
    ]


def write_ldl_pdf(ctx, path: Path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    rows, R = ctx["rows"], ctx["R"]
    W, H = landscape(A3)
    st = ParagraphStyle("c", fontSize=6.4, leading=7.4)
    sth = ParagraphStyle("h", fontSize=6.4, leading=7.4, textColor=colors.white, fontName="Helvetica-Bold")
    sp = ParagraphStyle("p", fontSize=8.5, leading=11)
    h1 = ParagraphStyle("h1", fontSize=13, leading=16, fontName="Helvetica-Bold", textColor=colors.HexColor("#1F3864"))

    def deco(c, doc):
        c.saveState()
        c.setFont("Helvetica-Bold", 10)
        c.drawString(12 * mm, H - 10 * mm, f"{P['name']}  -  ELECTRICAL LOAD LIST")
        c.setFont("Helvetica", 8)
        c.drawRightString(W - 12 * mm, H - 10 * mm, f"{LDL}   Rev {P['rev']}   {P['rev_desc']}   2026-10-02")
        c.line(12 * mm, H - 12 * mm, W - 12 * mm, H - 12 * mm)
        c.drawString(12 * mm, 7 * mm, f"{P['client']}  |  {P['stage']}  |  Prepared: Claude Code  |  Checked: -  |  "
                                      "Approved: -")
        c.drawRightString(W - 12 * mm, 7 * mm, f"Page {doc.page}")
        c.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=(W, H), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=16 * mm,
                            bottomMargin=13 * mm, title=f"{LDL} Electrical Load List")
    hdr = ["#", "Tag", "Description", "Volt", "Starter", "Bus", "Rated kW", "Abs. kW", "LF", "Eff", "PF", "Input kW",
           "kvar", "kVA", "Duty", "C kW", "I kW", "S kW"]
    cw = [8, 20, 100, 11, 17, 18, 13, 13, 9, 10, 9, 14, 13, 13, 17, 13, 13, 13]
    flow = [Paragraph("Electrical Load List - all consumers by voltage level and bus", h1), Spacer(1, 3 * mm)]
    for lvl, buses in (("4.16 kV MOTORS", B4), ("480 V / 208 V CONSUMERS", BLV)):
        for b in buses:
            data = [[Paragraph(h, sth) for h in hdr]]
            sel = [r for r in rows if r["bus"] == b]
            for r in sel:
                st_ = r["starter"] if r["kind"] == "motor" else ("Fdr/LTR" if r["volt"] == 208 else "Feeder")
                data.append([str(r["item"]), r["tag"], Paragraph(r["desc"], st), f"{r['volt']:.0f}", st_, r["bus"],
                             f"{r['rated_kw']:.1f}", f"{r['absorbed_kw']:.1f}", f"{r['lf']:.2f}", f"{r['eff']:.3f}",
                             f"{r['pf']:.2f}", f"{r['kw']:.1f}", f"{r['kvar']:.1f}", f"{r['kva']:.1f}",
                             DUTY[r["duty"]], f"{r['c_kw']:.1f}", f"{r['i_kw']:.1f}", f"{r['s_kw']:.1f}"])
            d = R["md"][b]
            data.append(["", "", Paragraph(f"<b>{b} totals; MD = C + 0.3 I + 0.1 S = {d['kW']:,.0f} kW / "
                                           f"{d['kvar']:,.0f} kvar / {d['kVA']:,.0f} kVA (PF {d['pf']:.2f})</b>", st),
                         "", "", "", "", "", "", "", "", f"{sum(r['kw'] for r in sel):,.0f}",
                         f"{sum(r['kvar'] for r in sel):,.0f}", "", "", f"{d['C_kW']:,.0f}", f"{d['I_kW']:,.0f}",
                         f"{d['S_kW']:,.0f}"])
            t = Table(data, colWidths=[c * mm for c in cw], repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                ("FONTSIZE", (0, 0), (-1, -1), 6.4), ("LEADING", (0, 0), (-1, -1), 7.4),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#999999")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#EEF3FA")]),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FFF2CC")),
                ("ALIGN", (6, 1), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 1.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2)]))
            flow += [Paragraph(f"<b>{lvl} - {b}</b>", sp), Spacer(1, 1.5 * mm), t, Spacer(1, 5 * mm)]
    flow.append(PageBreak())
    flow.append(Paragraph("Maximum demand summary and transformer loading", h1))
    flow.append(Spacer(1, 3 * mm))
    sh = ["Bus", "kV", "No.", "Conn. kW", "C kW", "I kW", "S kW", "MD kW", "MD kvar", "MD kVA", "PF"]
    data = [[Paragraph(h, sth) for h in sh]]
    for r in bus_summary(ctx):
        data.append([r[0], r[1]] + ["" if v is None else (f"{v:,.0f}" if isinstance(v, float) and j < 8 else
                                                           f"{v:.2f}" if j == 8 else str(v))
                                    for j, v in enumerate(r[2:])])
    t = Table(data, colWidths=[x * mm for x in (80, 14, 14, 22, 22, 22, 22, 22, 22, 22, 14)])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                           ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                           ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                           ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FFF2CC"))]))
    flow += [t, Spacer(1, 5 * mm)]
    trm, trl = R["tr_mv"], R["tr_lv"]
    for tx in (trm, trl):
        flow.append(Paragraph(
            f"<b>{tx['tag']}</b>: {tx['pri_kV']}/{tx['sec_kV']} kV, {tx['onan_kVA']:,} / {tx['onaf_kVA']:,} kVA "
            f"ONAN/ONAF, Z {tx['z_pct']} %. Contingency MD (both sections on one transformer) {tx['md_kVA']:,.0f} kVA "
            f"= {tx['load_onan_pct']:.0f} % of ONAN; MD x 1.25 = {tx['req_onaf_kVA']:,.0f} kVA = "
            f"{tx['load_onaf_future_pct']:.0f} % of ONAF.", sp))
    flow.append(Spacer(1, 4 * mm))
    flow.append(Paragraph("<b>UPS consumers (120 V, via UPS-101A/B; not included in bus totals above except as "
                          "UPS rectifier input)</b>", sp))
    data = [[Paragraph(h, sth) for h in ["Code", "Description", "kVA total", "Feed A kVA", "Feed B kVA"]]]
    for r in ctx["ups_rows"]:
        if r["tag"].endswith("-A"):
            data.append([r["tag"][:-2], Paragraph(r["desc"].split(" - feed")[0], st), f"{r['kva_rated']:g}",
                         f"{r['kva_rated'] / 2:g}", f"{r['kva_rated'] / 2:g}"])
    t = Table(data, colWidths=[x * mm for x in (25, 120, 22, 22, 22)])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                           ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]))
    flow += [t, Spacer(1, 5 * mm), Paragraph("<b>Basis</b>", sp)]
    for ln in basis_lines(ctx):
        flow.append(Paragraph(ln, sp))
    doc.build(flow, onFirstPage=deco, onLaterPages=deco)


# ============================================================================================ cable schedule
CBL_HDR = ["No.", "Cable tag", "From", "To", "Service", "kV", "Design I (A)", "Length (m)", "Length basis",
           "Cable (cores x size, type)", "Size mm2", "Runs", "Ampacity derated (A)", "VD running %",
           "VD start (cable) %", "Bus dip at start %", "Terminal dip %", "SC min mm2", "Protective device",
           "CT", "Relays (ANSI)", "Notes"]


def write_cbl_xlsx(ctx, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cable Schedule"
    r0 = _title(ws, CBL, "Power Cable Schedule", len(CBL_HDR))
    rows = []
    for c in ctx["cbl"]:
        rows.append([c["no"], c["tag"], c["frm"], c["to"], c["service"], c["kV"], _r(c["I_A"]), c["L_m"],
                     c["length_src"], c["cable"], c["size"], c["runs"], _r(c["ampacity"]), _r(c["vd_run"], 2),
                     _r(c["vd_start"], 2), _r(c.get("bus_dip"), 2), _r(c.get("term_dip"), 2), _r(c["s_sc_min"]),
                     c["device"], c.get("ct", ""), c.get("relays", ""), c.get("notes", "")])
    end = _table(ws, r0, CBL_HDR, rows, [5, 16, 22, 26, 40, 5, 9, 8, 14, 52, 7, 5, 10, 8, 8, 8, 8, 8, 44, 8, 34, 40],
                 fmt={14: "0.00", 15: "0.00", 16: "0.00", 17: "0.00", 6: "0.00"})
    ws.cell(end + 1, 2, "TOTAL LENGTH (m)").font = Font(bold=True)
    ws.cell(end + 1, 8, f"=SUM(H{r0 + 1}:H{end - 1})").font = Font(bold=True)
    # quantities by cable type for estimate
    q = {}
    for c in ctx["cbl"]:
        key = c["cable"]
        q[key] = q.get(key, 0.0) + c["L_m"] * (c["runs"] or 1)
    ws2 = wb.create_sheet("Quantities")
    r0 = _title(ws2, CBL, "Cable Quantities by Type (length x parallel runs)", 3)
    _table(ws2, r0, ["Cable type", "Total length (m)", "With 10 % wastage (m)"],
           [[k, v, v * 1.1] for k, v in sorted(q.items())], [80, 16, 20])
    ws3 = wb.create_sheet("Basis")
    r0 = _title(ws3, CBL, "Cable Sizing Basis", 1)
    for i, t in enumerate(cable_basis(ctx), r0):
        ws3.cell(i, 1, t).alignment = Alignment(wrap_text=True)
    ws3.column_dimensions["A"].width = 150
    wb.save(path)
    return q


def cable_basis(ctx):
    loc = ctx["loc"]
    der = C.DERATE["ambient"] * C.DERATE["group"]
    return [
        "1. Conductors: copper, XLPE 90 C, IEC 60228 metric sizes (mm2). LV 0.6/1 kV Cu/XLPE/SWA/PVC multicore with "
        "separate earth core (IEC 60364-5-54 sizing); 4.16 kV: 3.6/6 kV screened 3-core Cu/XLPE/CWS/SWA/PVC; "
        "13.8 kV: 8.7/15 kV single-core in trefoil. Equivalent UL/NEC types (MC-HL / TC-ER) acceptable; "
        "NEC 310 ampacity check at detailed design.",
        f"2. Ampacity: base values IEC 60364-5-52 Table B.52.12 method E (LV) / IEC 60502-2 Annex B (MV) at 30 C in "
        f"air on ladder tray; derating {C.DERATE['ambient']} (40 C ambient incl. solar) x {C.DERATE['group']} "
        f"(grouping) = {der:.2f}. Required ampacity >= 1.25 x FLC for motors (NEC 430.22) and >= 1.25 x rated "
        "current for feeders (continuous load).",
        "3. Voltage drop: dV% = sqrt(3) I L (R cos(phi) + X sin(phi)) / V x 100 with R at 90 C. Running <= 5 %. "
        f"Starting (DOL, {C.LRC} x FLC at PF {C.LR_PF_LV} LV / {C.LR_PF} MV): bus dip + cable drop <= 15 % at motor "
        "terminals; bus dip <= 10 %.",
        f"4. Short-circuit withstand: S >= sqrt(I^2 t) / k, k = {C.K_XLPE:g} (Cu/XLPE 90->250 C). MV: prospective "
        "bus fault current with t = 0.25 s (motor feeders, instantaneous 50 element) / 0.5 s (incomers, transformer "
        "feeders). LV: let-through I^2t of current-limiting MCCB by rating (manufacturer typical, 65 kA class).",
        f"5. Minimum sizes: LV power {C.LV_MIN_MM2:g} mm2; 4.16 kV {C.MV_MIN_MM2:g} mm2. Parallel LV runs >= 95 mm2.",
        f"6. Route length = Manhattan distance from SS-100 reference point ({loc.ss[0]:g}, {loc.ss[1]:g}) m to the "
        f"consumer + {C.RISER_M:g} m riser/termination allowance (+{C.RACK_TOP_EXTRA_M:g} m for air-cooler fan "
        "motors on top of the pipe rack), rounded up to 5 m. Coordinates source: " + loc.source + ".",
        "7. 13.8 kV incomer route from the refinery main substation assumed 600 m (OSBL, to be confirmed by "
        "refinery electrical master plan).",
    ]


# ============================================================================================ calculation
def _t(header, rows):
    s = "| " + " | ".join(header) + " |\n|" + "|".join("---" for _ in header) + "|\n"
    for r in rows:
        s += "| " + " | ".join(str(x) for x in r) + " |\n"
    return s


def calc_md(ctx):
    rows, R, ud, cab = ctx["rows"], ctx["R"], ctx["ud"], ctx["cab"]
    trm, trl, sw, sc = R["tr_mv"], R["tr_lv"], R["swgr"], R["sc"]
    U, D = ud["ups"], ud["dc"]
    net = R["net"]
    md = R["md"]
    L = []
    a = L.append
    a(f"# {CAL} - Electrical System Sizing Calculations\n")
    a(f"Rev {P['rev']} - {P['rev_desc']}. Unit: {P['name']} (100,000 BPSD). Generated from `cfu/elec` "
      "(load list, studies and cable schedule are computed from data/equipment.json; no hand-typed process data).\n")
    a("## 1. Purpose and scope\n")
    a("FEED-level sizing of the unit electrical distribution in substation SS-100: maximum demand, transformer "
      "ratings, short-circuit levels and switchgear ratings, largest-motor starting voltage dip, cable sizing method "
      "and results, UPS and DC battery sizing, and the basis for not providing an emergency diesel generator. "
      "Hazardous-area classification is deferred until the plot plan is final.\n")
    a("## 2. Basis and references\n")
    a(_t(["Item", "Value"], [
        ["Utility supply", f"2 x {C.KV_UT} kV feeders from refinery main substation, 3-ph, {C.HZ} Hz"],
        ["Utility fault level", f"{C.UT_FAULT_KA} kA at 13.8 kV ({net['S_utility_MVA']:.0f} MVA), X/R {C.UT_XR:g} (assumed)"],
        ["Distribution voltages", f"{C.KV_MV} kV MV motors >= {C.MV_MOTOR_KW:g} kW; {C.V_LV} V LV; 208Y/120 V lighting; "
                                  f"{C.V_UPS} V UPS"],
        ["System grounding", "13.8 kV per utility; 4.16 kV low-resistance 400 A / 10 s; 480 V high-resistance 5 A "
                             "(alarm, continued operation); 208Y/120 V solidly grounded"],
        ["Codes", "NFPA 70 (NEC), IEEE 141, 242, 399, 485, 1100, IEEE C57.12.00/.10, C37.20.1/.2/.7, C37.2, "
                  "IEC 60364-5-52, IEC 60502-2, IEC 60909 (peak factor), API RP 540"],
        ["Ambient", f"{basis.SITE['amb_design_C']} C design, {basis.SITE['amb_min_C']} C min, elevation "
                    f"{basis.SITE['elevation_m']} m"],
    ]))
    a("## 3. System configuration\n")
    a("Secondary-selective (double-ended) arrangement at every level: two 13.8 kV incomers to SWG-101A/B with a "
      "normally-open tie; two 13.8/4.16 kV transformers TR-101/102 feeding SWG-102A/B (N.O. tie); two "
      "4.16/0.48 kV transformers TR-103/104 feeding the 480 V switchgear/MCC-101A/B (N.O. tie). Ties are "
      "interlocked 2-out-of-3 so transformers are never paralleled; on loss of one source the incomer opens and the "
      "tie closes automatically (open transition). Every transformer therefore carries the full demand of both bus "
      "sections in the contingency case. Pump pairs (2 x 100 %) are split across the two bus sections. See "
      "CFU-000-EL-SLD-001..004.\n")
    a("## 4. Electrical load and maximum demand\n")
    a("Maximum demand MD = 1.0 x C + 0.3 x I + 0.1 x S (continuous / intermittent / standby), applied separately to "
      "kW and kvar (IEEE 141 practice). Motor input = absorbed kW / efficiency (/ 0.97 for VFDs). Detailed list: "
      f"{LDL}.\n")
    hr = ["Bus", "Connected kW", "C kW", "I kW", "S kW", "MD kW", "MD kvar", "MD kVA", "PF"]
    rr = []
    for b in BLV + B4:
        d = md[b]
        rr.append([b, f"{d['connected_kW']:,.0f}", f"{d['C_kW']:,.0f}", f"{d['I_kW']:,.0f}", f"{d['S_kW']:,.0f}",
                   f"{d['kW']:,.0f}", f"{d['kvar']:,.0f}", f"{d['kVA']:,.0f}", f"{d['pf']:.2f}"])
    for b in B4:
        bb = R["bus4"][b]
        rr.append([f"{b} incl. LV", "", "", "", "", f"{bb['kW']:,.0f}", f"{bb['kvar']:,.0f}", f"{bb['kVA']:,.0f}",
                   f"{bb['kW'] / bb['kVA']:.2f}"])
    t = R["total"]
    rr.append(["**Unit total at 13.8 kV**", "", "", "", "", f"**{t['kW']:,.0f}**", f"**{t['kvar']:,.0f}**",
               f"**{t['kVA']:,.0f}**", f"{t['pf']:.2f}"])
    a(_t(hr, rr))
    a("LV transformer through-load includes 1 % active / 4 % reactive transformer losses; MV transformer 0.8 % / 6 %.\n")
    a("## 5. Transformer sizing\n")
    a("Criteria: (a) ONAN rating >= contingency maximum demand (both bus sections on one transformer, tie closed), "
      "with ONAN loading <= 95 %; (b) ONAF (fan-cooled, +25 % for <= 10 MVA per IEEE C57.12.10) >= 1.25 x MD, "
      "i.e. the 25 % future margin is available with forced cooling. Off-circuit taps +/-2 x 2.5 %.\n")
    a(_t(["Transformer", "Ratio", "MD kVA", "1.25 x MD", "Selected ONAN/ONAF kVA", "ONAN load %", "Z %",
          "Secondary FLC (ONAF) A"],
         [[trm["tag"], "13.8/4.16 kV", f"{trm['md_kVA']:,.0f}", f"{trm['req_onaf_kVA']:,.0f}",
           f"{trm['onan_kVA']:,} / {trm['onaf_kVA']:,}", f"{trm['load_onan_pct']:.0f}", trm["z_pct"],
           f"{trm['I_sec_onaf_A']:,.0f}"],
          [trl["tag"], "4.16/0.48 kV", f"{trl['md_kVA']:,.0f}", f"{trl['req_onaf_kVA']:,.0f}",
           f"{trl['onan_kVA']:,} / {trl['onaf_kVA']:,}", f"{trl['load_onan_pct']:.0f}", trl["z_pct"],
           f"{trl['I_sec_onaf_A']:,.0f}"]] +
         [[f"LTR-101{ab}", "480-208Y/120 V", f"{R['ltr'][ab]['kva_load']:.0f}", f"{R['ltr'][ab]['kva_load'] * 1.25:.0f}",
           f"{R['ltr'][ab]['kva']:g} (dry, AN)", f"{100 * R['ltr'][ab]['kva_load'] / R['ltr'][ab]['kva']:.0f}", 4.0, ""]
          for ab in "AB"]))
    a(f"The 13.8/4.16 kV unit rating ({trm['onan_kVA'] / 1000:g} MVA) is governed by the 95 % ONAN loading limit: the "
      f"next smaller standard size (5 MVA) would be loaded to {100 * trm['md_kVA'] / 5000:.0f} % with no ONAN "
      "margin. The 480 V transformers are at the practical upper limit for 480 V unit substations (5000 A bus); "
      "see Section 11.\n")
    a("## 6. Short-circuit levels\n")
    a("Method: IEEE 141 / ANSI E/X hand calculation on a 100 MVA base, prefault voltage 1.0 pu. Sources: utility "
      f"({net['S_utility_MVA']:.0f} MVA, X/R {C.UT_XR:g}); running motors as subtransient sources - MV motors "
      f"X\" = {C.XD2} pu on motor kVA, LV motors grouped at 4 x FLC (X\" = 0.25 pu); VFD-fed motors excluded. "
      "Peak ip = kappa x sqrt(2) x Ik\", kappa = 1.02 + 0.98 e^(-3R/X) (IEC 60909).\n")
    a(_t(["Element", "Z (pu, 100 MVA)", "X/R"],
         [["Utility 13.8 kV", f"{abs(net['Zut']):.4f}", f"{C.UT_XR:g}"],
          [f"TR-101/102 ({trm['onan_kVA'] / 1000:g} MVA, {trm['z_pct']} %)", f"{abs(net['Ztm']):.4f}", f"{trm['xr']:g}"],
          [f"TR-103/104 ({trl['onan_kVA'] / 1000:g} MVA, {trl['z_pct']} %)", f"{abs(net['Ztl']):.4f}", f"{trl['xr']:g}"]]))
    rr = []
    for case, v in sc["cases"].items():
        for lv in ("13.8 kV", "4.16 kV", "0.48 kV"):
            rr.append([case, lv, f"{v[lv]['Ik_kA']:.1f}", f"{v[lv]['XR']:.1f}", f"{v[lv]['ip_kA']:.1f}"])
    a(_t(["Case", "Bus", "Ik\" kA sym", "X/R", "ip kA peak"], rr))
    a(_t(["Equipment", "Max. calculated Ik\" kA", "Selected rating (kA sym, >= 1.1 x calc.)", "Bus continuous A"],
         [["SWG-101 13.8 kV", f"{sc['worst_kA']['13.8 kV']:.1f}", f"{sc['rating_kA']['13.8 kV']:g}",
           sw["SWG-101"]["bus_A"]],
          ["SWG-102 4.16 kV", f"{sc['worst_kA']['4.16 kV']:.1f}", f"{sc['rating_kA']['4.16 kV']:g}",
           sw["SWG-102"]["bus_A"]],
          ["MCC-101 480 V", f"{sc['worst_kA']['0.48 kV']:.1f}", f"{sc['rating_kA']['0.48 kV']:g}",
           sw["MCC-101"]["bus_A"]]]))
    a("The 13.8 kV rating is set by the utility (31.5 kA) plus motor contribution; 40 kA switchgear is specified. "
      "Ratings assume the tie is never closed with both incomers in service (2-out-of-3 interlock).\n")
    a("## 7. Motor starting voltage dip\n")
    big = R["start"][0]
    a(f"Largest DOL motor: {big['motor']} ({big['kW']:g} kW, 4.16 kV; P-101A/B and P-102A/B are identical 710 kW "
      f"units). Locked-rotor {C.LRC} x FLC at PF {C.LR_PF} -> {big['LR_kVA']:,.0f} kVA. Network: utility + one "
      "13.8/4.16 kV transformer (ONAN impedance); pre-start bus load modelled as constant impedance with pre-start bus "
      "voltage 1.00 pu (tap setting); motor cable impedance included for the terminal voltage. Limits: 10 % at the "
      "4.16 kV bus, 15 % at the motor terminals.\n")
    a(_t(["Case", "Pre-start load kVA", "Bus dip %", "Terminal dip %", "Result"],
         [[s["case"], f"{s['pre_kVA']:,.0f}", f"{s['bus_dip_pct']:.1f}", f"{s['term_dip_pct']:.1f}",
           "OK" if s["ok"] else "EXCEEDS"] for s in R["start"]]))
    a("DOL starting of the 710 kW pumps is therefore acceptable; no soft-starter / autotransformer is required. "
      "Motor-acceleration time and pump torque margin to be confirmed with vendor curves (IEEE 399 dynamic study at "
      "detailed design). LV motors: bus dip on MCC-101A/B for each DOL start is computed in the cable schedule "
      f"(largest {max((c['bus_dip'] or 0) for c in ctx['cbl'] if c.get('bus_dip') is not None and c['kV'] < 1):.1f} %).\n")
    a("## 8. Cable sizing\n")
    for ln in cable_basis(ctx):
        a("- " + ln[3:])
    a("")
    # worked examples
    for tag in ("CBL-P-101A", "CBL-P-104A"):
        c = cab[tag]
        a(f"**Worked example {tag}:** {c['service']}; design FLC {c['I_A']:.0f} A; route {c['L_m']:.0f} m "
          f"({c['length_src']}); required ampacity {1.25 * c['I_A']:.0f} A; selected {c['cable']}; derated ampacity "
          f"{c['ampacity']:.0f} A; VD running {c['vd_run']:.2f} %; VD starting (cable) {c['vd_start']:.1f} %; bus dip "
          f"{c['bus_dip']:.1f} % -> terminal {c['term_dip']:.1f} %; SC minimum {c['s_sc_min']:.0f} mm2.\n")
    a(_t(["Summary", "Value"], [
        ["Number of cables", len(ctx["cbl"])],
        ["Total route length (m)", f"{sum(c['L_m'] * (c['runs'] or 1) for c in ctx['cbl']):,.0f}"],
        ["Max running VD %", f"{max(c['vd_run'] for c in ctx['cbl']):.2f}"],
        ["Max terminal dip at start %",
         f"{max((c.get('term_dip') or 0) for c in ctx['cbl']):.1f}"],
    ]))
    a("## 9. UPS and battery sizing\n")
    a(_t(["Step", "Value"], [
        ["UPS consumers (DCS 22, SIS/BMS 14, F&G 6, telecom 8, analysers 6, misc. 4 kVA)", f"{U['load_kVA']:.0f} kVA"],
        ["Design load incl. 20 % future", f"{U['design_kVA']:.0f} kVA"],
        ["Rating (design load <= 80 % of rating)", f"{U['rating_kVA']:g} kVA each, {U['config']}"],
        ["Battery power = 72 kVA x 0.9 / 0.94", f"{U['battery_kW']:.1f} kW"],
        ["Cells / nominal / float / end voltage", f"{U['cells']} x 2 V / {U['v_nom']:.0f} / {U['v_float']:.0f} / "
                                                   f"{U['v_end']:.0f} V (1.75 V/cell)"],
        ["Max discharge current at end voltage", f"{U['I_max_A']:.0f} A"],
        ["Kt for 30 min to 1.75 V/cell (VRLA, 25 C)", f"{U['Kt']} Ah/A"],
        ["Capacity = I x Kt x aging 1.25 x design margin 1.10 x temperature 1.00", f"{U['ah_req']:.0f} Ah"],
        ["Selected battery (per UPS)", f"{U['ah_sel']} Ah, {U['autonomy_min']} min"],
    ]))
    a("UPS input (rectifier) load on each MCC section includes 50 % share of the consumers plus 5 kW battery "
      "recharge. 125 V DC switchgear control supply (IEEE 485 duty cycle):\n")
    a(_t(["Section", "Value"], [
        ["Duty cycle", D["duty"]],
        ["Section 1 / 2 / 3 capacity (Ah)", " / ".join(f"{x:.1f}" for x in D["sections"])],
        ["Required = max x aging 1.25 x margin 1.10", f"{D['ah_req']:.0f} Ah -> {D['ah_sel']} Ah, 60 cells VRLA"],
        ["Chargers", D["chargers"]],
    ]))
    a("## 10. Emergency diesel generator - basis for omission\n")
    a("No emergency generator is provided. The unit is supplied by two independent 13.8 kV feeders, each able to "
      "carry the whole unit load, from the refinery main substation. On a total power failure the unit is designed "
      "to fail safe: the SIS is de-energise-to-trip, emergency isolation valves fail closed, the fired-heater BMS "
      "trips H-101/H-201 and snuffing steam is manual; there is no rotating equipment requiring post-trip power "
      "(no lube/seal-oil consoles - API 682 seals, ring-oil or rolling-bearing pumps; vacuum by steam ejectors). "
      "The essential loads (DCS, SIS/BMS, F&G, PAGA/telecom) are on the 2 x 100 % UPS with 30 min autonomy, which "
      "covers safe shutdown and operator response; escape lighting uses self-contained 90-min fittings. If the "
      "HAZOP identifies an essential motor load (e.g. an emergency cooling or flushing pump), it will be supplied "
      "from the refinery emergency power network rather than a dedicated unit generator.\n")
    a("## 11. Assumptions, holds and issues\n")
    for t_ in issues_list(ctx):
        a("- " + t_)
    if ctx.get("hac"):
        a("")
        a(ctx["hac"]["md"])
    return "\n".join(L)


def issues_list(ctx):
    R = ctx["R"]
    trl = R["tr_lv"]
    out = [
        "HOLD: utility X/R (15) and minimum fault level at 13.8 kV to be confirmed by the refinery power study; "
        "motor-start case 3 assumes 20 kA minimum.",
        "HOLD: 13.8 kV feeder route length from the refinery main substation assumed 600 m.",
        f"The 480 V double-ended substation needs {trl['onan_kVA']:,}/{trl['onaf_kVA']:,} kVA transformers and a "
        f"{R['swgr']['MCC-101']['bus_A']} A bus ({R['sc']['worst_kA']['0.48 kV']:.0f} kA calculated, "
        f"{R['sc']['rating_kA']['0.48 kV']:g} kA rated). This is at the practical limit for 480 V; recommended at "
        "detailed design to split LV into two double-ended substations (e.g. CDU / VDU+air coolers) of ~2000 kVA "
        "each, or to move the 160 kW pumps to 4.16 kV. Kept as one board here per the FEED key SLD.",
        "Long 480 V motor feeders: SS-100 is in the SW corner while the VDU pumps are ~170 m east, so the 15 % "
        "terminal-dip criterion governs (e.g. " + ", ".join(
            f"{c['to']} {c['size']:g} mm2 at {c['L_m']:.0f} m" for c in sorted(
                [c for c in ctx["cbl"] if (c.get("term_dip") or 0) > 14.0], key=lambda c: -c["term_dip"])[:4]) +
        "). Moving the >= 110 kW LV pumps to 4.16 kV, or a satellite LV substation near the VDU, would reduce copper; "
        "to be reviewed with the plot plan.",
        "Allowance loads (lighting, HVAC, heat tracing, MOVs, welding, UPS, CP, analyser house) are FEED estimates "
        "and must be replaced by vendor / discipline data.",
        "Desalter transformer load factor (0.35) assumed; desalter vendor to confirm grid power.",
        ("Cable lengths use equipment coordinates from " + ctx["loc"].source + " with Manhattan routing from SS-100 "
         "plus allowances; actual tray/trench routes to be confirmed at detailed design (lengths regenerate "
         "automatically when the layout changes)." if ctx["loc"].xy else
         "Cable lengths are route estimates from CONVENTIONS.md area blocks; re-run when data/layout.json is issued "
         "(the generator uses layout coordinates automatically when present)."),
        ("Hazardous-area classification: CFU-000-EL-HAC-001/002/003 (Section 12); motor Ex protection per "
         "location in data/electrical.json (loads[].area_class)." if ctx.get("hac") else
         "Hazardous-area classification (and Ex motor/cable gland requirements) deferred to a later pass."),
    ]
    return out + [f"Data note: {w}" for w in ctx["issues"]]


# ============================================================================================ JSON
def electrical_json(ctx, qty):
    rows, R, ud = ctx["rows"], ctx["R"], ctx["ud"]
    trm, trl, sw, sc = R["tr_mv"], R["tr_lv"], R["swgr"], R["sc"]

    def dd(d):
        return {k: _r(v, 1) if isinstance(v, float) else v for k, v in d.items()}

    buses = []
    for b, kv, swg, parent in [(B13[0], 13.8, "SWG-101", "Refinery MSS feeder 1"),
                               (B13[1], 13.8, "SWG-101", "Refinery MSS feeder 2"),
                               (B4[0], 4.16, "SWG-102", "TR-101"), (B4[1], 4.16, "SWG-102", "TR-102"),
                               (BLV[0], 0.48, "MCC-101", "TR-103"), (BLV[1], 0.48, "MCC-101", "TR-104")]:
        e = dict(id=b, kV=kv, board=swg, fed_from=parent, bus_A=sw[swg]["bus_A"], kA_rating=sw[swg]["kA"],
                 type=sw[swg]["type"], tie="N.O. tie to " + (b[:-1] + ("B" if b.endswith("A") else "A")))
        if b in R["md"]:
            e["max_demand"] = dd({k: R["md"][b][k] for k in ("C_kW", "I_kW", "S_kW", "kW", "kvar", "kVA", "pf",
                                                             "connected_kW")})
        if b in R["bus4"]:
            e["max_demand_incl_downstream"] = dd(R["bus4"][b])
        if b in R["bus13"]:
            e["max_demand_incl_downstream"] = dd(R["bus13"][b])
        e["loads"] = [r["tag"] for r in rows if r["bus"] == b]
        buses.append(e)
    for s, ab in enumerate("AB"):
        buses.append(dict(id=f"UDB-101{ab}", kV=0.208, board="UDB-101", fed_from=f"UPS-101{ab}",
                          loads=[r["tag"] for r in ctx["ups_rows"] if r["bus"] == BUPS[s]]))
    txs = []
    for t, tags in ((trm, ("TR-101", "TR-102")), (trl, ("TR-103", "TR-104"))):
        for i, tg in enumerate(tags):
            txs.append(dict(tag=tg, pri_kV=t["pri_kV"], sec_kV=t["sec_kV"], onan_kVA=t["onan_kVA"],
                            onaf_kVA=t["onaf_kVA"], z_pct=t["z_pct"], vector=t["vector"], grounding=t["grounding"],
                            type=t["type"], from_bus=(B13 if t is trm else B4)[i], to_bus=(B4 if t is trm else BLV)[i],
                            md_contingency_kVA=_r(t["md_kVA"]), loading_onan_pct=_r(t["load_onan_pct"])))
    for ab in "AB":
        txs.append(dict(tag=f"LTR-101{ab}", pri_kV=0.48, sec_kV=0.208, kVA=R["ltr"][ab]["kva"], type="Dry type",
                        from_bus=f"MCC-101{ab}"))
    for s in range(1, 3):
        for i, ab in enumerate("AB"):
            txs.append(dict(tag=f"DT-101{ab}{s}", pri_kV=0.48, sec_kV="13-23", kVA=150,
                            type="Desalter HV transformer (vendor package)", from_bus=BLV[s - 1]))
    loads = []
    for r in rows:
        c = ctx["cab"].get(f"CBL-{r['tag']}")
        sig = None
        if r["kind"] == "motor":
            sig = dict(DI=["running", "tripped", "remote_selected"] + (["vfd_fault"] if r["vfd"] else []),
                       DO=["start", "stop"],
                       AI=(["motor_current"] if r["rated_kw"] >= 30 else []) + (["speed_feedback"] if r["vfd"] else []),
                       AO=["speed_reference"] if r["vfd"] else [],
                       comms="IEC 61850 (MV relays) / Modbus-TCP (MCC IEDs)")
        loads.append(dict(tag=r["tag"], description=r["desc"], equipment=r["eq"], category=r["group"],
                          voltage_V=r["volt"], bus=r["bus"], starter=r["starter"] if r["kind"] == "motor" else "Feeder",
                          vfd=r["vfd"], rated_kW=_r(r["rated_kw"], 2), absorbed_kW=_r(r["absorbed_kw"], 2),
                          input_kW=_r(r["kw"], 2), kVA=_r(r["kva"], 2), duty=r["duty"], pair=r.get("pair"),
                          FLC_A=_r(r.get("I_fl")), mccb_A=r.get("mccb"), cable=c["cable"] if c else None,
                          cable_tag=c["tag"] if c else None, cable_length_m=c["L_m"] if c else None,
                          ct=r.get("ct"), relays=r.get("relays"), dcs_signals=sig,
                          area_class=(ctx.get("hac") or {}).get("area", {}).get(r["tag"])))
    nmot = sum(1 for r in rows if r["kind"] == "motor")
    return dict(
        doc=dict(load_list=LDL, calc=CAL, cable_schedule=CBL,
                 sld=["CFU-000-EL-SLD-001", "CFU-000-EL-SLD-002", "CFU-000-EL-SLD-003", "CFU-000-EL-SLD-004"],
                 rev=P["rev"]),
        basis=dict(utility_kV=C.KV_UT, mv_kV=C.KV_MV, lv_V=C.V_LV, ups_V=C.V_UPS, hz=C.HZ,
                   utility_fault_kA=C.UT_FAULT_KA, mv_motor_threshold_kW=C.MV_MOTOR_KW,
                   diversity=dict(C=1.0, I=0.3, S=0.1), transformer_future_margin=C.FUTURE,
                   location_source=ctx["loc"].source, substation="SS-100", far="FAR-100"),
        max_demand=dict(total_13_8kV=dd(R["total"]), lv_total=dd({k: R["md_lv_total"][k] for k in
                                                                  ("kW", "kvar", "kVA", "pf", "connected_kW")}),
                        mv_motors_total=dd({k: R["md_mv_total"][k] for k in ("kW", "kvar", "kVA", "pf",
                                                                               "connected_kW")}),
                        per_bus={b: dd({k: R["md"][b][k] for k in ("C_kW", "I_kW", "S_kW", "kW", "kvar", "kVA", "pf")})
                                 for b in BLV + B4}),
        buses=buses, transformers=txs,
        switchgear={k: dict(v) for k, v in sw.items()},
        short_circuit=dict(max_kA=dd(sc["worst_kA"]), max_peak_kA=dd(sc["worst_ip_kA"]), rating_kA=sc["rating_kA"]),
        motor_starting=[dd(s) for s in R["start"]],
        ups=dd({k: v for k, v in ud["ups"].items()}), dc=dict(ud["dc"], sections=[_r(x) for x in ud["dc"]["sections"]]),
        emergency_generator=dict(provided=False, basis="Two independent 13.8 kV feeders; fail-safe unit; essential "
                                                       "loads on 2 x 100 % UPS 30 min; see CFU-000-EL-CAL-001 s.10"),
        loads=loads,
        ups_loads=[dict(tag=r["tag"], description=r["desc"], bus=r["bus"], kVA_total=r["kva_rated"],
                        kVA_feed=_r(r["kva"], 2)) for r in ctx["ups_rows"]],
        quantities=dict(
            mv_switchgear_cubicles_13_8kV=2 + 2 + 2 + 1 + 2,       # incomers, TR feeders, spares, tie, VT
            mv_switchgear_cubicles_4_16kV=2 + sum(1 for r in rows if r["volt"] > 1000) + 2 + 2 + 1 + 2,
            mv_vfds=sum(1 for r in rows if r["vfd"] and r["volt"] > 1000),
            lv_vfds=sum(1 for r in rows if r["vfd"] and r["volt"] < 1000),
            lv_motor_starters=sum(1 for r in rows if r["kind"] == "motor" and r["volt"] < 1000 and not r["vfd"]),
            lv_feeder_buckets=sum(1 for r in rows if r["kind"] != "motor" and r["volt"] == C.V_LV) + 2 + 4,
            motors_total=nmot, motors_mv=sum(1 for r in rows if r["kind"] == "motor" and r["volt"] > 1000),
            ups_kVA=f"2 x {ud['ups']['rating_kVA']}", ups_battery_Ah=ud["ups"]["ah_sel"],
            dc_battery_Ah=ud["dc"]["ah_sel"],
            cable_m_by_type={k: _r(v, 0) for k, v in qty.items()},
            cable_m_total=_r(sum(qty.values()), 0),
            cable_m_mv=_r(sum(v for k, v in qty.items() if "kV" in k and "0.6/1" not in k), 0),
            cable_m_lv=_r(sum(v for k, v in qty.items() if "0.6/1" in k), 0)),
        cables=[{k: (_r(v, 2) if isinstance(v, float) else v) for k, v in c.items() if k not in ("R_ohm", "X_ohm")}
                for c in ctx["cbl"]],
        hazardous_area=(dict(doc=["CFU-000-EL-HAC-001", "CFU-000-EL-HAC-002", "CFU-000-EL-HAC-003"],
                             basis="API RP 505 / NFPA 70 Art. 505 / EI 15; IIA T3 general, IIB in H2S/FG services",
                             release_sources=len(ctx["hac"]["src"]),
                             clearances={c["tag"]: round(c["dist"], 1) for c in ctx["hac"]["clr"]})
                        if ctx.get("hac") else None),
        issues=issues_list(ctx),
    )


# ============================================================================================ driver
def build_all(ctx, out: Path):
    from ..docgen import render
    out.mkdir(parents=True, exist_ok=True)
    write_ldl_xlsx(ctx, out / f"{LDL}_Electrical-Load-List.xlsx")
    write_ldl_pdf(ctx, out / f"{LDL}_Electrical-Load-List.pdf")
    qty = write_cbl_xlsx(ctx, out / f"{CBL}_Cable-Schedule.xlsx")
    render(calc_md(ctx), out / f"{CAL}_Electrical-Sizing-Calculations", CAL, "Electrical System Sizing Calculations")
    js = electrical_json(ctx, qty)
    (C.DATA / "electrical.json").write_text(json.dumps(js, indent=1, default=str))
    return js
