"""Hazardous area classification (API RP 505 Zone system / NEC Art. 505; fluid categories per EI 15).

Point-source method: every release source from data/layout.json (pump seals, vessel/exchanger flanges,
control-valve stations, sample points, vents) gets Zone 1 / Zone 2 envelopes (plan shape + vertical extent).
Outputs: CFU-000-EL-HAC-001 plan (A1 1:500), CFU-000-EL-HAC-002 sections, CFU-000-EL-HAC-003 schedule (xlsx).
"""
from __future__ import annotations

import json
import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ..drawing.sheet import Sheet
from . import common as C

GRADE_EL = 100.0
# --- envelope sizes (m) - API RP 505 figures, heavier-than-air, non-enclosed adequately ventilated
PUMP_R, PUMP_H = 7.5, 7.5          # pump seals (secondary): Zone 2 7.5 m radius, grade to 7.5 m
LPG_EXT_R, LPG_EXT_H = 15.0, 0.6   # additional Zone 2 for highly volatile liquids (cat. A), 0.6 m above grade
SOUR_R = 3.0                       # small sour-water pumps (low HC content)
FLANGE_R = 3.0                     # vessels / exchangers / columns / air-cooler headers / CV stations
DRUM_A_R = 7.5                     # pressurised cat. A drums (LPG / unstabilised naphtha)
SP_Z1, SP_Z2 = 1.0, 3.0            # sample points (primary grade)
VENT_Z1, VENT_Z2 = 1.5, 4.5        # atmospheric vent (hotwell)

CAT_TXT = {
    "A": "Cat. A - flammable liquid which on release vaporises rapidly and substantially (LPG, unstabilised naphtha)",
    "B": "Cat. B - flammable liquid at a temperature sufficient for boiling on release (crude, naphtha)",
    "C": "Cat. C - flammable liquid that may be above its flash point or form a mist on release (kerosene, "
         "diesel, gas oils, residue)",
    "G": "Cat. G(i) - refinery fuel gas / off-gas (may contain H2, H2S)",
    "N": "Non-flammable (water, brine, caustic, steam) - not a release source",
}

# fluid assignment (equipment tag base -> (fluid, category, gas group))
FLUID = {
    "C-101": ("Crude vapours / atm. distillates", "B", "IIA"),
    "C-102": ("Kerosene", "C", "IIA"), "C-103": ("Diesel", "C", "IIA"), "C-104": ("AGO", "C", "IIA"),
    "C-105": ("LPG / naphtha (stabiliser)", "A", "IIA"), "C-106": ("Light / heavy naphtha", "B", "IIA"),
    "C-201": ("Vacuum gas oils / residue (sub-atmospheric)", "C", "IIA"),
    "D-101": ("Crude + wash water, 136 C, 11.5 barg", "B", "IIA"),
    "D-102": ("Unstabilised naphtha, LPG, sour water (H2S)", "A", "IIB"),
    "D-103": ("Refinery fuel gas (H2 < 30 vol% assumed)", "G", "IIB"),
    "D-104": ("Flare KO liquids / gas", "G", "IIB"),
    "D-105": ("LPG", "A", "IIA"), "D-106": ("Light naphtha", "B", "IIA"),
    "D-201": ("Hotwell: sour water, slop oil, NCG with H2S", "G", "IIB"),
    "D-202": ("Vacuum off-gas (H2S)", "G", "IIB"),
    "E-115": ("Atm. overhead vapour (H2S)", "B", "IIB"), "E-114": ("Naphtha / LPG", "A", "IIA"),
    "E-116": ("Stabiliser bottoms naphtha", "B", "IIA"), "E-117": ("Naphtha", "B", "IIA"),
    "E-118": ("Wash water / brine", "N", "-"),
    "E-202": ("Ejector condensers: steam, HC, H2S", "G", "IIB"),
    "E-203": ("Ejector condensers: steam, HC, H2S", "G", "IIB"),
    "E-204": ("Ejector condensers: steam, HC, H2S", "G", "IIB"),
    "A-101": ("Atm. overhead vapour / naphtha (H2S)", "B", "IIB"), "A-103": ("Kerosene", "C", "IIA"),
    "A-104": ("Diesel", "C", "IIA"), "A-105": ("AGO", "C", "IIA"), "A-106": ("LPG / stabiliser OH", "A", "IIA"),
    "A-107": ("Light naphtha", "B", "IIA"), "A-108": ("Heavy naphtha", "B", "IIA"),
    "A-201": ("LVGO", "C", "IIA"), "A-202": ("HVGO", "C", "IIA"),
    "P-101": ("Crude (with light ends)", "B", "IIA"), "P-102": ("Desalted crude, 135 C", "B", "IIA"),
    "P-103": ("Naphtha reflux", "B", "IIA"), "P-104": ("Unstabilised naphtha (C3/C4)", "A", "IIA"),
    "P-105": ("Sour water (H2S, HC traces)", "S", "IIB"), "P-106": ("TPA (kerosene range)", "C", "IIA"),
    "P-107": ("MPA", "C", "IIA"), "P-108": ("BPA", "C", "IIA"), "P-109": ("Kerosene", "C", "IIA"),
    "P-110": ("Diesel", "C", "IIA"), "P-111": ("AGO", "C", "IIA"),
    "P-112": ("Atm. residue > AIT (360 C)", "C", "IIA"), "P-114": ("Wash water", "N", "-"),
    "P-115": ("LPG", "A", "IIA"), "P-116": ("Light naphtha", "B", "IIA"), "P-117": ("Heavy naphtha", "B", "IIA"),
    "P-118": ("Desalter mud-wash water", "N", "-"),
    "P-201": ("LVGO", "C", "IIA"), "P-202": ("HVGO", "C", "IIA"), "P-203": ("Slop wax", "C", "IIA"),
    "P-204": ("Vacuum residue > AIT", "C", "IIA"), "P-205": ("Hotwell sour water", "S", "IIB"),
    "P-206": ("Hotwell slop oil", "C", "IIB"),
    "X-101": ("Demulsifier (aromatic solvent)", "C", "IIA"), "X-102": ("Caustic (aqueous)", "N", "-"),
    "X-103": ("Neutralising amine (flammable)", "C", "IIA"), "X-104": ("Corrosion inhibitor (solvent)", "C", "IIA"),
}
UNCLASSIFIED = {"H-101": "Fired heater - ignition source; not classified (RP 505 fired-equipment practice)",
                "H-201": "Fired heater - ignition source; not classified",
                "E-120": "Air preheater (air / flue gas only)", "K-101": "FD fan (air)", "K-102": "ID fan (flue gas)",
                "SS-100": "Substation - non-classified, pressurised", "FAR-100": "Field auxiliary room - "
                "non-classified, pressurised"}
SAMPLE_AT = {"P-101": "Crude", "P-109": "Kerosene product", "P-110": "Diesel product", "P-111": "AGO product",
             "P-116": "LN product", "P-117": "HN product", "P-115": "LPG (closed-loop sampler)",
             "P-201": "LVGO product", "P-202": "HVGO product", "P-204": "VR product", "P-112": "Atm. residue"}


def _base(tag):
    m = re.match(r"([A-Z]-\d{3})", tag)
    return m.group(1) if m else tag


def _fluid(tag):
    b = _base(tag)
    if b in FLUID:
        return FLUID[b]
    if b.startswith("E-"):
        return ("Hydrocarbon (crude / products)", "C", "IIA")
    return None


# ---------------------------------------------------------------------------------------------- geometry
class Env:
    """Zone envelope: plan shape (circle about (cx,cy) radius R, or rounded rectangle = bbox offset by R)
    and vertical extent z0..z1 (EL m)."""

    def __init__(self, zone, kind, z0, z1, R, cx=None, cy=None, bbox=None):
        self.zone, self.kind, self.z0, self.z1, self.R = zone, kind, z0, z1, R
        self.cx, self.cy, self.bbox = cx, cy, bbox

    def contains(self, x, y, z=None):
        if z is not None and not (self.z0 - 1e-6 <= z <= self.z1 + 1e-6):
            return False
        if self.kind == "circle":
            return math.hypot(x - self.cx, y - self.cy) <= self.R
        x0, y0, x1, y1 = self.bbox
        dx = max(x0 - x, 0, x - x1)
        dy = max(y0 - y, 0, y - y1)
        return math.hypot(dx, dy) <= self.R

    def dist_rect(self, r):
        """Clearance from this envelope to an axis-aligned rectangle r (negative = overlap) and closest points."""
        X0, Y0, X1, Y1 = r
        if self.kind == "circle":
            px, py = min(max(self.cx, X0), X1), min(max(self.cy, Y0), Y1)
            d = math.hypot(px - self.cx, py - self.cy)
            if d < 1e-9:
                return -self.R, (px, py), (px, py)
            ux, uy = (px - self.cx) / d, (py - self.cy) / d
            return d - self.R, (self.cx + ux * self.R, self.cy + uy * self.R), (px, py)
        x0, y0, x1, y1 = self.bbox
        gx = max(X0 - x1, x0 - X1, 0)
        gy = max(Y0 - y1, y0 - Y1, 0)
        # closest points on the two rectangles
        ax = x1 if X0 >= x1 else x0 if X1 <= x0 else (max(x0, X0) + min(x1, X1)) / 2
        bx = X0 if X0 >= x1 else X1 if X1 <= x0 else ax
        ay = y1 if Y0 >= y1 else y0 if Y1 <= y0 else (max(y0, Y0) + min(y1, Y1)) / 2
        by = Y0 if Y0 >= y1 else Y1 if Y1 <= y0 else ay
        d = math.hypot(gx, gy)
        if d < 1e-9:
            return -self.R, (ax, ay), (bx, by)
        ux, uy = (bx - ax) / d, (by - ay) / d
        return d - self.R, (ax + ux * self.R, ay + uy * self.R), (bx, by)

    def cut_x(self, xs):
        """y-interval of the plan shape on the section line x = xs (None if not cut)."""
        if self.kind == "circle":
            dx = abs(xs - self.cx)
            if dx >= self.R:
                return None
            h = math.sqrt(self.R ** 2 - dx ** 2)
            return self.cy - h, self.cy + h
        x0, y0, x1, y1 = self.bbox
        dx = max(x0 - xs, 0, xs - x1)
        if dx >= self.R:
            return None
        h = math.sqrt(self.R ** 2 - dx ** 2)
        return y0 - h, y1 + h

    def extent_txt(self):
        return f"{'R' if self.kind == 'circle' else 'offset '}{self.R:g} m, EL {self.z0:.1f}-{self.z1:.1f}"


def _offset_bbox(b, d=0.0):
    return [b[0] - d, b[1] - d, b[2] + d, b[3] + d]


def build_sources(L, eq, loops):
    items = {e["tag"]: e for e in L["equipment"]}
    eqd = {e["tag"]: e for e in eq}
    src = []

    def opdata(tag):
        e = eqd.get(tag) or eqd.get(_base(tag) + "A/B") or eqd.get(_base(tag)) or {}
        T = e.get("op_T", e.get("T"))
        P = e.get("op_P", "")
        return (str(T) if T is not None else ""), str(P)

    def add(sid, kind, desc, equip, fluid, cat, group, grade, envs, basis, x, y, el, remarks=""):
        T, P = opdata(equip)
        src.append(dict(id=sid, kind=kind, desc=desc, equip=equip, fluid=fluid, cat=cat, group=group,
                        tclass="T3", grade=grade, vent="Open, adequate (natural)", envs=envs, basis=basis,
                        x=x, y=y, el=el, T=T, P=P, remarks=remarks))

    n = 0
    for it in L["equipment"]:
        tag = it["tag"]
        b = _base(tag)
        if b in UNCLASSIFIED or it["type"] in ("Fired heater", "Fan", "Air preheater"):
            continue
        fl = _fluid(tag)
        if fl is None or fl[1] == "N":
            continue
        fluid, cat, grp = fl
        bb = it["bbox"]
        top = it["z_base"] + (it.get("height") or 2.0)
        if it["type"] == "Pump":
            n += 1
            dn = it["nozzles"].get("discharge", [it["x"], it["y"], 101.0])
            seal = (it["x"], it["y"])
            R = SOUR_R if cat == "S" else PUMP_R
            envs = [Env("Z2", "circle", GRADE_EL, GRADE_EL + (SOUR_R + 1.0 if cat == "S" else PUMP_H), R, *seal)]
            basis = "RP 505 pump figure (HTA, open area); API 682 seals" if cat != "S" else \
                "Low HC content sour water; EI 15 small release"
            if cat == "A":
                envs.append(Env("Z2X", "circle", GRADE_EL, GRADE_EL + LPG_EXT_H, LPG_EXT_R, *seal))
                basis += "; additional Zone 2 15 m x 0.6 m (highly volatile liquid)"
            add(f"RS-{n:03d}", "Pump seal", f"{tag} mechanical seal", tag, fluid, "C" if cat == "S" else cat, grp,
                "Secondary", envs, basis, seal[0], seal[1], GRADE_EL + 0.8)
            if b in SAMPLE_AT and tag.endswith("A"):
                n += 1
                sx, sy = dn[0] + 1.5, dn[1] + (2.5 if it["y"] > 75 else -2.5)
                add(f"RS-{n:03d}", "Sample point", f"{SAMPLE_AT[b]} sample point at {b}A/B discharge", tag, fluid,
                    cat, grp, "Primary",
                    [Env("Z1", "circle", GRADE_EL, GRADE_EL + 2.0, SP_Z1, sx, sy),
                     Env("Z2", "circle", GRADE_EL, GRADE_EL + 4.0, SP_Z2, sx, sy)],
                    "Sample point (manual, intermittent) - EI 15 / RP 505", sx, sy, GRADE_EL + 1.0)
            continue
        if it["type"] == "Air cooler":
            n += 1
            add(f"RS-{n:03d}", "Header box plugs / flanges", f"{tag} headers on rack structure",
                it.get("parent_tag", tag), fluid, cat, grp, "Secondary",
                [Env("Z2", "rrect", it["z_base"] - FLANGE_R, top + FLANGE_R, FLANGE_R, bbox=bb)],
                "RP 505 elevated flanged equipment, 3 m around (no grade extension: covered by rack/pump sources)",
                it["x"], it["y"], it["z_base"])
            continue
        if it["type"] == "Ejector":
            continue        # covered by the ejector structure source (E-202..204)
        n += 1
        R = DRUM_A_R if (cat == "A" and it["type"] == "Drum") else FLANGE_R
        kind = "Vessel flanges / instruments"
        if it["type"] == "Shell & tube":
            kind = "Exchanger flanges"
        if it["shape"] == "vcyl":
            r0 = it["L"] / 2
            envs = [Env("Z2", "circle", GRADE_EL, top + R, r0 + R, it["x"], it["y"])]
        else:
            envs = [Env("Z2", "rrect", GRADE_EL, top + R, R, bbox=bb)]
        basis = "RP 505 vessel / flanged equipment, open area, 3 m" if R == FLANGE_R else \
            "Pressurised cat. A drum: 7.5 m (EI 15 hazard radius, LPG flange / instrument releases)"
        if cat == "A" and it["type"] == "Drum":
            envs.append(Env("Z2X", "circle", GRADE_EL, GRADE_EL + LPG_EXT_H, LPG_EXT_R, it["x"], it["y"]))
            basis += "; additional Zone 2 15 m x 0.6 m"
        if tag in ("E-202", "E-203", "E-204"):
            if tag != "E-202":
                continue
            st = L["structures"]["ejector_structure"]
            envs = [Env("Z2", "rrect", GRADE_EL, st["top"] + FLANGE_R, FLANGE_R,
                        bbox=[st["x0"], st["y0"], st["x1"], st["y1"]])]
            kind, tag = "Ejector / condenser flanges", "J-201..203 / E-202..204 (ST-201)"
            basis = "Flanged ejector set on structure, 3 m around structure"
        add(f"RS-{n:03d}", kind, f"{tag} {it.get('service', '')}".strip(), tag, fluid, cat, grp, "Secondary", envs,
            basis, it["x"], it["y"], it["z_base"])
        # primary-grade sources on vessels
        if b == "D-101":
            n += 1
            sx, sy = bb[0] - 1.0, it["y"]
            add(f"RS-{n:03d}", "Sample point", f"{tag} interface try-cocks / profiler sample", tag, fluid, cat, grp,
                "Primary", [Env("Z1", "circle", GRADE_EL, GRADE_EL + 2.5, SP_Z1, sx, sy),
                            Env("Z2", "circle", GRADE_EL, GRADE_EL + 4.0, SP_Z2, sx, sy)],
                "Sample point (manual) - EI 15", sx, sy, GRADE_EL + 1.5)
        if tag == "D-102":
            n += 1
            sx, sy = it["x"] + 4.0, bb[1] - 1.0
            add(f"RS-{n:03d}", "Sample point", "D-102 boot sour-water / naphtha sample", tag, fluid, cat, "IIB",
                "Primary", [Env("Z1", "circle", GRADE_EL, GRADE_EL + 2.0, SP_Z1, sx, sy),
                            Env("Z2", "circle", GRADE_EL, GRADE_EL + 4.0, SP_Z2, sx, sy)],
                "Sample point (manual) - EI 15", sx, sy, GRADE_EL + 1.0)
        if tag == "D-201":
            n += 1
            vx, vy, vz = it["x"] + 2.0, it["y"], it["z_base"] + it.get("height", 1.3) + 3.0
            add(f"RS-{n:03d}", "Atmospheric vent", "D-201 hotwell seal-leg / atmospheric vent outlet", tag, fluid,
                cat, "IIB", "Primary",
                [Env("Z1", "circle", vz - VENT_Z1, vz + VENT_Z1, VENT_Z1, vx, vy),
                 Env("Z2", "circle", GRADE_EL, vz + VENT_Z2, VENT_Z2, vx, vy)],
                "RP 505 atmospheric vent: Zone 1 1.5 m, Zone 2 4.5 m to grade (HTA)", vx, vy, vz,
                remarks="Vent outlet 3 m above drum top; route to H-201 off-gas system preferred (HAZOP)")
    # control-valve stations from the control-loop list
    stations = {}
    skip = re.compile(r"steam|BFW|condensate|wash water|brine|Caustic|Demulsifier|neutraliser|stroke", re.I)
    for lp in loops:
        fin = lp.get("final", "")
        vs = re.findall(r"\b([A-Z]{1,3}V-\d{4}[AB]?)", fin)
        if not vs or skip.search(lp["service"] + " " + fin + " " + lp["measured"]):
            continue
        if re.search(r"fuel gas|FG to burners", lp["service"] + lp["measured"], re.I):
            continue          # FG valve trains at heaters: see note (fired-equipment practice)
        m = re.findall(r"\b([CDPEXAH]-\d{3}[AB]?)", lp["service"] + " " + lp["measured"] + " " + fin)
        if not m:
            continue
        key = _base(m[0])
        stations.setdefault(key, []).extend(vs)
    for key, vs in sorted(stations.items()):
        cands = [it for it in L["equipment"] if _base(it["tag"]) == key]
        if not cands:
            continue
        it = cands[0]
        if key.startswith("H-"):
            x, y = it["x"], 62.0                   # heater pass-flow valves at rack-side manifold
            fl = ("Heater charge (crude / atm. residue)", "C", "IIA")
        elif key.startswith("P-"):
            x, y = it["x"] + 1.5, it["y"] + (4.5 if it["y"] > 75 else -4.5)
            fl = _fluid(key) or ("Hydrocarbon", "C", "IIA")
        else:
            bb = it["bbox"]
            x = (bb[0] + bb[2]) / 2
            y = bb[1] - 2.0 if it["y"] > 75 else bb[3] + 2.0
            fl = _fluid(key) or ("Hydrocarbon", "C", "IIA")
        if fl[1] == "N":
            continue
        n += 1
        add(f"RS-{n:03d}", "Control-valve station", f"Control valves {', '.join(sorted(set(vs)))} ({key})", key,
            fl[0], "C" if fl[1] == "S" else fl[1], fl[2], "Secondary",
            [Env("Z2", "circle", GRADE_EL, GRADE_EL + 4.0, FLANGE_R, x, y)],
            "RP 505 valve stations (stem packing, flanges), 3 m", x, y, GRADE_EL + 1.0,
            remarks="Indicative location (P&ID / piping layout not yet issued)")
    return src


# ---------------------------------------------------------------------------------------------- analysis
def zone_at(src, x, y, z=None):
    best, grp = "Unclassified", "-"
    rank = {"Z1": 2, "Z2": 1, "Z2X": 1}
    r = 0
    for s in src:
        for e in s["envs"]:
            if e.contains(x, y, z) and rank[e.zone] > r:
                r = rank[e.zone]
                best = "Zone 1" if e.zone == "Z1" else "Zone 2"
            if e.contains(x, y, z):
                grp = max(grp, s["group"]) if grp != "-" else s["group"]
    return best, grp


def clearances(src, L):
    out = []
    items = {e["tag"]: e for e in L["equipment"]}
    targets = [(t, items[t]["bbox"]) for t in ("H-101", "H-201", "E-120") if t in items]
    for b in L["structures"]["buildings"]:
        targets.append((b["id"], [b["x0"], b["y0"], b["x1"], b["y1"]]))
    for tag, r in targets:
        best = None
        for s in src:
            for e in s["envs"]:
                d, pa, pb = e.dist_rect(r)
                if best is None or d < best[0]:
                    best = (d, pa, pb, s, e)
        out.append(dict(tag=tag, bbox=r, dist=best[0], pa=best[1], pb=best[2], src=best[3]["id"],
                        src_desc=best[3]["desc"], zone=best[4].zone))
    return out


# ---------------------------------------------------------------------------------------------- drawing helpers
Z2_FILL, Z2_LINE = "#FFE699", "#C55A11"
Z1_FILL, Z1_LINE = "#F4B183", "#C00000"
ZX_FILL = "#FFF5D6"


def _patterns(d):
    defs = d.defs
    p2 = d.pattern(id="hz2", size=(2.4, 2.4), patternUnits="userSpaceOnUse", patternTransform="rotate(45)")
    p2.add(d.rect((0, 0), (2.4, 2.4), fill=Z2_FILL))
    p2.add(d.line((0, 0), (0, 2.4), stroke=Z2_LINE, stroke_width=0.35))
    defs.add(p2)
    p1 = d.pattern(id="hz1", size=(1.6, 1.6), patternUnits="userSpaceOnUse", patternTransform="rotate(45)")
    p1.add(d.rect((0, 0), (1.6, 1.6), fill=Z1_FILL))
    p1.add(d.line((0, 0), (0, 1.6), stroke=Z1_LINE, stroke_width=0.35))
    p1.add(d.line((0, 0), (1.6, 0), stroke=Z1_LINE, stroke_width=0.35))
    defs.add(p1)
    px = d.pattern(id="hzx", size=(2.0, 2.0), patternUnits="userSpaceOnUse")
    px.add(d.rect((0, 0), (2.0, 2.0), fill=ZX_FILL))
    px.add(d.circle((1.0, 1.0), 0.28, fill=Z2_LINE))
    defs.add(px)


STYLE = {"Z2X": ("url(#hzx)", Z2_LINE, "2,1"), "Z2": ("url(#hz2)", Z2_LINE, None), "Z1": ("url(#hz1)", Z1_LINE, None)}


def _d_circle(cx, cy, r):
    return (f"M {cx - r:.3f} {cy:.3f} A {r:.3f} {r:.3f} 0 1 1 {cx + r:.3f} {cy:.3f} "
            f"A {r:.3f} {r:.3f} 0 1 1 {cx - r:.3f} {cy:.3f} Z ")


def _d_rrect(x, y, w, h, r):
    r = min(r, w / 2, h / 2)
    return (f"M {x + r:.3f} {y:.3f} H {x + w - r:.3f} A {r:.3f} {r:.3f} 0 0 1 {x + w:.3f} {y + r:.3f} "
            f"V {y + h - r:.3f} A {r:.3f} {r:.3f} 0 0 1 {x + w - r:.3f} {y + h:.3f} H {x + r:.3f} "
            f"A {r:.3f} {r:.3f} 0 0 1 {x:.3f} {y + h - r:.3f} V {y + r:.3f} A {r:.3f} {r:.3f} 0 0 1 {x + r:.3f} "
            f"{y:.3f} Z ")


def _draw_union(d, g, paths, zone, lw=0.55):
    """Union of closed sub-paths (all clockwise): thick outlines first, then ONE compound path filled with
    nonzero winding on top -> inner outlines hidden and no seams."""
    if not paths:
        return
    fill, line, dash = STYLE[zone]
    for pd in paths:
        el = d.path(d=pd, fill="none", stroke=line, stroke_width=lw * 2)
        if dash:
            el["stroke-dasharray"] = dash
        g.add(el)
    solid = {"Z2": Z2_FILL, "Z1": Z1_FILL, "Z2X": ZX_FILL}[zone]
    g.add(d.path(d="".join(paths), fill=solid, stroke="none", fill_rule="nonzero"))
    g.add(d.path(d="".join(paths), fill=fill, stroke="none", fill_rule="nonzero"))


# ---------------------------------------------------------------------------------------------- HAC-001 plan
def hac_plan(ctx, L, src, clr, out: Path):
    from ..layout import plotplan as PP
    SC, X0, Y0 = PP.SC, PP.X0, PP.Y0
    P = PP.P
    notes = [
        "Classification per API RP 505 (Zone system) and NFPA 70 Art. 505; release-source extents from RP 505 "
        "figures for heavier-than-air vapours in open, adequately ventilated (natural) areas; fluid categories per "
        "EI 15 (Model Code of Safe Practice Pt 15).",
        "Point-source method (no blanket classification). Areas not shown hatched are unclassified.",
        "Gas group IIA, temperature class T3 generally; IIB where H2S / fuel gas / off-gas is present (overhead "
        "system D-102 / A-101 / E-115, hotwell D-201, D-202, D-103, D-104, ejector set, sour-water pumps).",
        "Fired heaters H-101 / H-201 are ignition sources and are not classified; fuel-gas valve trains are of "
        "welded construction with minimum flanges; electrical equipment within 3 m of FG flanges shall be Zone 2 "
        "rated. Distances from the nearest Zone 2 boundary are dimensioned.",
        "PSVs discharge to the closed flare system - no atmospheric release envelope. Drains are closed "
        "(to OWS / sour water); below-grade pits, trenches and catch basins within a Zone 2 are Zone 1.",
        "Pipe rack: only flanges, valves and instrument connections are sources (3 m Zone 2); straight welded runs "
        "are not sources. Control-valve stations are shown at indicative locations pending P&IDs / piping layout.",
        "Pump seals are secondary-grade releases (API 682 seals, open area) -> Zone 2 only; Zone 1 applies to "
        "primary-grade sources (sample points, try-cocks, atmospheric vents).",
        "SS-100 and FAR-100 are non-classified pressurised buildings; HVAC intakes shall be located in "
        "unclassified area (>= 3 m from any Zone 2 boundary). Gas detection at intakes.",
        "Equipment surface temperatures of process plant may exceed fluid AIT (hot oil > 260 C); T-class applies to "
        "electrical equipment. Release source schedule: CFU-000-EL-HAC-003; sections: CFU-000-EL-HAC-002.",
    ]
    sh = Sheet("A1", "HAZARDOUS AREA CLASSIFICATION", "PLAN - ZONE 1 / ZONE 2 (API RP 505)", "CFU-000-EL-HAC-001",
               scale="1:500", discipline="ELECTRICAL")
    from .sld import put_notes, table
    from .symbols import Pen
    put_notes(sh, notes, width=104)
    d = sh.dwg
    _patterns(d)
    g = d.g()
    d.add(g)
    p = Pen(sh)
    # base: plot, roads, rack, buildings
    a, b = P(0, 150)
    g.add(d.rect((a, b), (230 * SC, 150 * SC), fill="white", stroke="black", stroke_width=0.6))
    for r in L["structures"]["roads"]:
        a, b = P(r["x0"], r["y1"])
        g.add(d.rect((a, b), ((r["x1"] - r["x0"]) * SC, (r["y1"] - r["y0"]) * SC), fill="#EDEDED", stroke="#AAAAAA",
                     stroke_width=0.2))
    rk = L["structures"]["pipe_rack"]
    a, b = P(rk["x0"], rk["y1"])
    g.add(d.rect((a, b), ((rk["x1"] - rk["x0"]) * SC, (rk["y1"] - rk["y0"]) * SC), fill="none", stroke="#666666",
                 stroke_width=0.3, stroke_dasharray="4,1.5"))
    # zones
    gz = d.g()
    d.add(gz)
    for zone in ("Z2X", "Z2", "Z1"):
        shapes = []
        for s in src:
            for e in s["envs"]:
                if e.zone != zone:
                    continue
                if e.kind == "circle":
                    c, R = P(e.cx, e.cy), e.R * SC
                    shapes.append(_d_circle(c[0], c[1], R))
                else:
                    x0, y0, x1, y1 = e.bbox
                    a_, b_ = P(x0 - e.R, y1 + e.R)
                    w_, h_ = (x1 - x0 + 2 * e.R) * SC, (y1 - y0 + 2 * e.R) * SC
                    shapes.append(_d_rrect(a_, b_, w_, h_, e.R * SC))
        _draw_union(d, gz, shapes, zone)
    # equipment on top
    ge = d.g()
    d.add(ge)
    for it in L["equipment"]:
        try:
            PP.draw_item(sh, d, ge, it)
        except (KeyError, TypeError):
            x0, y0, x1, y1 = it["bbox"]
            a, b = P(x0, y1)
            ge.add(d.rect((a, b), ((x1 - x0) * SC, (y1 - y0) * SC), fill="white", stroke="black", stroke_width=0.3))
    for bl in L["structures"]["buildings"]:
        a, b = P(bl["x0"], bl["y1"])
        ge.add(d.rect((a, b), ((bl["x1"] - bl["x0"]) * SC, (bl["y1"] - bl["y0"]) * SC), fill="#E2F0D9",
                      stroke="#375623", stroke_width=0.5))
        cx, cy = P((bl["x0"] + bl["x1"]) / 2, (bl["y0"] + bl["y1"]) / 2)
        p.text(bl["id"], cx, cy - 0.5, 2.4, "middle", bold=True, color="#375623")
        p.text("NON-CLASSIFIED", cx, cy + 2.6, 1.7, "middle", color="#375623")
    for t in ("H-101", "H-201"):
        it = next(e for e in L["equipment"] if e["tag"] == t)
        cx, cy = P(it["x"], it["bbox"][1])
        p.text("UNCLASSIFIED - IGNITION SOURCE", cx, cy - 2.2, 1.6, "middle", bold=True, color="#C00000")
    # sample point / vent markers
    for s in src:
        if s["kind"] in ("Sample point", "Atmospheric vent"):
            c = P(s["x"], s["y"])
            p.text("SP" if s["kind"] == "Sample point" else "V", c[0] + 2.2, c[1] - 1.2, 1.5, bold=True,
                   color=Z1_LINE)
        if s["kind"] == "Control-valve station":
            c = P(s["x"], s["y"])
            ge.add(d.rect((c[0] - 0.9, c[1] - 0.9), (1.8, 1.8), fill="white", stroke="#7F6000", stroke_width=0.3))
            ge.add(d.line((c[0] - 0.9, c[1] - 0.9), (c[0] + 0.9, c[1] + 0.9), stroke="#7F6000", stroke_width=0.3))
            ge.add(d.line((c[0] - 0.9, c[1] + 0.9), (c[0] + 0.9, c[1] - 0.9), stroke="#7F6000", stroke_width=0.3))
    # clearance dimensions
    for c in clr:
        pa, pb = P(*c["pa"]), P(*c["pb"])
        ge.add(d.line(pa, pb, stroke="#0050A0", stroke_width=0.35, stroke_dasharray="1.5,0.8"))
        for q in (pa, pb):
            ge.add(d.circle(q, 0.6, fill="#0050A0"))
        mx, my = (pa[0] + pb[0]) / 2, (pa[1] + pb[1]) / 2
        p.text(f"{c['dist']:.1f} m", mx + 1.2, my - 0.8, 2.0, bold=True, color="#0050A0")
    # section marks
    for xs, nm, (y0_, y1_) in SECTIONS:
        a, b = P(xs, y1_ + 3)
        c, e = P(xs, y0_ - 3)
        ge.add(d.line((a, b), (c, e), stroke="#000000", stroke_width=0.5, stroke_dasharray="6,1.5,1.5,1.5"))
        for (qx, qy), up in (((a, b), True), ((c, e), False)):
            ge.add(d.polygon([(qx, qy), (qx - 3.5, qy - (1.6 if up else -1.6)), (qx - 3.5, qy + (1.6 if up else -1.6))],
                             fill="black"))
            p.text(nm, qx + 1.5, qy + (-1.5 if up else 4.0), 3.2, bold=True)
    # north arrow & grid ticks
    PP.north_arrow(sh, d, 480, 36)
    for xg in range(0, 231, 10):
        a, b = P(xg, 0)
        g.add(d.line((a, b), (a, b + 1.5), stroke="black", stroke_width=0.2))
        if xg % 50 == 0:
            p.text(f"E{xg}", a, b + 4.5, 1.8, "middle")
    for yg in range(0, 151, 10):
        a, b = P(0, yg)
        g.add(d.line((a - 1.5, b), (a, b), stroke="black", stroke_width=0.2))
        if yg % 50 == 0:
            p.text(f"N{yg}", a - 2.2, b + 0.7, 1.8, "end")
    p.text("SCALE 1:500 (A1)  -  0     10     20     30     40     50 m", 36, 384, 2.2, bold=True)
    for i in range(6):
        g.add(d.rect((70 + i * 20, 386.5), (20, 1.6), fill="black" if i % 2 == 0 else "white", stroke="black",
                     stroke_width=0.2))
    # ---- right column: legend + fluid categories + summary
    xl = 505
    p.rect(xl, 18, 318, 66, lw=0.35)
    p.text("LEGEND", xl + 3, 23.5, 2.8, bold=True)
    leg = [("Z1", "ZONE 1 - flammable atmosphere likely in normal operation (primary-grade sources)"),
           ("Z2", "ZONE 2 - not likely in normal operation; short duration if it occurs (secondary grade)"),
           ("Z2X", "ADDITIONAL ZONE 2, 0.6 m above grade (cat. A / highly volatile liquids, 15 m)")]
    for i, (z, t) in enumerate(leg):
        yy = 28 + i * 9
        fill, line, dash = STYLE[z]
        el = d.rect((xl + 4, yy), (14, 6.5), fill=fill, stroke=line, stroke_width=0.6)
        if dash:
            el["stroke-dasharray"] = dash
        d.add(el)
        p.text(t, xl + 21, yy + 4.5, 2.1)
    yy = 28 + 3 * 9
    p.rect(xl + 4, yy, 14, 6.5, fill="#E2F0D9", lw=0.4)
    p.text("NON-CLASSIFIED BUILDING (pressurised)  /  unhatched area = unclassified", xl + 21, yy + 4.5, 2.1)
    yy += 9
    d.add(d.rect((xl + 9.6, yy + 2.1), (2.4, 2.4), fill="white", stroke="#7F6000", stroke_width=0.3))
    p.text("CV station (indicative)    SP sample point    V atmospheric vent    - - -  dimension to nearest Zone 2",
           xl + 21, yy + 4.5, 2.1)
    # fluid categories
    rows = [[k, v.split(" - ", 1)[1] if " - " in v else v] for k, v in CAT_TXT.items()]
    table(p, xl, 96, [("CAT.", 12, "l"), ("FLUID CATEGORY (EI 15)", 306, "l")], rows,
          "FLUID CATEGORIES", size=2.0, rh=4.6)
    # condensed release-source summary
    summ = {}
    for s in src:
        k = (s["kind"], s["cat"], s["group"], "/".join(sorted({e.zone.replace("Z2X", "Z2 ext") for e in s["envs"]})),
             "; ".join(sorted({f"{e.R:g} m" for e in s["envs"]})))
        summ.setdefault(k, []).append(s["equip"])
    srows = []
    for (kind, cat, grp, zones, ext), eqs in sorted(summ.items()):
        tags = sorted({_base(e) if e.startswith(("P-", "A-")) else e for e in eqs})
        txt = ", ".join(tags)
        if len(txt) > 64:
            txt = txt[:61] + "..."
        srows.append([kind, cat, grp, zones.replace("Z", "Zone "), ext, txt])
    table(p, xl, 136, [("SOURCE TYPE", 52, "l"), ("CAT", 10, "c"), ("GRP", 11, "c"), ("ZONES", 32, "l"),
                       ("EXTENT", 30, "l"), ("EQUIPMENT", 183, "l")], srows,
          "RELEASE SOURCE SUMMARY (full schedule CFU-000-EL-HAC-003)", size=1.85, rh=4.1)
    # clearance table under plot
    crow = [[c["tag"], f"{c['dist']:.1f}", c["src"], c["src_desc"][:60], "OK" if c["dist"] > 0 else "OVERLAP"]
            for c in clr]
    table(p, 36, 398, [("UNCLASSIFIED ITEM", 30, "l"), ("CLEARANCE m", 22, "r"), ("NEAREST SOURCE", 26, "l"),
                       ("DESCRIPTION", 110, "l"), ("CHECK", 18, "c")], crow,
          "CLEARANCE OF IGNITION SOURCES / BUILDINGS TO NEAREST CLASSIFIED AREA", size=2.1, rh=4.8)
    # protection summary
    prow = [[r[0], r[1], r[2]] for r in PROTECTION[:7]]
    table(p, 250, 398, [("EQUIPMENT", 46, "l"), ("ZONE 1", 82, "l"), ("ZONE 2", 82, "l")], prow,
          "MINIMUM EQUIPMENT PROTECTION (IEC 60079-14 / NEC 505)", size=1.9, rh=4.6)
    d.elements.remove(sh.g)          # texts / labels on top of the zone layers
    d.elements.append(sh.g)
    sh.save(out / "CFU-000-EL-HAC-001")
    return out / "CFU-000-EL-HAC-001.svg"


PROTECTION = [
    ("LV motors (480 V)", "Ex db or Ex eb (tE, stall relay), IIB T3 (Gb)", "Ex ec (non-sparking), IIA/IIB T3 (Gc); "
                                                                          "AEx ec per NEC 505"),
    ("MV motors (4.16 kV)", "Not located in Zone 1 (relocate) / Ex pxb", "Ex ec with stator discharge risk "
                                                                         "assessment / pre-start purge, IIA T3"),
    ("VFD-fed motors", "Ex db certified with the drive (converter duty)", "Ex ec certified for converter duty; "
                                                                          "T-class verified with VFD"),
    ("Instruments", "Ex ia / ib (intrinsic safety), Ex db", "Ex ia / ic, Ex db, Ex ec"),
    ("Lighting fittings", "Ex db / eb, IIB T3", "Ex ec / nR, IIA T3 (IIB in H2S areas)"),
    ("Junction boxes / glands", "Ex eb boxes; Ex db barrier glands for Ex d", "Ex eb / ec boxes and glands"),
    ("Heat tracing", "Ex 60079-30-1, stabilised design <= T3", "Ex 60079-30-1, self-regulating, T3"),
    ("Desalter transformers", "n/a (Zone 2 only)", "Vendor package certified for Zone 2 (Ex o / ec), HV entry to "
                                                   "vessel via Ex bushings"),
    ("Welding / receptacles", "Not permitted", "Ex de interlocked receptacles; hot-work permit"),
    ("Buildings SS-100 / FAR-100", "-", "Non-classified; pressurised (NFPA 496 / IEC 60079-13), intakes in "
                                        "unclassified area, gas detection"),
]

# sections: (x of cut, name, (y_min, y_max) of view)
SECTIONS = [(79.0, "A", (56.0, 112.0)), (64.0, "B", (60.0, 116.0))]


# ---------------------------------------------------------------------------------------------- HAC-002 sections
def hac_sections(ctx, L, src, out: Path):
    from .sld import put_notes, table
    from .symbols import Pen
    notes = [
        "Sections are cut at the section lines shown on CFU-000-EL-HAC-001 (looking west). Equipment within "
        "+/- 12 m of the cut is projected (light outline); zones are shown only where the cut plane intersects them.",
        "Vertical extents (heavier-than-air, open area): pump seals Zone 2 to 7.5 m above grade within 7.5 m "
        "radius; vessels / columns / exchangers Zone 2 3 m beyond the equipment, from grade to 3 m above the top "
        "source; air-cooler headers 3 m around the bundles; sample points Zone 1 1 m / Zone 2 3 m; cat. A "
        "additional Zone 2 15 m radius x 0.6 m above grade.",
        "Pipe-rack tiers (EL 106.0 / 108.5 / 111.0) are unclassified except within 3 m of flanges / valves and "
        "where inside the pump / equipment envelopes shown.",
        "Scale 1:200 (A1). Elevations in m, grade EL 100.000.",
    ]
    sh = Sheet("A1", "HAZARDOUS AREA CLASSIFICATION", "SECTIONS A-A (D-102 / OH PUMP ROW) & B-B (LPG AREA)",
               "CFU-000-EL-HAC-002", scale="1:200", discipline="ELECTRICAL")
    put_notes(sh, notes, width=104)
    d = sh.dwg
    _patterns(d)
    p = Pen(sh)
    S = 5.0                            # mm per m
    items = L["equipment"]
    rk = L["structures"]["pipe_rack"]
    zmax = 150.0
    for k, (xs, nm, (ya, yb)) in enumerate(SECTIONS):
        ox = 28 + k * 395              # sheet x of y = ya
        oz = 300.0                     # sheet y of grade

        def Q(y, z, ox=ox, ya=ya):
            return ox + (y - ya) * S, oz - (z - GRADE_EL) * S

        W = (yb - ya) * S
        g = d.g()
        d.add(g)
        # zones
        for zone in ("Z2X", "Z2", "Z1"):
            shapes = []
            for s in src:
                for e in s["envs"]:
                    if e.zone != zone:
                        continue
                    cut = e.cut_x(xs)
                    if not cut:
                        continue
                    y0, y1 = max(cut[0], ya), min(cut[1], yb)
                    if y1 <= y0:
                        continue
                    z1 = min(e.z1, zmax)
                    a_, b_ = Q(y0, z1)
                    w_, h_ = (y1 - y0) * S, (z1 - e.z0) * S
                    if e.zone == "Z1" and e.z0 > GRADE_EL + 0.5:      # elevated vent sphere
                        cy_, cz_ = (y0 + y1) / 2, (e.z0 + e.z1) / 2
                        c = Q(cy_, cz_)
                        shapes.append(_d_circle(c[0], c[1], (y1 - y0) / 2 * S))
                    else:
                        shapes.append(_d_rrect(a_, b_, w_, h_, 0.01))
            _draw_union(d, g, shapes, zone, lw=0.5)
        ge = d.g(fill="none", stroke="black", stroke_width=0.4)
        d.add(ge)
        # grade
        a_, b_ = Q(ya, GRADE_EL)
        ge.add(d.line((a_ - 5, b_), (a_ + W + 5, b_), stroke_width=0.8))
        for i in range(int(W / 4) + 2):
            ge.add(d.line((a_ - 5 + i * 4, b_), (a_ - 5 + i * 4 - 2, b_ + 2), stroke_width=0.25))
        # rack
        for el in (106.0, 108.5, 111.0):
            ge.add(d.line(Q(rk["y0"], el), Q(rk["y1"], el), stroke_width=0.6))
        for yy in (rk["y0"], rk["y1"]):
            ge.add(d.line(Q(yy, GRADE_EL), Q(yy, 113.0), stroke_width=0.6))
        p.text("PIPE RACK PR-100", *Q((rk["y0"] + rk["y1"]) / 2, 111.6), size=2.0, anchor="middle")
        for el in (106.0, 108.5, 111.0):
            p.text(f"EL {el:.1f}", *Q(rk["y1"] + 0.4, el + 0.3), size=1.7)
        # equipment
        for it in items:
            x0, y0_, x1, y1_ = it["bbox"]
            if x1 < xs - 12 or x0 > xs + 12 or y1_ < ya or y0_ > yb:
                continue
            cut = x0 <= xs <= x1
            z0 = it["z_base"]
            h = it.get("height") or 2.0
            if it["type"] == "Fired heater":
                continue
            y0c, y1c = max(y0_, ya), min(y1_, yb)
            a_, b_ = Q(y0c, z0 + h)
            col = "black" if cut else "#8C8C8C"
            if it["shape"] == "vcyl":
                ge.add(d.rect((a_, b_), ((y1c - y0c) * S, h * S), rx=1.5, fill="white" if cut else "none",
                              stroke=col, stroke_width=0.5 if cut else 0.3,
                              stroke_dasharray=None if cut else "2,1"))
            elif it["shape"] == "hcyl":
                ge.add(d.rect((a_, b_), ((y1c - y0c) * S, h * S), rx=min(h * S / 2, 4), fill="white" if cut else "none",
                              stroke=col, stroke_width=0.5 if cut else 0.3, stroke_dasharray=None if cut else "2,1"))
                if z0 > GRADE_EL + 1:
                    for yy in (y0c + 0.6, y1c - 0.6):
                        ge.add(d.line(Q(yy, GRADE_EL), Q(yy, z0), stroke=col, stroke_width=0.3))
            else:
                ge.add(d.rect((a_, b_), ((y1c - y0c) * S, h * S), fill="white" if cut else "none", stroke=col,
                              stroke_width=0.5 if cut else 0.3, stroke_dasharray=None if cut else "2,1"))
            if cut or it["type"] in ("Column", "Drum"):
                tx, tz = Q((y0c + y1c) / 2, z0 + h + 0.8)
                p.text(it.get("parent_tag", it["tag"]) if it["type"] == "Air cooler" else it["tag"], tx, tz, 2.0,
                       "middle", bold=cut, color=col)
        # elevation scale
        for el in range(100, int(zmax) + 1, 5):
            a_, b_ = Q(ya, el)
            ge.add(d.line((a_ - 8, b_), (a_ - 6.5, b_), stroke_width=0.3))
            p.text(f"{el}", a_ - 9, b_ + 0.7, 1.8, "end")
        ge.add(d.line(Q(ya, GRADE_EL), Q(ya, zmax), stroke_width=0.2, stroke_dasharray="1,1"))
        for yy in range(int(math.ceil(ya / 5) * 5), int(yb) + 1, 5):
            a_, b_ = Q(yy, GRADE_EL)
            p.text(f"N{yy}", a_, b_ + 5.0, 1.8, "middle")
        p.text(f"SECTION {nm}-{nm}  (cut at E {xs:g}, looking west)  SCALE 1:200", Q(ya, zmax)[0], 32, 3.4, bold=True)
        p.text("SOUTH", *Q(ya, GRADE_EL - 2.2), size=2.0)
        p.text("NORTH", *Q(yb, GRADE_EL - 2.2), size=2.0, anchor="end")
    # legend + protection table
    xl = 28
    for i, (z, t) in enumerate([("Z1", "ZONE 1"), ("Z2", "ZONE 2"), ("Z2X", "ADDITIONAL ZONE 2 (0.6 m)")]):
        fill, line, dash = STYLE[z]
        el = d.rect((xl + i * 70, 318), (12, 6), fill=fill, stroke=line, stroke_width=0.6)
        if dash:
            el["stroke-dasharray"] = dash
        d.add(el)
        p.text(t, xl + i * 70 + 15, 322.5, 2.3, bold=True)
    p.text("Section equipment: solid = cut, dashed grey = projected within +/- 12 m", xl + 215, 322.5, 2.1)
    table(p, xl, 340, [("EQUIPMENT", 48, "l"), ("ZONE 1", 108, "l"), ("ZONE 2", 150, "l")],
          [list(r) for r in PROTECTION], "EQUIPMENT PROTECTION BY ZONE (IEC 60079-14 / NFPA 70 Art. 505; EPL Gb / Gc)",
          size=2.1, rh=5.0)
    d.elements.remove(sh.g)          # texts / labels on top of the zone layers
    d.elements.append(sh.g)
    sh.save(out / "CFU-000-EL-HAC-002")
    return out / "CFU-000-EL-HAC-002.svg"


# ---------------------------------------------------------------------------------------------- HAC-003 xlsx
def hac_schedule(src, clr, path: Path):
    HDR = PatternFill("solid", fgColor="1F3864")
    thin = Side(style="thin", color="999999")
    BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "Release Sources"
    ws.cell(1, 1, "100 kBPSD Crude & Vacuum Distillation Unit - Hazardous Area Release Source Schedule").font = \
        Font(bold=True, size=13)
    ws.cell(2, 1, "CFU-000-EL-HAC-003  Rev A (Issued for review (FEED)), 2026-10-02  |  API RP 505 / EI 15 / "
                  "NFPA 70 Art. 505").font = Font(italic=True, size=9)
    hdr = ["Item", "Source ID", "Source type", "Description", "Equipment", "Fluid", "Fluid cat. (EI 15)",
           "Op. T (C)", "Op. P (barg)", "Grade of release", "Ventilation", "Zone(s)", "Horizontal extent (m)",
           "Vertical extent (EL m)", "Gas group", "Temp. class", "Plan location x / y (m)", "Basis", "Remarks"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(4, j, h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF", size=9), HDR, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, s in enumerate(src, 1):
        zones = ", ".join(("Zone 1" if e.zone == "Z1" else "Zone 2" if e.zone == "Z2" else "Zone 2 (additional)")
                          for e in s["envs"])
        hor = "; ".join(f"{'Z1' if e.zone == 'Z1' else 'Z2'}: {('R ' if e.kind == 'circle' else 'offset ')}{e.R:g}"
                        for e in s["envs"])
        ver = "; ".join(f"{e.z0:.1f}-{e.z1:.1f}" for e in s["envs"])
        row = [i, s["id"], s["kind"], s["desc"], s["equip"], s["fluid"], s["cat"], s["T"], s["P"], s["grade"],
               s["vent"], zones, hor, ver, s["group"], s["tclass"], f"{s['x']:.1f} / {s['y']:.1f}", s["basis"],
               s["remarks"]]
        for j, v in enumerate(row, 1):
            c = ws.cell(4 + i, j, v)
            c.border, c.font = BOX, Font(size=9)
            c.alignment = Alignment(wrap_text=j in (4, 6, 18, 19), vertical="top")
    for j, w in enumerate([5, 9, 20, 40, 14, 30, 8, 10, 10, 11, 18, 22, 20, 18, 7, 7, 13, 50, 40], 1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "C5"
    ws2 = wb.create_sheet("Unclassified items")
    for j, h in enumerate(["Item", "Status", "Clearance to nearest Zone 2 (m)", "Nearest source", "Description"], 1):
        c = ws2.cell(1, j, h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HDR
    for i, c in enumerate(clr, 2):
        st = UNCLASSIFIED.get(c["tag"], "")
        for j, v in enumerate([c["tag"], st, round(c["dist"], 1), c["src"], c["src_desc"]], 1):
            ws2.cell(i, j, v)
    for j, w in enumerate([12, 60, 16, 12, 60], 1):
        ws2.column_dimensions[chr(64 + j)].width = w
    ws3 = wb.create_sheet("Fluid categories")
    for i, (k, v) in enumerate(CAT_TXT.items(), 1):
        ws3.cell(i, 1, k)
        ws3.cell(i, 2, v)
    ws3.column_dimensions["B"].width = 120
    ws4 = wb.create_sheet("Equipment protection")
    for j, h in enumerate(["Equipment", "Zone 1", "Zone 2"], 1):
        c = ws4.cell(1, j, h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HDR
    for i, r in enumerate(PROTECTION, 2):
        for j, v in enumerate(r, 1):
            ws4.cell(i, j, v)
    for j, w in enumerate([28, 60, 80], 1):
        ws4.column_dimensions[chr(64 + j)].width = w
    wb.save(path)


# ---------------------------------------------------------------------------------------------- md section
def hac_md(src, clr):
    L = ["## 12. Hazardous area classification and equipment protection\n",
         "Classification CFU-000-EL-HAC-001/002 (schedule HAC-003) per API RP 505 / NFPA 70 Art. 505, point-source "
         "method, fluid categories per EI 15. Gas group IIA T3 generally, IIB in H2S / fuel-gas / off-gas services. "
         f"{len(src)} release sources: "
         + ", ".join(f"{k} {sum(1 for s in src if s['kind'] == k)}" for k in sorted({s['kind'] for s in src})) + ".\n",
         "| Unclassified item | Clearance to nearest classified area (m) | Nearest source |\n|---|---|---|"]
    for c in clr:
        L.append(f"| {c['tag']} | {c['dist']:.1f} | {c['src']} {c['src_desc'][:50]} |")
    L.append("\n| Equipment | Zone 1 | Zone 2 |\n|---|---|---|")
    for r in PROTECTION:
        L.append(f"| {r[0]} | {r[1]} | {r[2]} |")
    L.append("")
    return "\n".join(L)


# ---------------------------------------------------------------------------------------------- driver
def build(ctx, out: Path):
    L = json.loads((C.DATA / "layout.json").read_text())
    eq = json.loads((C.DATA / "equipment.json").read_text())
    loops = json.loads((C.DATA / "control_loops.json").read_text())["loops"]
    src = build_sources(L, eq, loops)
    clr = clearances(src, L)
    out.mkdir(parents=True, exist_ok=True)
    files = [hac_plan(ctx, L, src, clr, out), hac_sections(ctx, L, src, out)]
    hac_schedule(src, clr, out.parent / "CFU-000-EL-HAC-003_Release-Source-Schedule.xlsx")
    # area classification at each motor / consumer location (for electrical.json)
    loc = ctx["loc"]
    items = {e["tag"]: e for e in L["equipment"]}
    area = {}
    for r in ctx["rows"]:
        if r["kind"] != "motor" and not r["tag"].startswith("DT-"):
            continue
        (x, y), srcp = loc.point(r["tag"] if r["kind"] == "motor" else r["eq"])
        z = GRADE_EL + 1.0
        if r["group"] == "Air-cooler fans":
            z = 113.0
        if r["tag"].startswith("DT-"):
            it = items.get(r["eq"])
            z = (it["z_base"] + it["height"] + 0.5) if it else 106.0
        zn, grp = zone_at(src, x, y, z)
        area[r["tag"]] = dict(zone=zn, gas_group=grp if zn != "Unclassified" else "-",
                              temp_class="T3" if zn != "Unclassified" else "-",
                              protection=("Ex ec (Gc)" if zn == "Zone 2" else "Ex db / eb (Gb)" if zn == "Zone 1"
                                          else "Industrial (unclassified)"))
    return dict(src=src, clr=clr, files=files, area=area, md=hac_md(src, clr))
