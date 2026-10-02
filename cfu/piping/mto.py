"""Piping MTO CFU-000-PI-MTO-001 (xlsx) + valve list used by the study report."""
from __future__ import annotations

import re
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import specs
from .iso import MATS
from .model import key_of

ALLOW = {"critical": 0.05, "header": 0.10, "study": 0.10}
CV_PREFIX = ("FV", "LV", "PV", "TV", "HV", "PDV", "XV", "FFV")


def _sch(r, nps):
    w = specs.wall_calc(nps, r.line["design_P_barg"], r.line["design_T_C"], r.cls)
    if nps == r.nps:
        return getattr(r, "_sch_override", {}).get(nps, w["sch"])
    return w["sch"]


def aggregate(R):
    pipe, fit, flg, joints, welds, sup = OrderedDict(), OrderedDict(), OrderedDict(), OrderedDict(), [], []

    def add(d, k, q):
        d[k] = d.get(k, 0) + q

    for r in R.routes:
        if r.line_no in R.props:
            r._sch_override = {r.nps: R.props[r.line_no]["sch"]}
        rating = specs.cls(r.cls)["rating"]
        for bid, els in r.elements.items():
            for e in els:
                k, n = e["kind"], e.get("nps")
                if k == "pipe":
                    add(pipe, (r.cls, n, _sch(r, n), r.level), e["d1"] - e["d0"])
                elif k == "elbow":
                    add(fit, (r.cls, "ELBOW 45 LR" if e.get("angle", 90) < 60 else ("ELBOW 90 SR" if e.get("SR") else "ELBOW 90 LR"),
                              n, None), 1)
                elif k == "tee":
                    nb = e.get("nps_b")
                    add(fit, (r.cls, "TEE EQUAL" if nb == n else "TEE REDUCING", n, nb if nb != n else None), 1)
                elif k == "reducer":
                    it = e["item"]
                    add(fit, (r.cls, "REDUCER ECC" if it.get("ecc") else "REDUCER CONC", max(it["nps"], it["nps2"]),
                              min(it["nps"], it["nps2"])), 1)
                elif k == "cap":
                    add(fit, (r.cls, "CAP", n, None), 1)
                elif k == "flange":
                    add(flg, (r.cls, rating, "WN RF" + (" ORIFICE" if e.get("orifice") else ""), n), 1)
            for o in getattr(r, "_olets", {}).get(bid, []):
                nb = o.get("nps_b", 1)
                add(fit, (r.cls, "WELDOLET" if nb >= 2 else "SOCKOLET 3000#", nb, None), 1)
        for j in r.joints:
            add(joints, (r.cls, rating, j["nps"]), 1)
        welds.append(dict(line=r.line_no, cls=r.cls, level=r.level,
                          shop=sum(1 for w in r.welds if w["type"] == "S"),
                          field=sum(1 for w in r.welds if w["type"] == "F"),
                          idia_shop=round(sum(w["nps"] for w in r.welds if w["type"] == "S"), 1),
                          idia_field=round(sum(w["nps"] for w in r.welds if w["type"] == "F"), 1),
                          spools=len(r.spools)))
        cnt = {}
        for s in r.supports:
            cnt[s["type"]] = cnt.get(s["type"], 0) + 1
        sup.append(dict(line=r.line_no, level=r.level, **{k: cnt.get(k, 0) for k in "RGAS"}))
    return pipe, fit, flg, joints, welds, sup


def instrument_valves(R, line):
    return [i["tag"] for i in R.P.instr if i.get("line_no") == line["line_no"] and i["tag"].startswith(CV_PREFIX)
            and "valve" in i.get("type", "").lower()]


def valve_list(R):
    """(source, class, rating, kind, nps, qty) - isos exact; study/unrouted lines estimated from P&ID line list."""
    out = OrderedDict()

    def add(src, l, kind, nps, q):
        k = (src, l["cls"], specs.cls(l["cls"])["rating"], kind, nps)
        out[k] = out.get(k, 0) + q

    routed = {r.line_no: r for r in R.routes}
    for r in R.routes:
        l = r.line
        if r.level == "critical":
            for bid, els in r.elements.items():
                for e in els:
                    if e["kind"] in ("gate", "globe", "check", "cv", "strainer", "spec"):
                        add("ISO", l, {"cv": "control", "spec": "spectacle blind", "strainer": "strainer (temp.)"}.get(e["kind"], e["kind"]),
                            e.get("nps") or r.nps, 1)
            for o in [o for v in getattr(r, "_olets", {}).values() for o in v]:
                tg = (o.get("item") or {}).get("tag") or ""
                if tg in ("HPV", "LPD"):
                    add("ISO", l, "gate (vent/drain)", o.get("nps_b", 1), 2 if l["cls"] in ("B2", "B3", "C1", "S2") else 1)
            continue
    for l in R.P.lines:
        if l["line_no"] in routed and routed[l["line_no"]].level == "critical":
            continue
        src = "ROUTED (study)" if l["line_no"] in routed else "P&ID estimate"
        n = l["size_in"]
        f, t = l["from"], l["to"]
        r = routed.get(l["line_no"])
        sv = getattr(r, "study_valves", None) if r else None
        if sv:
            for kind, nps, q in sv:
                add(src, l, {"strainer": "strainer (temp.)"}.get(kind, kind), nps, q)
        else:
            if re.match(r"^P-\d{3}", t):
                add(src, l, "gate", n, 2)
                add(src, l, "strainer (temp.)", n, 2)
            if re.match(r"^P-\d{3}", f):
                add(src, l, "check", n, 2)
                add(src, l, "gate", n, 2)
            if re.match(r"^C-\d{3} (tray|bottom|.*pan|boot)", f):
                add(src, l, "gate", n, 1)
            if "OSBL" in f + t or f.startswith("TK"):
                add(src, l, "gate", n, 1)
                add(src, l, "spectacle blind", n, 1)
            if "header" in f.lower():
                add(src, l, "gate", n, 1)
            if f.startswith("PSV"):
                add(src, l, "gate (PSV outlet, CSO)", n, 1)
        for tag in instrument_valves(R, l):
            cvn = specs.step(n, -1) if n > 2 else n
            add(src, l, "control" if not tag.startswith("XV") else "on/off (SDV)", cvn, 1)
            if not tag.startswith("XV"):
                add(src, l, "gate", n, 2)
                add(src, l, "globe", specs.step(n, -1) if n > 2 else n, 1)
        if n >= 2:
            add(src, l, "gate (vent/drain)", 0.75 if l["cls"] not in ("B2", "B3") else 1, 2)
    return out


def line_crosscheck(R):
    routed = {r.line_no: r for r in R.routes}
    in_manifold = {}
    for r in R.routes:
        for b in r.branches:
            if getattr(b, "line_ref", None):
                in_manifold[b.line_ref] = r
    unr = dict(R.unrouted)
    all_nos = {l["line_no"] for l in R.P.lines}
    rows = []
    for l in R.P.lines:
        ln = l["line_no"]
        if ln in routed:
            r = routed[ln]
            rows.append([ln, l["cls"], l["size_in"], f"{l['from']} -> {l['to']}", "ROUTED", r.level,
                         getattr(r, "iso", "") or "", round(r.total_length(), 1), f"Cat {r.cat[0]}", ""])
        elif ln in in_manifold:
            r = in_manifold[ln]
            rows.append([ln, l["cls"], l["size_in"], f"{l['from']} -> {l['to']}", "ROUTED (manifold branch)",
                         "critical", getattr(r, "iso", ""), "", "", f"part of {r.line_no}"])
        else:
            rows.append([ln, l["cls"], l["size_in"], f"{l['from']} -> {l['to']}", "NOT ROUTED", "", "", "", "",
                         unr.get(ln, "connection shown as tie-in on related iso / detailed design")])
    orphans = [r.line_no for r in R.routes if r.line_no not in all_nos]
    return rows, orphans


def _sheet(wb, title, hdr, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths or [14] * len(hdr)):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"
    return ws


def write(R, path):
    pipe, fit, flg, joints, welds, sup = aggregate(R)
    vl = valve_list(R)
    xrows, orphans = line_crosscheck(R)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    tot_pipe = sum(pipe.values())
    tot_pipe_allow = sum(v * (1 + ALLOW[k[3]]) for k, v in pipe.items())
    n_fit = sum(fit.values())
    n_flg = sum(flg.values())
    n_v = sum(q for k, q in vl.items() if k[0] in ("ISO", "ROUTED (study)") and "vent" not in k[3])
    n_v_all = sum(vl.values())
    idia_s = sum(w["idia_shop"] for w in welds)
    idia_f = sum(w["idia_field"] for w in welds)
    R.mto_totals = dict(pipe_m=round(tot_pipe, 1), pipe_m_allow=round(tot_pipe_allow, 1), fittings=int(n_fit),
                        flanges=int(n_flg), joints=int(sum(joints.values())), valves_routed=int(n_v),
                        valves_all=int(n_v_all), weld_idia_shop=round(idia_s), weld_idia_field=round(idia_f),
                        welds=int(sum(w["shop"] + w["field"] for w in welds)), spools=int(sum(w["spools"] for w in welds)),
                        supports=int(sum(s["R"] + s["G"] + s["A"] + s["S"] for s in sup)),
                        routed_lines=len(R.routes), lines=len(R.P.lines))
    meta = [("Document", "CFU-000-PI-MTO-001 Piping Material Take-Off"), ("Revision", "A - Issued for review (FEED)"),
            ("Date", "2026-10-02"), ("Basis", "Routed 3D model data/routing.json (critical = isometric level, "
                                              "study/header = auto-routed FEED level)"),
            ("Allowances", "Pipe: +5 % critical (iso), +10 % study/header; fittings/valves net"),
            ("", "")]
    for a, b in meta:
        ws.append([a, b])
    ws.append(["ITEM", "QUANTITY", "UNIT"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for a, b, c in [("Routed lines", len(R.routes), "no."), ("Lines in line list", len(R.P.lines), "no."),
                    ("Pipe (net)", round(tot_pipe, 1), "m"), ("Pipe incl. allowance", round(tot_pipe_allow, 1), "m"),
                    ("Fittings (BW + olets)", n_fit, "pcs"), ("Flanges", n_flg, "pcs"),
                    ("Bolted joints (gasket + bolt set)", sum(joints.values()), "sets"),
                    ("Valves on routed lines (excl. vents/drains)", n_v, "pcs"),
                    ("Valves incl. P&ID estimate for unrouted lines + vents/drains", n_v_all, "pcs"),
                    ("Butt welds - shop", sum(w["shop"] for w in welds), "no."), ("Butt welds - field", sum(w["field"] for w in welds), "no."),
                    ("Weld inch-dia - shop", round(idia_s), "in-dia"), ("Weld inch-dia - field", round(idia_f), "in-dia"),
                    ("Spools", sum(w["spools"] for w in welds), "no."),
                    ("Pipe supports", R.mto_totals["supports"], "no.")]:
        ws.append([a, b, c])
    ws.append([])
    ws.append(["Pipe by class (m incl. allowance)"])
    bycls = {}
    for k, v in pipe.items():
        bycls[k[0]] = bycls.get(k[0], 0) + v * (1 + ALLOW[k[3]])
    for c_, v in sorted(bycls.items()):
        ws.append([c_, round(v, 1), "m"])
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 60
    # pipe
    agg = OrderedDict()
    for (c_, n, sch, lev), v in pipe.items():
        k = (c_, n, sch)
        a = agg.setdefault(k, [0.0, 0.0])
        a[0] += v
        a[1] += v * (1 + ALLOW[lev])
    rows = []
    for (c_, n, sch), (net, gross) in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        m = MATS[specs.cls(c_)["mat"]]
        rows.append([c_, specs.nps_str(n), sch, m["pipe"] if n <= 24 else m["big"], round(net, 1), round(gross, 1)])
    _sheet(wb, "Pipe", ["Class", "NPS", "Schedule / WT", "Material", "Net length m", "Length incl. allowance m"], rows,
           [8, 8, 14, 34, 14, 22])
    rows = [[c_, t, specs.nps_str(n), specs.nps_str(n2) if n2 else "", MATS[specs.cls(c_)["mat"]]["fit"], q]
            for (c_, t, n, n2), q in sorted(fit.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2]))]
    _sheet(wb, "Fittings", ["Class", "Type", "NPS 1", "NPS 2", "Material", "Qty"], rows, [8, 18, 8, 8, 20, 8])
    rows = [[c_, f"CL{rt}", t, specs.nps_str(n), MATS[specs.cls(c_)["mat"]]["flg"], q]
            for (c_, rt, t, n), q in sorted(flg.items(), key=lambda kv: (kv[0][0], kv[0][3]))]
    rows += [[c_, f"CL{rt}", "GASKET SPW + BOLT SET", specs.nps_str(n), MATS[specs.cls(c_)["mat"]]["bolt"], q]
             for (c_, rt, n), q in sorted(joints.items(), key=lambda kv: (kv[0][0], kv[0][2]))]
    _sheet(wb, "Flanges-Gaskets-Bolts", ["Class", "Rating", "Item", "NPS", "Material", "Qty"], rows, [8, 8, 24, 8, 22, 8])
    rows = [[src, c_, f"CL{rt}", kind, specs.nps_str(n), MATS[specs.cls(c_)["mat"]]["valve"], q]
            for (src, c_, rt, kind, n), q in vl.items()]
    _sheet(wb, "Valves", ["Source", "Class", "Rating", "Type", "NPS", "Body / trim", "Qty"], rows,
           [16, 8, 8, 22, 8, 20, 8])
    rows = [[w["line"], w["cls"], w["level"], w["shop"], w["field"], w["idia_shop"], w["idia_field"], w["spools"]]
            for w in welds]
    _sheet(wb, "Welds-Spools", ["Line", "Class", "Level", "Shop welds", "Field welds", "Inch-dia shop",
                                "Inch-dia field", "Spools"], rows, [26, 8, 10, 10, 10, 12, 12, 8])
    rows = [[s["line"], s["level"], s["R"], s["G"], s["A"], s["S"], s["R"] + s["G"] + s["A"] + s["S"]] for s in sup]
    _sheet(wb, "Supports", ["Line", "Level", "Rest", "Guide", "Anchor", "Spring", "Total"], rows,
           [26, 10, 8, 8, 8, 8, 8])
    ws2 = _sheet(wb, "Line-List-Crosscheck", ["Line no. (lines.json)", "Class", "NPS", "From -> To", "Status", "Level",
                                              "Isometric", "Routed m", "Stress", "Remark"], xrows,
                 [26, 7, 6, 46, 22, 9, 20, 9, 8, 60])
    ws2.append([])
    ws2.append(["Routed lines not in lines.json:", ", ".join(orphans) or "none"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return R.mto_totals
