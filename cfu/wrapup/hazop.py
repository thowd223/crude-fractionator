"""Preliminary HAZOP: node list and key deviations / safeguards (CFU-000-PR-RPT-002)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .. import docgen

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "deliverables" / "06-wrapup"

NODES = [
    ("N01", "Crude charge P-101A/B and cold preheat E-101..E-105", "PID-001", "130 C, 22 barg"),
    ("N02", "Desalters D-101A/B incl. wash water and brine", "PID-002", "136 C, 11.5 barg"),
    ("N03", "Hot preheat E-106..E-111 and P-102A/B", "PID-003", "274 C, 20 barg"),
    ("N04", "Atmospheric heater H-101 incl. fuel gas, APH, fans", "PID-004/005", "366 C COT"),
    ("N05", "C-101 flash zone, wash zone and stripping section, P-112", "PID-006", "361 C, 1.5 barg"),
    ("N06", "C-101 pumparounds TPA/MPA/BPA and side strippers C-102/103/104", "PID-007/008", "150-310 C"),
    ("N07", "C-101 overhead A-101, E-115, D-102, P-103/104/105", "PID-009", "45-135 C, 0.7-1.2 barg"),
    ("N08", "Stabiliser C-105, E-114, E-116, A-106, D-105, P-115", "PID-010", "11.5 barg, LPG"),
    ("N09", "Naphtha splitter C-106, E-117, A-107, A-108, D-106", "PID-011", "0.8 barg"),
    ("N10", "Vacuum heater H-201", "PID-012", "398 C COT"),
    ("N11", "Vacuum column C-201 incl. wash bed, quench, P-204", "PID-013/014", "60 mbar(a), 386 C"),
    ("N12", "Ejectors J-201..203, condensers E-202..204, D-201, D-202", "PID-015", "20-1100 mbar(a)"),
    ("N13", "Fuel gas KO D-103, unit flare KO D-104, steam headers", "PID-016", "Utilities"),
]

# node, deviation, cause, consequence, safeguards, recommendation
ROWS = [
    ("N01", "No flow", "P-101 trip, FV-1001 fails closed, tank valve closed",
     "Loss of crude to H-101 passes -> tube overheating / coking, potential tube rupture and fire",
     "FAL-1001; SIF-101 low-low pass flow trips H-101 fuel (SIL 2); P-101B auto-start", "Verify SIF-101 response time vs. tube metal heat-up (LOPA)"),
    ("N01", "More pressure", "Blocked outlet downstream of P-101", "Overpressure of cold train shells", "Pump shut-off < design P of cold train (check); PSV-1010 thermal", "Confirm shut-off head of P-101 vs. 30 barg exchanger design"),
    ("N02", "Less level (interface)", "LV-1007 fails open", "Oil carry-under to brine -> WWT upset; low interface shorts grids",
     "LAL-1007; SIF-107 trips desalter transformers; brine oil analyser", "Provide brine oil-in-water analyser with alarm"),
    ("N02", "More temperature", "TV-1004 bypass fails closed / high crude T", "Crude vaporises in desalter, grid arcing, pressure rise", "TAH-1004; PIC-1009 backpressure; PSV-1002/1003", "Set desalter P >= crude bubble P + 1.5 bar at max T"),
    ("N03", "Leak", "Tube/flange leak on hot exchangers (> auto-ignition)", "Fire in hot train (crude 274 C above AIT)",
     "5Cr / 9Cr materials, fireproofing, deluge on E-111 bank, ROSOV on P-102 suction", "Hot flange management programme; flange guards"),
    ("N04", "Less flow (one pass)", "Pass FV fails closed, coking", "Pass tube overheating and rupture -> firebox fire",
     "FIC-1011..1018 + TDIC-1019; SIF-101 low-low pass flow per pass 2oo3; TI tube skins", "Tube-skin TCs on each pass outlet with alarms"),
    ("N04", "More fuel", "PV-1021 fails open, PIC fault", "Overfiring, high COT, coking, flame impingement", "SIF-103 high-high FG pressure, SIF-106 high-high COT; TAH-1020", ""),
    ("N04", "Loss of flame", "Low FG pressure, liquid carry-over in FG", "Fuel accumulation, firebox explosion on relight",
     "D-103 FG KO with LAHH; SIF-102/104 (BMS per NFPA 85/API 556); purge interlock", "BMS to API 556 with 5-volume purge"),
    ("N04", "Less draft", "ID fan K-102 trip", "Positive firebox pressure, flue gas leakage, burner instability", "SIF-105; auto-open stack damper / natural-draft fallback; K-102B", "Confirm natural-draft operating capacity (~70 %)"),
    ("N05", "More level", "Loss of P-112 / LV-1082 fails closed", "Liquid into flash zone -> slugging, tray damage, overflash loss",
     "LAH-1082, SIF-108 (LAHH) stops crude charge", ""),
    ("N05", "More water", "Desalter upset, water in crude", "Rapid vaporisation in column -> pressure surge, tray damage",
     "Desalter interface control; PAH C-101; PSV-1001", "Operating procedure for water-slug response"),
    ("N06", "No flow (PA)", "P-106/107/108 trip", "Loss of heat removal -> column top T and P rise, overhead overload, relief",
     "Spare pumps auto-start; TAH top; PSV-1001 (reflux-failure case)", "Include total PA failure in flare load study"),
    ("N07", "Less cooling", "A-101 fan failure, CW failure on E-115", "Column pressure rise, relief to flare",
     "PAH-1032; 4 fans on 2 buses; PSV-1001 sized for reflux failure", ""),
    ("N07", "Corrosion", "HCl / NH4Cl salt deposition, low water dew-point margin", "Overhead leaks, loss of containment (H2S)",
     "Neutraliser + filming amine (X-103), pH AIC-1037, wash water, Monel top, Ti E-115 tubes, 14 C dew-point margin", "Corrosion monitoring: ER probes + chloride analyser on boot water"),
    ("N08", "More pressure", "Loss of A-106, reboiler overheat, blocked LPG", "Overpressure of C-105 (LPG)", "PIC-1091; SIF-110 cuts HP steam; PSV-1005/1006", ""),
    ("N08", "Leak", "LPG pump seal failure", "Flammable cloud, VCE potential", "Dual seals Plan 53B, gas detection, ROSOV on D-105 outlet, 30 m spacing to heaters", "Consequence modelling of LPG release (QRA input)"),
    ("N08", "Less level", "LV/FV on LPG product fails open, LIC-1092 fault", "Gas blow-by into LPG treating, overpressure downstream",
     "LALL-1092; SIF-109 closes XV-1093 (SIL 1)", ""),
    ("N09", "Less temperature (reboiler)", "MP steam loss", "Off-spec LN/HN, no safety consequence", "TAL-1104; quality alarms", ""),
    ("N10", "Less flow / coking", "Loss of coil steam, low pass flow", "Tube coking, hot spots, rupture", "FIC-2007 coil steam, SIF-201", "Online spalling / pigging provision"),
    ("N11", "Air ingress", "Flange leak under vacuum", "Internal fire / explosion in column, loss of vacuum", "O2 analyser on off-gas, leak testing, tightness class, hot-oil temperature limits", "Add O2 analyser AI-2032 on J-201 suction"),
    ("N11", "Less flow (wash oil)", "FV-2020 fails closed", "Wash-bed coking -> HVGO metals/CCR off-spec, bed plugging", "FAL-2020 (critical alarm), min-stop on FV-2020", ""),
    ("N11", "More temperature (bottoms)", "Loss of quench", "VR coking in boot, P-204 damage", "TIC-2026 / FIC-2027, TAH-2026", ""),
    ("N12", "Loss of vacuum", "Motive steam failure, CW failure", "Column pressure rise, off-spec, possible relief", "PAH-2010, alarm on motive steam P", "Vacuum-loss emergency procedure (reduce H-201 firing)"),
    ("N12", "H2S release", "Hotwell vent / off-gas line", "Toxic exposure", "Off-gas routed to H-201 burners with KO D-202; H2S detectors", "Backup off-gas route to flare on H-201 trip"),
    ("N13", "Liquid in fuel gas", "Upstream carry-over", "Burner flame-out / flooding", "D-103 LAHH -> SIF-102", ""),
]


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Nodes"
    ws.append(["Node", "Description", "P&ID", "Design intent"])
    for r in NODES:
        ws.append(list(r))
    ws2 = wb.create_sheet("Worksheet")
    ws2.append(["Node", "Deviation", "Causes", "Consequences", "Safeguards", "Recommendations"])
    for r in ROWS:
        ws2.append(list(r))
    for w in (ws, ws2):
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
        for col, wd in zip("ABCDEF", (7, 22, 40, 45, 50, 40)):
            w.column_dimensions[col].width = wd
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT / "CFU-000-PR-RPT-002_Preliminary-HAZOP.xlsx")

    md = """# 1 Scope and method
This is a preliminary (FEED-stage) HAZOP of the CDU/VDU P&IDs. It uses guideword-deviation analysis
(IEC 61882) on 13 nodes. It identifies the main hazards, checks that the proposed safeguards (BPCS, SIS,
relief and mechanical) are adequate, and feeds the SIL assessment (LOPA). It does not replace the formal
multidisciplinary HAZOP of the issued-for-design P&IDs.

# 2 Nodes
| Node | Description | P&ID | Design intent |
|---|---|---|---|
""" + "".join(f"| {a} | {b} | {c} | {d} |\n" for a, b, c, d in NODES) + """
# 3 Worksheet (principal deviations)
| Node | Deviation | Causes | Consequences | Safeguards | Recommendation |
|---|---|---|---|---|---|
""" + "".join(f"| {' | '.join(r)} |\n" for r in ROWS) + """
# 4 Key findings
- The fired heaters dominate the risk profile. SIF-101 and SIF-201 (low-low pass flow, SIL 2) and a BMS to
  API 556 are essential.
- The hot train and hot pumps handle crude above its auto-ignition temperature. Hot pumps P-102, P-108, P-112 and
  P-204 get remotely operated shut-off valves (SIF-204), fireproofing to API 2218 and deluge.
- Overhead corrosion (HCl/NH4Cl) is controlled by desalting, caustic, neutraliser, wash water, Monel and Ti
  materials, and corrosion monitoring.
- The LPG section is kept at least 30 m from the fired heaters (plot plan), with gas detection and ROSOVs.
- The governing unit relief case (C-101 reflux/pumparound failure, about 220 t/h) needs a global flare-load study
  covering power and cooling-water failure.
"""
    docgen.render(md, OUT / "CFU-000-PR-RPT-002_Preliminary-HAZOP", "CFU-000-PR-RPT-002", "Preliminary HAZOP Study")
