"""Phase maturity plan (CFU-000-PM-DEP-003 report + workbook, portal/phases.html).

Extends the design dependency map (activities.py) across the front-end-loading phases. The same activities run
in every phase; what changes is how mature each one has to be at each gate. So the plan has three parts:

- a maturity scale (LEVELS) and a target matrix (TARGET): the level each activity must reach at each gate;
- the activities that only exist in some phases (PHASE_ACTIVITIES: business case, select decision, execution
  planning, procurement, construction-facing detailed design);
- the maturity each link needs. By default, taking an activity to level L needs each input at level L. CONSUME
  caps that for activities that only need early inputs (estimates, HAZOPs, studies), LINK_REQ overrides single
  links, and LOOP_REQ turns each iteration loop into a forward link between levels (P&IDs reach "defined" only
  after the preliminary HAZOP), so the network of (activity, level) steps has no cycles.

The plan is checked for consistency (no gate asks for an activity at a level its inputs cannot support at that
gate) and compared with what this package has issued (achieved levels) to give gate readiness and the
ready / blocked list for every phase.
"""
from __future__ import annotations

import html
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import basis, docgen
from .activities import A, ACTIVITIES, FEEDBACK

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "deliverables" / "06-wrapup"
PORTAL = OUT / "portal"
DOCNO, TITLE = "CFU-000-PM-DEP-003", "Phase Maturity Plan"

DISCIPLINES = ["Project", "Process", "Safety", "Mechanical", "Procurement", "Layout", "Piping", "Civil", "I&C",
               "Electrical", "Construction", "Cost"]

PHASES = [
    dict(id="FEL1", name="FEL 1 · Appraise", gate="Gate 1 · opportunity confirmed", est="Class 5",
         aim="Confirm there is a business case and a technically feasible concept worth developing."),
    dict(id="FEL2", name="FEL 2 · Select", gate="Gate 2 · concept selected", est="Class 4",
         aim="Compare the alternatives and select one configuration, technology and site to take forward."),
    dict(id="FEL3", name="FEL 3 · Define (FEED)", gate="Gate 3 · final investment decision", est="Class 3",
         aim="Define the selected concept well enough to fund it: HAZOP'd P&IDs, datasheets, long-lead orders."),
    dict(id="DD", name="Detailed design", gate="Gate 4 · issued for construction", est="Class 2",
         aim="Complete the design with certified vendor data and issue everything needed to build."),
]
PH = [p["id"] for p in PHASES]

LEVELS = {
    0: "Not started",
    1: "Concept",            # sketches, factored data, assumption lists
    2: "Preliminary",        # issued for review (IFR) / rev A; enough to estimate and screen
    3: "Defined",            # issued for design / for HAZOP / for purchase; frozen basis for the next phase
    4: "Final",              # IFC, approved or certified; one-off studies and decisions when complete
}

# Activities that do not appear in the FEED dependency map: front-end decisions, execution planning,
# procurement and the detailed-design work that turns FEED into construction drawings.
PHASE_ACTIVITIES = [
    # ---------------------------------------------------------------- FEL 1
    A("FEL-01", "Business case, capacity and product slate", "Project", "assumed",
      ext=["Market and margin study", "Owner corporate strategy"],
      note="Taken as given in this package: the BOD states capacity, crude and products."),
    A("FEL-02", "Configuration options and screening", "Process", "assumed",
      after=[("FEL-01", "capacity, products"), ("PRJ-01", "design basis")],
      note="Taken as given: a conventional CDU/VDU configuration was assumed, no options were screened."),
    A("FEL-03", "Site screening and selection", "Layout", "assumed",
      after=[("FEL-01", "location constraints")], ext=["Candidate sites, land, logistics, permitting risk"],
      note="Taken as given: generic US Gulf Coast site in the BOD."),
    A("FEL-04", "Class 5 estimate and economics", "Cost", "assumed",
      after=[("FEL-01", "capacity, margins"), ("FEL-02", "options"), ("PRC-02", "major equipment list"),
             ("LAY-01", "outline plot area")],
      note="Superseded in this package by the Class 4 estimate (CST-01)."),
    # ---------------------------------------------------------------- FEL 2
    A("SEL-02", "Technology and licensor selection", "Process", "assumed",
      after=[("FEL-02", "screened options")], ext=["Licensor proposals"],
      note="Taken as given: open-art distillation, no licensor."),
    A("SEL-01", "Alternatives evaluation and select decision", "Process", "assumed",
      after=[("FEL-02", "options"), ("SEL-02", "selected technology"), ("PRC-01", "case balances"),
             ("FEL-03", "selected site")],
      note="Taken as given: the package develops a single case."),
    A("SEL-03", "Utility and offsite basis", "Process", "partial",
      after=[("PRJ-01", "utility conditions"), ("PRC-01", "utility demands")],
      note="The BOD gives battery-limit utility conditions only; no utility or offsite balances."),
    A("SEL-04", "Permitting and environmental basis", "Project", "open",
      after=[("PRC-01", "emissions and effluents"), ("LAY-01", "site layout")],
      ext=["Air permit, wastewater and other regulatory requirements"]),
    A("PRJ-05", "Execution plan and contracting strategy", "Project", "open",
      after=[("FEL-01", "schedule driver"), ("SEL-01", "selected scope")],
      ext=["Owner execution and contracting preferences"]),
    A("CON-01", "Constructability, modularisation and path of construction", "Construction", "open",
      after=[("LAY-01", "plot plan"), ("LAY-02", "layout model"), ("PRJ-05", "execution strategy")],
      ext=["Construction contractor input, logistics and heavy-lift survey"]),
    # ---------------------------------------------------------------- FEL 3 / detailed design
    A("PRO-01", "Requisitions and technical bid evaluations", "Procurement", "open",
      after=[("MEC-02", "datasheets"), ("PRJ-05", "contracting strategy")],
      note="Long-lead items (heaters, columns, compressors, large pumps) are requisitioned in FEL 3."),
    A("CON-02", "Work packaging and system boundaries", "Construction", "open",
      after=[("CON-01", "path of construction"), ("PRC-08", "P&IDs"), ("PIP-04", "isometrics")],
      note="Engineering and construction work packages, test packs, commissioning systems."),
    A("CIV-02", "Foundation and structural steel detailing", "Civil", "open",
      after=[("CIV-01", "civil and structural design"), ("VEN-01", "certified loads and anchor bolts"),
             ("PIP-05", "support loads")]),
    A("ICS-09", "Instrument installation design", "I&C", "open",
      after=[("ICS-05", "I/O list"), ("PIP-02", "instrument positions"), ("VEN-01", "vendor instruments")],
      note="Location plans, hook-ups, JB and cabinet wiring."),
    A("ELE-06", "Electrical layouts", "Electrical", "open",
      after=[("ELE-04", "cable schedule"), ("LAY-02", "layout model"), ("PIP-02", "piping model")],
      note="Raceway, grounding, lighting and substation layouts."),
    A("CST-03", "Control estimate", "Cost", "open",
      after=[("CST-02", "Class 3 estimate"), ("PRO-01", "awarded prices"), ("PIP-04", "piping MTO"),
             ("CIV-01", "civil quantities")]),
]

# Extra links that only matter once the phases are modelled (from, to, what flows).
EXTRA_LINKS = [
    ("FEL-01", "PRJ-01", "business case"),
    ("FEL-03", "LAY-01", "selected site"),
    ("SEL-01", "PRC-01", "selected configuration"),
    ("SEL-01", "CST-01", "selected scope"),
    ("PRO-01", "VEN-01", "purchase orders"),
    ("CIV-02", "PRJ-04", "foundation and steel drawings"),
    ("ICS-09", "PRJ-04", "installation drawings"),
    ("ELE-06", "PRJ-04", "electrical layouts"),
    ("CON-02", "PRJ-04", "work packages"),
    ("PRO-01", "PRJ-04", "purchase orders"),
]

# Target level at each gate: (FEL1, FEL2, FEL3, DD).
TARGET = {
    "FEL-01": (3, 4, 4, 4), "FEL-02": (4, 4, 4, 4), "FEL-03": (2, 4, 4, 4), "FEL-04": (4, 4, 4, 4),
    "PRJ-01": (2, 3, 4, 4), "PRJ-02": (1, 2, 4, 4), "PRC-00": (0, 1, 4, 4), "PRC-01": (1, 2, 3, 4),
    "PRC-02": (1, 2, 3, 4), "PRC-03": (0, 1, 3, 4), "PRC-04": (0, 1, 3, 4), "PRC-05": (1, 3, 4, 4),
    "PRC-06": (1, 2, 3, 4), "PRC-07": (0, 1, 4, 4), "PIP-01": (0, 1, 3, 4), "PRC-08": (0, 1, 3, 4),
    "SEL-02": (0, 4, 4, 4), "SEL-01": (0, 4, 4, 4), "SEL-03": (0, 2, 3, 4), "SEL-04": (0, 2, 3, 4),
    "PRJ-05": (1, 2, 4, 4),
    "SAF-01": (0, 0, 4, 4), "SAF-02": (0, 1, 3, 4), "SAF-03": (0, 0, 0, 4), "SAF-04": (0, 0, 4, 4),
    "MEC-01": (0, 1, 3, 4), "MEC-02": (0, 1, 3, 4), "MEC-03": (0, 0, 2, 4), "VEN-01": (0, 0, 2, 4),
    "PRO-01": (0, 0, 3, 4),
    "LAY-01": (1, 2, 3, 4), "LAY-02": (0, 1, 3, 4), "LAY-03": (0, 0, 2, 4),
    "PIP-02": (0, 0, 2, 4), "PIP-03": (0, 0, 2, 4), "PIP-04": (0, 0, 2, 4), "PIP-05": (0, 0, 1, 4),
    "CIV-01": (0, 1, 2, 4), "CIV-02": (0, 0, 0, 4),
    "ICS-01": (0, 1, 3, 4), "ICS-02": (0, 0, 2, 4), "ICS-03": (0, 0, 2, 4), "ICS-04": (0, 0, 2, 4),
    "ICS-05": (0, 0, 2, 4), "ICS-06": (0, 0, 1, 4), "ICS-07": (0, 1, 3, 4), "ICS-08": (0, 0, 0, 4),
    "ICS-09": (0, 0, 0, 4),
    "ELE-01": (0, 1, 3, 4), "ELE-02": (0, 0, 3, 4), "ELE-03": (0, 1, 3, 4), "ELE-04": (0, 0, 2, 4),
    "ELE-05": (0, 0, 1, 4), "ELE-06": (0, 0, 0, 4),
    "CON-01": (0, 1, 3, 4), "CON-02": (0, 0, 1, 4),
    "CST-01": (0, 4, 4, 4), "CST-02": (0, 0, 4, 4), "CST-03": (0, 0, 0, 4),
    "PRJ-03": (1, 2, 3, 4), "PRJ-04": (0, 0, 0, 4),
}

# Highest input maturity an activity ever needs: an int caps every level, a dict caps single levels
# (0 = no input needed at that level).
CONSUME = {
    "FEL-02": 2, "FEL-03": 2, "FEL-04": 1, "SEL-01": 2, "SEL-02": 2,
    "PRC-05": 2, "PRC-07": 3, "SAF-01": 2, "SAF-03": 3, "SAF-04": 2, "VEN-01": 3,
    "CST-01": 2, "CST-02": 2, "CST-03": 3,
    "PRJ-03": {1: 0, 2: 0, 3: 0},   # the register lists whatever exists; it needs a complete set only at IFC
}

# Single-link overrides: (from, to) -> {level of 'to': level needed of 'from'} (0 = not needed).
LINK_REQ = {
    ("PRC-00", "PRC-01"): {1: 0, 2: 1},         # shortcut model until the rigorous simulation exists
    ("SEL-01", "PRC-01"): {1: 0, 2: 0},         # case balances come before the select decision
    ("SEL-01", "CST-01"): {4: 4},               # the Class 4 estimate prices the selected case
    ("SEL-02", "SEL-01"): {4: 4},
    ("SEL-01", "PRJ-05"): {1: 0},               # FEL 1 execution outline comes before the select decision
    ("PRC-04", "PRC-06"): {1: 0, 2: 1},         # principal loops only on early PFDs
    ("PIP-03", "CIV-01"): {1: 0},
    ("ICS-05", "ICS-07"): {1: 0, 3: 2},         # FEED architecture sized from preliminary I/O counts
    ("MEC-01", "CST-01"): {4: 1},               # factored weights are enough for Class 4
}

# Iteration loops as forward links between levels: (from, to) -> (level of 'to', level needed of 'from').
LOOP_REQ = {
    ("MEC-01", "PRC-02"): (3, 2),
    ("MEC-01", "LAY-01"): (3, 2),
    ("PIP-02", "PRC-02"): (4, 3),
    ("PIP-02", "LAY-01"): (4, 3),
    ("ICS-04", "PRC-08"): (4, 3),
    ("SAF-01", "PRC-08"): (3, 4),               # P&IDs are "defined" once the preliminary HAZOP is closed
    ("ELE-04", "ELE-03"): (4, 3),
    ("ICS-07", "ELE-01"): (4, 3),
    ("SAF-02", "PRC-03"): (4, 3),
    ("SAF-03", "PRC-08"): (4, 4),               # IFC P&IDs carry the closed formal HAZOP actions
    ("SAF-04", "LAY-01"): (3, 4),               # the plot plan is frozen after the siting study
    ("VEN-01", "MEC-01"): (4, 4),
    ("VEN-01", "LAY-01"): (4, 4),
    ("VEN-01", "ELE-01"): (4, 4),
    ("VEN-01", "ICS-05"): (4, 4),
}

STATUS = {"done": "Issued", "partial": "Issued, provisional", "open": "Not started",
          "assumed": "Assumed by the design basis"}


# ---------------------------------------------------------------------------------------------- model
class Plan:
    def __init__(self):
        self.acts = ACTIVITIES + PHASE_ACTIVITIES
        self.A = {a["id"]: a for a in self.acts}
        missing = sorted(set(self.A) ^ set(TARGET))
        if missing:
            raise ValueError(f"TARGET and the activity list differ: {missing}")
        if len(self.A) != len(self.acts):
            raise ValueError("duplicate activity id")
        self.links = [(p, a["id"], w) for a in self.acts for p, w in a["after"]] + EXTRA_LINKS
        bad = [(p, i) for p, i, _ in self.links if p not in self.A or i not in self.A]
        bad += [k for k in list(LINK_REQ) + list(LOOP_REQ) if k[0] not in self.A or k[1] not in self.A]
        bad += [(f, t) for f, t, _ in FEEDBACK if (f, t) not in LOOP_REQ]
        if bad:
            raise ValueError(f"unknown or unmapped links: {bad}")
        self.what = {(p, i): w for p, i, w in self.links}
        self.what.update({(f, t): w for f, t, w in FEEDBACK})
        self.top = {i: max(TARGET[i]) for i in self.A}
        self.achieved = {i: self._achieved(i) for i in self.A}
        self.req = self._requirements()
        self.phase = {(i, L): self._phase_of(i, L) for i in self.A for L in range(1, self.top[i] + 1)}
        self.depth = self._depths()
        self.issues = self._consistency()

    def _achieved(self, i):
        a, t = self.A[i], TARGET[i]
        st = a["status"]
        if st in ("done", "assumed"):
            return t[2] or max(t[:2])
        if st == "partial":
            return max(t[1], 1)
        return 0

    def need(self, src, dst, L):
        """Level of `src` that `dst` needs to reach level L (0 = none)."""
        if (src, dst) in LOOP_REQ:
            lv, need = LOOP_REQ[(src, dst)]
            return need if L >= lv else 0
        cap = CONSUME.get(dst, L)
        r = min(L, cap.get(L, L) if isinstance(cap, dict) else cap)
        o = LINK_REQ.get((src, dst), {})
        return min(r, o[L]) if L in o else r

    def _requirements(self):
        req = {}
        pairs = [(p, i) for p, i, _ in self.links] + list(LOOP_REQ)
        for i in self.A:
            for L in range(1, self.top[i] + 1):
                r = {}
                for p, d in pairs:
                    if d == i:
                        n = self.need(p, i, L)
                        if n:
                            r[p] = max(r.get(p, 0), n)
                req[(i, L)] = r
        return req

    def _phase_of(self, i, L):
        return next(PH[k] for k, t in enumerate(TARGET[i]) if t >= L)

    def _depths(self):
        """Longest path to each (activity, level) step; raises on a cycle."""
        depth, state = {}, {}

        def preds(n):
            i, L = n
            out = [(p, l) for p, l in self.req[n].items()]
            if L > 1:
                out.append((i, L - 1))
            return out

        def visit(n, path):
            if state.get(n) == 2:
                return depth[n]
            if state.get(n) == 1:
                raise ValueError("maturity cycle: " + " -> ".join(f"{a}@{b}" for a, b in path + [n]))
            state[n] = 1
            d = 1 + max((visit(m, path + [n]) for m in preds(n)), default=0)
            state[n] = 2
            depth[n] = d
            return d

        for n in self.req:
            visit(n, [])
        return depth

    def _consistency(self):
        """A gate must not ask for a level whose inputs the same gate leaves less mature."""
        out = []
        for k, g in enumerate(PH):
            for i in self.A:
                L = TARGET[i][k]
                if not L:
                    continue
                for p, n in self.req[(i, L)].items():
                    if TARGET[p][k] < n:
                        out.append(f"{g}: {i} at level {L} needs {p} at level {n}, but the gate targets "
                                   f"{p} at level {TARGET[p][k]}")
        return out

    # ------------------------------------------------------------------ views
    def gate(self, k):
        rows = []
        for i in self.order():
            t = TARGET[i][k]
            if t:
                rows.append(dict(id=i, target=t, achieved=self.achieved[i], met=self.achieved[i] >= t,
                                 assumed=self.A[i]["status"] == "assumed"))
        return rows

    def order(self):
        first = {i: next(k for k, t in enumerate(TARGET[i]) if t) for i in self.A}
        return sorted(self.A, key=lambda i: (DISCIPLINES.index(self.A[i]["disc"]), first[i], i))

    def steps(self, ph):
        """Steps (activity, level) that belong to a phase, in network order, with their state today."""
        out = []
        for n in sorted((n for n in self.req if self.phase[n] == ph), key=lambda n: (self.depth[n], n)):
            i, L = n
            done = self.achieved[i] >= L
            wait = {p: l for p, l in self.req[n].items() if self.achieved[p] < l}
            if L > 1 and self.achieved[i] < L - 1:
                wait[i] = L - 1
            out.append(dict(id=i, level=L, depth=self.depth[n], done=done,
                            state="done" if done else ("ready" if not wait else "blocked"),
                            needs=self.req[n], wait=wait))
        return out


# ---------------------------------------------------------------------------------------------- outputs
def lv(n):
    return f"{n} {LEVELS[n]}"


def workbook(plan):
    wb = Workbook()
    hdr_fill, hdr_font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
    fills = {0: None, 1: "F3E5D8", 2: "E9C9A8", 3: "D7955F", 4: "B5541F"}

    def sheet(title, hdr, widths, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(hdr)
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill
        for k, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(k + 1)].width = w
        ws.freeze_panes = "C2"
        return ws

    ws = sheet("Matrix", ["ID", "Activity", "Discipline", "Status"] + [p["id"] for p in PHASES] +
               ["Achieved", "Note"], (9, 40, 13, 22, 7, 7, 7, 7, 10, 60), True)
    for i in plan.order():
        a = plan.A[i]
        ws.append([i, a["name"], a["disc"], STATUS[a["status"]], *TARGET[i], plan.achieved[i], a["note"]])
        for k in range(4):
            c = ws.cell(row=ws.max_row, column=5 + k)
            if fills[c.value]:
                c.fill = PatternFill("solid", fgColor=fills[c.value])
                c.font = Font(color="FFFFFF" if c.value >= 3 else "000000", bold=True)
    ws.append([])
    ws.append(["", "Levels: " + "; ".join(lv(k) for k in LEVELS)])

    ws = sheet("Gates", ["Gate", "ID", "Activity", "Target", "Achieved", "Met", "Note"], (8, 9, 40, 16, 16, 8, 50))
    for k, p in enumerate(PHASES):
        for r in plan.gate(k):
            ws.append([p["id"], r["id"], plan.A[r["id"]]["name"], lv(r["target"]), lv(r["achieved"]),
                       "assumed" if r["assumed"] else ("yes" if r["met"] else "no"), plan.A[r["id"]]["note"]])

    ws = sheet("Steps", ["Phase", "Order", "ID", "Activity", "To level", "State today", "Needs", "Waiting on"],
               (8, 7, 9, 40, 16, 12, 60, 40))
    for p in PHASES:
        for r in plan.steps(p["id"]):
            ws.append([p["id"], r["depth"], r["id"], plan.A[r["id"]]["name"], lv(r["level"]), r["state"],
                       ", ".join(f"{q} at {l}" for q, l in r["needs"].items()),
                       ", ".join(f"{q} at {l}" for q, l in r["wait"].items())])

    ws = sheet("Link maturity", ["From", "To", "Type", "What flows"] + [f"To at {L}" for L in range(1, 5)],
               (9, 9, 10, 46, 9, 9, 9, 9))
    for p, i, w in plan.links:
        ws.append([p, i, "input", w] + [plan.need(p, i, L) if L <= plan.top[i] else "" for L in range(1, 5)])
    for (p, i), (L, n) in LOOP_REQ.items():
        ws.append([p, i, "loop", plan.what[(p, i)]] + [n if x == L else "" for x in range(1, 5)])
    ws.append([])
    ws.append(["", "", "", "Cells give the level of 'From' needed for 'To' to reach each level (0 = not needed)."])

    ws = sheet("Consistency", ["Issue"], (120,))
    for s in plan.issues or ["none: every gate target is supported by its inputs at the same gate"]:
        ws.append([s])
    for w in wb.worksheets:
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT / f"{DOCNO}_{TITLE.replace(' ', '-')}.xlsx")


def report(plan):
    n_steps = len(plan.req)
    md = [f"""# 1 Purpose
This plan extends the design dependency map (CFU-000-PM-DEP-001) across the four phases of execution: FEL 1
(appraise), FEL 2 (select), FEL 3 (define / FEED) and detailed design. Most activities run in every phase. What
changes is how mature each one must be at each gate, so the plan is a matrix of target maturity by activity and
gate, plus the maturity each link needs.

Use it to:
- see what each gate requires, and what is still short;
- plan each phase as the steps that take activities from one level to the next, in the order the links allow;
- check that the gate targets are consistent (no gate asks for a result whose inputs are not mature enough).

# 2 Maturity scale
| Level | Name | Meaning |
|---|---|---|
| 0 | {LEVELS[0]} | — |
| 1 | {LEVELS[1]} | Sketches, factored or assumed data, option lists |
| 2 | {LEVELS[2]} | Issued for review (rev A); enough to estimate and screen |
| 3 | {LEVELS[3]} | Issued for design, HAZOP or purchase; the frozen basis for the next phase |
| 4 | {LEVELS[4]} | Issued for construction, approved or certified; one-off studies and decisions when complete |

# 3 Phases and gates
| Phase | Gate | Estimate | Purpose | Activities needed |
|---|---|---|---|---|
"""]
    for k, p in enumerate(PHASES):
        md.append(f"| {p['name']} | {p['gate']} | {p['est']} | {p['aim']} | {len(plan.gate(k))} |\n")
    md.append(f"""
The plan has {len(plan.acts)} activities ({len(ACTIVITIES)} from the dependency map, {len(PHASE_ACTIVITIES)} that
only appear once the phases are modelled), {len(plan.links)} input links, {len(LOOP_REQ)} iteration loops and
{n_steps} maturity steps.

# 4 Target matrix
Target level at each gate. "Now" is the level this package has reached.

| ID | Activity | Discipline | FEL 1 | FEL 2 | FEL 3 | DD | Now |
|---|---|---|---|---|---|---|---|
""")
    for i in plan.order():
        a = plan.A[i]
        t = TARGET[i]
        md.append(f"| {i} | {a['name']} | {a['disc']} | " + " | ".join(str(x) if x else "·" for x in t) +
                  f" | {plan.achieved[i]}{'*' if a['status'] == 'assumed' else ''} |\n")
    md.append("\n\\* assumed: decided in the design basis rather than developed in this package.\n")
    md.append("""
# 5 Link maturity rules
- By default, taking an activity to level L needs each of its inputs at level L.
- Studies, estimates and HAZOPs need early inputs only. Their cap: """ +
              ", ".join(f"{k} {v}" for k, v in CONSUME.items() if not isinstance(v, dict)) + """.
- Single-link exceptions:
""")
    for (p, i), o in LINK_REQ.items():
        md.append(f"  - {p} → {i}: " + ", ".join(f"at level {L} needs {n or 'nothing'}" for L, n in o.items())
                  + f" ({plan.what.get((p, i), '')})\n")
    md.append("- Iteration loops become forward links between levels, so the plan has no cycles:\n")
    for (p, i), (L, n) in LOOP_REQ.items():
        md.append(f"  - {i} reaches level {L} after {p} reaches level {n} ({plan.what[(p, i)]})\n")
    md.append("\n# 6 Gate readiness of this package\n")
    md.append("| Gate | Needed | Met | Assumed | Short |\n|---|---|---|---|---|\n")
    for k, p in enumerate(PHASES):
        g = plan.gate(k)
        met = sum(r["met"] and not r["assumed"] for r in g)
        ass = sum(r["assumed"] for r in g)
        md.append(f"| {p['gate']} | {len(g)} | {met} | {ass} | {len(g) - met - ass} |\n")
    for k, p in enumerate(PHASES[:3]):
        short = [r for r in plan.gate(k) if not r["met"]]
        md.append(f"\n**{p['gate']}: short**\n\n")
        md += [f"- {r['id']} {plan.A[r['id']]['name']}: at {lv(r['achieved']).lower()}, needs "
               f"{lv(r['target']).lower()}\n" for r in short] or ["- nothing\n"]
    md.append("\n# 7 Steps by phase\nEach step takes an activity to a level. Steps are in network order; "
              "\"ready\" means every input is already at the level the step needs.\n")
    for p in PHASES:
        st = plan.steps(p["id"])
        cnt = {s: sum(r["state"] == s for r in st) for s in ("done", "ready", "blocked")}
        md.append(f"\n## {p['name']}\n{len(st)} steps: {cnt['done']} done, {cnt['ready']} ready, "
                  f"{cnt['blocked']} blocked.\n\n| Order | ID | Activity | To level | Today | Waiting on |\n"
                  "|---|---|---|---|---|---|\n")
        for r in st:
            md.append(f"| {r['depth']} | {r['id']} | {plan.A[r['id']]['name']} | {lv(r['level'])} | {r['state']} | "
                      f"{', '.join(f'{q} at {l}' for q, l in r['wait'].items()) or '-'} |\n")
    md.append("\n# 8 Consistency check\n")
    md += [f"- {s}\n" for s in plan.issues] or ["- none: every gate target is supported by its inputs at the "
                                                "same gate, and the step network has no cycles\n"]
    md.append("""
# 9 Limits
- Targets are typical values for a refinery process unit and should be agreed with the owner's gate criteria
  (for example CII PDRI or IPA FEL index) before use.
- One level per activity: an activity whose deliverables mature at different rates (for example equipment
  datasheets for long-lead and bulk items) is shown at the level of most of its deliverables.
- "Achieved" is taken from the package status: issued = the FEL 3 target, provisional = the FEL 2 target.
- Construction, commissioning and handover phases are not modelled.
""")
    docgen.render("".join(md), OUT / f"{DOCNO}_{TITLE.replace(' ', '-')}", DOCNO, TITLE)


def page(plan):
    from .depmap import Net, doc_map
    from .portal import GH
    docs, _, _ = doc_map(Net())
    acts = []
    for i in plan.order():
        a = plan.A[i]
        acts.append(dict(id=i, name=a["name"], disc=a["disc"], status=a["status"], note=a["note"], ext=a["ext"],
                         target=TARGET[i], now=plan.achieved[i], isNew=i not in {x["id"] for x in ACTIVITIES}, docs=len(docs.get(i, [])),
                         levels=[dict(L=L, ph=plan.phase[(i, L)], needs=plan.req[(i, L)])
                                 for L in range(1, plan.top[i] + 1)]))
    gates = []
    for k, p in enumerate(PHASES):
        g = plan.gate(k)
        gates.append(dict(p, need=len(g), met=sum(r["met"] and not r["assumed"] for r in g),
                          assumed=sum(r["assumed"] for r in g),
                          short=[dict(id=r["id"], t=r["target"], a=r["achieved"]) for r in g if not r["met"]],
                          steps=plan.steps(p["id"])))
    data = dict(acts=acts, gates=gates, levels=LEVELS, status=STATUS, issues=plan.issues,
                names={i: plan.A[i]["name"] for i in plan.A})
    tpl = Path(__file__).with_name("phases_page.html").read_text()
    out = (tpl.replace("/*DATA*/null", json.dumps(data, separators=(",", ":")).replace("</", "<\\/"))
              .replace("{{N_ACT}}", str(len(plan.acts))).replace("{{N_STEPS}}", str(len(plan.req)))
              .replace("{{N_NEW}}", str(len(PHASE_ACTIVITIES))).replace("{{N_LOOPS}}", str(len(LOOP_REQ)))
              .replace("{{REV}}", html.escape(basis.PROJECT["rev"])).replace("{{GH}}", GH)
              .replace("{{DOC}}", f"{DOCNO}_{TITLE.replace(' ', '-')}"))
    PORTAL.mkdir(parents=True, exist_ok=True)
    (PORTAL / "phases.html").write_text(out)


def build():
    plan = Plan()
    OUT.mkdir(parents=True, exist_ok=True)
    workbook(plan)
    report(plan)
    page(plan)
    for s in plan.issues:
        print("  PHASES:", s)
    return plan


if __name__ == "__main__":
    build()
