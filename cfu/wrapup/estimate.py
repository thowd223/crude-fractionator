"""AACE Class 4 factored capital cost estimate (CFU-000-PM-EST-001).

Purchased-equipment costs: Towler & Sinnott (2nd ed.) correlations Ce = a + b*S^n, USGC Jan-2010
(CEPCI 532.9), escalated to 2026 with CEPCI 820 (assumed). Installed cost by the Towler/Hand
factor method for fluids processing; ISBL only (OSBL interfaces excluded).
"""
from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .. import docgen

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "deliverables" / "06-wrapup"
CEPCI_2010, CEPCI_2026 = 532.9, 820.0
ESC = CEPCI_2026 / CEPCI_2010
RHO_STEEL = 7850.0

FACT = dict(fer=0.3, fp=0.8, fi=0.3, fel=0.2, fc=0.3, fs=0.2, fl=0.1)
OFFSITE_EXCL = "Tankage, SWS, flare stack, CCR, main substation, OSBL racks"


def fm(moc: str) -> float:
    s = (moc or "").lower()
    if "monel" in s:
        return 1.9
    if "9cr" in s or "p9" in s:
        return 1.7
    if "clad" in s:
        return 1.6
    if "5cr" in s or "1.25cr" in s:
        return 1.4
    if "ss" in s or "316" in s or "304" in s:
        return 1.3
    return 1.0


def shell_mass(e, mech):
    if mech and e["tag"] in mech and mech[e["tag"]].get("shell_weight_kg"):
        return mech[e["tag"]]["shell_weight_kg"]
    D = e.get("D") or 2.0
    L = e.get("H") if e.get("orient") == "V" else e.get("L") or e.get("H") or 5.0
    P = e.get("des_P")
    P = 3.5 if not isinstance(P, (int, float)) else P
    t = max(P * 1e5 * D / (2 * 138e6 * 0.85 - 1.2 * P * 1e5) + 0.003, 0.008 + D / 1000 * 2)
    return math.pi * D * L * t * RHO_STEEL * 1.25  # heads, nozzles, internals supports


def ce_items(eq, mech):
    rows = []
    for e in eq:
        typ, tag = e["type"], e["tag"]
        n = 2 if tag.endswith("A/B") else 1
        cost, basis_s = 0.0, ""
        if typ in ("Column", "Drum", "Desalter"):
            W = shell_mass(e, mech)
            W = min(W, 250_000)
            vert = e.get("orient") == "V"
            a, b = (11600, 34) if vert else (10200, 31)
            cost = (a + b * W ** 0.85) * fm(e.get("moc"))
            basis_s = f"shell {W / 1000:.0f} t"
            if typ == "Column" and "tray" in (e.get("internals") or ""):
                import re
                ntr = int(re.findall(r"(\d+)\s+valve", e["internals"])[0]) if re.findall(r"(\d+)\s+valve", e["internals"]) else 20
                D = min(e["D"], 5.0)
                cost += ntr * (180 + 340 * D ** 1.9) * 1.3 * (e["D"] / D) ** 2
                basis_s += f", {ntr} trays"
            if tag == "C-201":
                vol = math.pi / 4 * e["D"] ** 2 * 9.0
                cost += vol * 7600
                basis_s += f", {vol:.0f} m3 packing"
            if typ == "Desalter":
                cost *= 2.2   # electrostatic internals, transformers
        elif typ == "Fired heater":
            Q = e["duty_kw"] / 1000
            cost = (43000 + 111000 * Q ** 0.8) * 1.25 * 2.5  # 9Cr tubes; x2.5 refinery API 560 scope (APH, stack, BMS, ULNB)
            basis_s = f"{Q:.1f} MW box"
        elif typ in ("Shell & tube", "Air preheater"):
            A = e.get("area_m2") or 500
            sh = e.get("n_shells") or 1
            per = max(min(A / sh, 1000), 10)
            cost = sh * (32000 + 70 * per ** 1.2) * fm(e.get("moc"))
            basis_s = f"{A:.0f} m2, {sh} shell(s)"
        elif typ == "Air cooler":
            cost = 1100 * e["area_m2"] + 25000 * e["fans"]
            basis_s = f"{e['area_m2']:.0f} m2 bare"
        elif typ == "Pump":
            q = e["flow_m3h"] / 3.6
            cost = n * ((8000 + 240 * q ** 0.9) * (1.6 if e.get("api610") == "BB2" else 1.2)
                        + (-950 + 1770 * e["motor_kw"] ** 0.6))
            basis_s = f"{n} x {e['flow_m3h']:.0f} m3/h, {e['motor_kw']} kW"
        elif typ == "Fan":
            cost = n * (4000 + 57 * (e["absorbed_kw"] * 10) ** 0.8 + (-950 + 1770 * e["motor_kw"] ** 0.6))
            basis_s = f"{n} x {e['motor_kw']} kW"
        elif typ == "Ejector":
            cost = 2 * 45000
            basis_s = "2 x 50 % train"
        elif typ == "Package":
            cost = 120000
            basis_s = "vendor package"
        rows.append(dict(tag=tag, type=typ, service=e["service"], basis=basis_s, fm=fm(e.get("moc")),
                         ce_2010=cost, ce_2026=cost * ESC))
    return rows


def build():
    eq = json.loads((ROOT / "data" / "equipment.json").read_text())
    mp = ROOT / "data" / "mech.json"
    mech = None
    if mp.exists():
        raw = json.loads(mp.read_text())
        items = raw if isinstance(raw, list) else raw.get("items", raw)
        if isinstance(items, dict):
            mech = {k: v for k, v in items.items() if isinstance(v, dict)}
        else:
            mech = {x["tag"]: x for x in items if isinstance(x, dict) and "tag" in x}
    rows = ce_items(eq, mech)
    Ce = sum(r["ce_2026"] for r in rows)
    # Towler: ISBL = sum Ce[(1+fp)fm + (fer+fel+fi+fc+fs+fl)] ; fm applied in Ce already for vessel shell,
    # piping fp scales with alloy fraction -> use 1.15 average
    isbl = sum(r["ce_2026"] * ((1 + FACT["fp"] * 1.15) + (FACT["fer"] + FACT["fel"] + FACT["fi"] + FACT["fc"]
                                                           + FACT["fs"] + FACT["fl"]) / max(r["fm"], 1)) for r in rows)
    de = 0.25 * isbl
    cont = 0.20 * (isbl + de)
    tic = isbl + de + cont
    by_type = {}
    for r in rows:
        by_type[r["type"]] = by_type.get(r["type"], 0) + r["ce_2026"]
    summary = [("Purchased equipment (2026 USGC)", Ce), ("Installed ISBL (Towler factors)", isbl),
               ("Design & engineering (25 %)", de), ("Contingency (20 %)", cont), ("Total installed cost, ISBL", tic)]
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for r in [("Item", "USD")] + summary:
        ws.append(list(r))
    ws.append([])
    ws.append(["Per bbl/d capacity", tic / 100000])
    ws2 = wb.create_sheet("Equipment")
    ws2.append(["Tag", "Type", "Service", "Sizing basis", "Material factor", "Ce 2010 USD", "Ce 2026 USD"])
    for r in rows:
        ws2.append([r["tag"], r["type"], r["service"], r["basis"], r["fm"], round(r["ce_2010"]), round(r["ce_2026"])])
    for w in (ws, ws2):
        for c in w[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
        w.column_dimensions["A"].width = 36
        for col in "BCDEFG":
            w.column_dimensions[col].width = 18
    wb.save(OUT / "CFU-000-PM-EST-001_Cost-Estimate.xlsx")
    M = 1e6
    md = f"""# 1 Basis
- AACE Class 4 (expected accuracy -30 % / +50 %), ISBL only, US Gulf Coast, 2026 USD (CEPCI {CEPCI_2026:.0f}, assumed).
- Purchased equipment: Towler & Sinnott correlations (CEPCI {CEPCI_2010}), with material factors per equipment
  metallurgy (5Cr 1.4, 9Cr 1.7, clad 1.6, Monel 1.9). Vessel shell weights come from the mechanical calculations
  (data/mech.json) where available{'' if mech else ' (not available: estimated from design pressure and dimensions)'}.
- Installation: Towler fluids-processing factors (erection {FACT['fer']}, piping {FACT['fp']}, instrumentation
  {FACT['fi']}, electrical {FACT['fel']}, civil {FACT['fc']}, structures {FACT['fs']}, lagging/paint {FACT['fl']}).
- Excluded: {OFFSITE_EXCL}; owner's costs, licence fees, land, escalation beyond 2026, catalyst/chemicals first fill.

# 2 Summary
| Item | USD million |
|---|---|
""" + "".join(f"| {k} | {v / M:,.1f} |\n" for k, v in summary) + f"""
Specific cost: **{tic / 100000:,.0f} USD per BPSD** of capacity, ISBL. Check this against the owner's own
benchmark data for grassroots CDU/VDU units before using it for budgeting. The factored method tends to
under-estimate large, alloy-heavy refinery units.

# 3 Purchased equipment by type
| Type | USD million | Share |
|---|---|---|
""" + "".join(f"| {k} | {v / M:,.2f} | {v / Ce:.0%} |\n" for k, v in sorted(by_type.items(), key=lambda x: -x[1])) + """
# 4 Equipment detail
| Tag | Service | Basis | Ce 2026 kUSD |
|---|---|---|---|
""" + "".join(f"| {r['tag']} | {r['service']} | {r['basis']} | {r['ce_2026'] / 1000:,.0f} |\n" for r in rows) + """
# 5 Accuracy and next steps
The heaters, the column shells and internals, and the hot exchanger bank account for more than 60 % of
equipment cost. Vendor budget quotes for H-101, H-201, C-101, C-201 and the desalters should be obtained to
move the estimate to Class 3.
"""
    docgen.render(md, OUT / "CFU-000-PM-EST-001_Cost-Estimate", "CFU-000-PM-EST-001", "Capital Cost Estimate (Class 4)")
    return tic
