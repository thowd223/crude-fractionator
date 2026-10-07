"""Design dependency map (CFU-000-PM-DEP-001 report + workbook, CFU-000-PM-DEP-002 network drawing,
portal/dependencies.html interactive map).

The activity network is declared in activities.py. It is checked against the code: `python -m
cfu.wrapup.depmap --trace` runs the full build with file access recorded and writes data/dataflow.json
(which module read which data file). Every traced read must come from an upstream activity, or from a
declared iteration loop (the module then reads the previous revision of that file).
"""
from __future__ import annotations

import html
import json
import sys
import textwrap
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import basis, docgen
from .activities import ACTIVITIES, DISCIPLINES, FEEDBACK, HELPERS

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
OUT = ROOT / "deliverables" / "06-wrapup"
PORTAL = OUT / "portal"
STATUS = {"done": "Issued (FEED)", "partial": "Issued, provisional", "open": "Not started"}
DEP_DOCS = [("CFU-000-PM-DEP-001", "Design Dependency Map"), ("CFU-000-PM-DEP-002", "Design Network Diagram")]


# ---------------------------------------------------------------------------------------------- graph
class Net:
    def __init__(self):
        self.A = {a["id"]: a for a in ACTIVITIES}
        self.order_ids = [a["id"] for a in ACTIVITIES]
        bad = [(a["id"], p) for a in ACTIVITIES for p, _ in a["after"] if p not in self.A]
        bad += [(f, t) for f, t, _ in FEEDBACK if f not in self.A or t not in self.A]
        if bad:
            raise ValueError(f"unknown activity references: {bad}")
        self.pred = {i: [p for p, _ in self.A[i]["after"]] for i in self.A}
        self.succ = {i: [] for i in self.A}
        for i in self.A:
            for p in self.pred[i]:
                self.succ[p].append(i)
        self.flow = {(p, a["id"]): w for a in ACTIVITIES for p, w in a["after"]}
        self.fb = [dict(src=f, dst=t, what=w) for f, t, w in FEEDBACK]
        self.wave = {}
        self._waves()
        self.up = {i: self._closure(i, self.pred) for i in self.A}
        self.down = {i: self._closure(i, self.succ) for i in self.A}
        self.seq = sorted(self.A, key=lambda i: (self.wave[i], DISCIPLINES.index(self.A[i]["disc"]), i))
        self.writer = {f: a["id"] for a in ACTIVITIES for f in a["writes"]}

    def _waves(self):
        state = {}

        def visit(i, path):
            if state.get(i) == 2:
                return self.wave[i]
            if state.get(i) == 1:
                raise ValueError("forward cycle: " + " -> ".join(path + [i]))
            state[i] = 1
            w = 1 + max((visit(p, path + [i]) for p in self.pred[i]), default=0)
            state[i] = 2
            self.wave[i] = w
            return w

        for i in self.A:
            visit(i, [])
        # iteration loops must point backwards, otherwise they are ordinary forward links
        for f in self.fb:
            if self.wave[f["dst"]] > self.wave[f["src"]]:
                raise ValueError(f"feedback {f['src']}->{f['dst']} points forward; declare it in 'after'")

    @staticmethod
    def _closure(i, nbr):
        seen, stack = set(), list(nbr[i])
        while stack:
            j = stack.pop()
            if j not in seen:
                seen.add(j)
                stack.extend(nbr[j])
        return seen


# ---------------------------------------------------------------------------------------------- docs
def doc_map(net):
    from . import register
    rows = {r["no"]: r for r in register()}
    for no, title in DEP_DOCS:
        rows.setdefault(no, dict(no=no, title=title, folder="06-wrapup"))
    pref = sorted(((p, a["id"]) for a in ACTIVITIES for p in a["docs"]), key=lambda x: -len(x[0]))
    owner, unmapped = {}, []
    for no in sorted(rows):
        hit = next((aid for p, aid in pref if no.startswith(p)), None)
        if hit:
            owner[no] = hit
        else:
            unmapped.append(no)
    docs = {i: sorted(no for no, a in owner.items() if a == i) for i in net.A}
    titles = {no: (r["title"] if r["title"] != "(drawing)" else "") for no, r in rows.items()}
    return docs, titles, unmapped


# ---------------------------------------------------------------------------------------------- code check
def trace():
    """Run the full build, record every data/*.json read and write with the module responsible."""
    import builtins
    import pathlib

    recs = set()

    def who():
        f = sys._getframe(2)
        while f:
            try:
                rel = Path(f.f_code.co_filename).resolve().relative_to(ROOT).as_posix()
            except ValueError:
                rel = None
            if rel and ((rel.startswith("cfu/") and rel not in HELPERS) or rel == "build.py"):
                return rel
            f = f.f_back
        return "?"

    def note(path, mode):
        try:
            p = Path(path).resolve()
        except (TypeError, OSError):
            return
        if p.parent == DATA and p.suffix == ".json" and p.name != "dataflow.json":
            recs.add((who(), p.name, "write" if any(c in str(mode) for c in "wax") else "read"))

    b_open, p_open = builtins.open, pathlib.Path.open

    def o(file, mode="r", *a, **k):
        if isinstance(file, (str, Path)):
            note(file, mode)
        return b_open(file, mode, *a, **k)

    def po(self, mode="r", *a, **k):
        note(self, mode)
        return p_open(self, mode, *a, **k)

    builtins.open, pathlib.Path.open = o, po
    try:
        sys.path.insert(0, str(ROOT))
        import build
        build.main(["build.py"])
    finally:
        builtins.open, pathlib.Path.open = b_open, p_open
    out = [dict(module=m, data=d, op=op) for m, d, op in sorted(recs)]
    (DATA / "dataflow.json").write_text(json.dumps(out, indent=1))
    return out


def stage(module):
    """Position in build.STAGES of the stage that runs a module (data written by a later stage is a previous
    revision when this module reads it)."""
    sys.path.insert(0, str(ROOT))
    import build
    best, idx = "", 0
    for k, (_, fn) in enumerate(build.STAGES):
        if not isinstance(fn, str):
            continue
        mod = fn.split(":")[0].replace(".", "/")
        for pref in (mod + ".py", mod + "/"):
            if module.startswith(pref) and len(pref) > len(best):
                best, idx = pref, k
    return idx


def code_check(net):
    mod = {c: a["id"] for a in ACTIVITIES for c in a["code"]}
    fb_pairs = {(f["src"], f["dst"]) for f in net.fb}
    flow_p = DATA / "dataflow.json"
    rows, issues = [], []
    if not flow_p.exists():
        return rows, ["data/dataflow.json not found: run python -m cfu.wrapup.depmap --trace"]
    flow = json.loads(flow_p.read_text())
    wmod = {r["data"]: r["module"] for r in flow if r["op"] == "write"}
    for r in flow:
        if r["op"] != "read":
            w = net.writer.get(r["data"])
            if r["module"] in mod and w and mod[r["module"]] != w:
                issues.append(f"{r['module']} writes {r['data']}, owned by {w}")
            continue
        act, prod = mod.get(r["module"]), net.writer.get(r["data"])
        if act is None:
            res = "module not mapped to an activity"
        elif prod is None:
            res = "data file has no declared owner"
        elif prod == act:
            res = "own data"
        elif prod in net.up[act]:
            res = "upstream"
        elif any(p == prod and (d == act or d in net.up[act]) for p, d in fb_pairs) or (prod, act) in fb_pairs:
            res = "iteration (reads previous revision)"
        else:
            res = "UNDECLARED"
        if res in ("upstream", "own data") and r["data"] in wmod and stage(wmod[r["data"]]) > stage(r["module"]):
            res = "build order: reads previous revision"
        rows.append(dict(module=r["module"], activity=act or "-", data=r["data"], producer=prod or "-", result=res))
        if res in ("UNDECLARED", "module not mapped to an activity", "data file has no declared owner"):
            issues.append(f"{r['module']} ({act or '?'}) reads {r['data']} from {prod or '?'}: {res}")
    # every generator module should belong to an activity (package __init__ files only dispatch)
    for p in sorted((ROOT / "cfu").rglob("*.py")):
        rel = p.relative_to(ROOT).as_posix()
        if rel in mod or rel in HELPERS or p.name == "__init__.py" or rel in (
                "cfu/pid/__init__.py",):
            continue
        issues.append(f"{rel} is not assigned to an activity")
    return rows, issues


# ---------------------------------------------------------------------------------------------- geometry
CW, NW, NH, RH, LANE_W, PAD = 196, 168, 62, 76, 112, 14


def geometry(net):
    lanes, y = [], PAD + 26
    pos = {}
    for d in DISCIPLINES:
        ids = [i for i in net.seq if net.A[i]["disc"] == d]
        if not ids:
            continue
        cols = {}
        for i in ids:
            cols.setdefault(net.wave[i], []).append(i)
        rows = max(len(v) for v in cols.values())
        for w, v in cols.items():
            for k, i in enumerate(v):
                pos[i] = (LANE_W + (w - 1) * CW + (CW - NW) / 2, y + 7 + k * RH)
        lanes.append(dict(disc=d, y=y, h=rows * RH))
        y += rows * RH
    W = LANE_W + max(net.wave.values()) * CW + PAD
    H = y + PAD
    edges = []
    for (p, i), what in net.flow.items():
        x0, y0 = pos[p][0] + NW, pos[p][1] + NH / 2
        x1, y1 = pos[i][0], pos[i][1] + NH / 2
        dx = max((x1 - x0) * 0.5, 30)
        edges.append(dict(src=p, dst=i, what=what, kind="flow",
                          d=f"M{x0:.0f},{y0:.0f} C{x0 + dx:.0f},{y0:.0f} {x1 - dx:.0f},{y1:.0f} {x1:.0f},{y1:.0f}"))
    for f in net.fb:
        x0, y0 = pos[f["src"]][0] + NW / 2, pos[f["src"]][1]
        x1, y1 = pos[f["dst"]][0] + NW / 2, pos[f["dst"]][1]
        lift = 34 + 0.08 * abs(x0 - x1)
        top = min(y0, y1) - lift
        edges.append(dict(src=f["src"], dst=f["dst"], what=f["what"], kind="loop",
                          d=f"M{x0:.0f},{y0:.0f} C{x0:.0f},{top:.0f} {x1:.0f},{top:.0f} {x1:.0f},{y1:.0f}"))
    return dict(pos=pos, lanes=lanes, W=W, H=H, edges=edges, waves=max(net.wave.values()))


def wrap(s, n=25, lines=3):
    out = textwrap.wrap(s, n)
    if len(out) > lines:
        out = out[:lines]
        out[-1] = out[-1][: n - 1] + "…"
    return out


# ---------------------------------------------------------------------------------------------- drawing
def drawing(net, G):
    from ..drawing.sheet import Sheet
    sh = Sheet("A1", "DESIGN DEPENDENCY NETWORK", "ACTIVITIES, LINKS AND ITERATION LOOPS", "CFU-000-PM-DEP-002",
               discipline="PROJECT", notes=[
                   "Columns are sequence waves: an activity can start when every activity it needs (solid arrows) has issued.",
                   "Dashed arrows are iteration loops: a later result revises an earlier activity.",
                   "Box style: solid = issued (FEED); double = issued, provisional; dashed grey = not started.",
                   "Generated from cfu/wrapup/activities.py and checked against data/dataflow.json (build trace).",
                   "See CFU-000-PM-DEP-001 for inputs, impacts and the change-impact matrix."])
    x0, y0, x1, y1 = sh.area
    s = min((x1 - x0) / G["W"], (y1 - y0) / G["H"])
    d, g = sh.dwg, sh.dwg.g(transform=f"translate({x0},{y0}) scale({s})", font_family="DejaVu Sans, Arial, sans-serif")
    sh.dwg.add(g)
    for k, ln in enumerate(G["lanes"]):
        g.add(d.rect((0, ln["y"]), (G["W"], ln["h"]), fill="#f2f5f7" if k % 2 == 0 else "#ffffff", stroke="#9aa7b0",
                     stroke_width=0.6))
        g.add(d.text(ln["disc"].upper(), insert=(10, ln["y"] + 20), font_size=13, font_weight="bold", fill="#1d2b36"))
    for w in range(1, G["waves"] + 1):
        g.add(d.text(f"WAVE {w}", insert=(LANE_W + (w - 0.5) * CW, 24), font_size=12, text_anchor="middle",
                     fill="#1d2b36", font_weight="bold"))
    for e in G["edges"]:
        loop = e["kind"] == "loop"
        g.add(d.path(d=e["d"], fill="none", stroke="#c8501e" if loop else "#2f5d7c", stroke_width=1.1 if loop else 0.9,
                     stroke_dasharray="5,3" if loop else "none", opacity=0.75))
        x, y = map(float, e["d"].split(" ")[-1].split(","))
        if not loop:
            g.add(d.polygon([(x, y), (x - 6, y - 3), (x - 6, y + 3)], fill="#2f5d7c"))
        else:
            g.add(d.polygon([(x, y), (x - 3, y - 6), (x + 3, y - 6)], fill="#c8501e"))
    for i, (x, y) in G["pos"].items():
        a = net.A[i]
        st = a["status"]
        g.add(d.rect((x, y), (NW, NH), fill="#ffffff" if st != "open" else "#eceff1",
                     stroke="#1d2b36" if st != "open" else "#7d8a93", stroke_width=1.4 if st == "done" else 1.0,
                     stroke_dasharray="4,3" if st == "open" else "none"))
        if st == "partial":
            g.add(d.rect((x + 3, y + 3), (NW - 6, NH - 6), fill="none", stroke="#1d2b36", stroke_width=0.6))
        g.add(d.text(i, insert=(x + 8, y + 16), font_size=11, font_weight="bold", fill="#c8501e" if st == "open" else "#1d2b36"))
        for k, ln in enumerate(wrap(a["name"])):
            g.add(d.text(ln, insert=(x + 8, y + 30 + k * 12), font_size=10.5, fill="#1d2b36"))
    sh.save(OUT / "CFU-000-PM-DEP-002_Design-Network-Diagram")


# ---------------------------------------------------------------------------------------------- workbook
def workbook(net, docs, titles, chk_rows, chk_issues):
    wb = Workbook()
    hdr_fill, hdr_font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")

    def sheet(title, hdr, widths, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(hdr)
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill
        for k, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(k + 1)].width = w
        ws.freeze_panes = "A2"
        return ws

    ws = sheet("Sequence", ["Wave", "ID", "Activity", "Discipline", "Status", "Needs (direct)", "External inputs",
                            "Data owned", "Deliverables", "Feeds (direct)", "Downstream activities",
                            "Downstream deliverables", "Note"], (6, 9, 34, 12, 18, 48, 40, 22, 30, 30, 10, 10, 50), True)
    for i in net.seq:
        a = net.A[i]
        dd = sum(len(docs[j]) for j in net.down[i])
        ws.append([net.wave[i], i, a["name"], a["disc"], STATUS[a["status"]],
                   "\n".join(f"{p}: {w}" for p, w in a["after"]), "\n".join(a["ext"]), ", ".join(a["writes"]),
                   ", ".join(docs[i]), ", ".join(sorted(net.succ[i])), len(net.down[i]), dd, a["note"]])
    ws = sheet("Links", ["From", "To", "Type", "What flows / what is revised"], (9, 9, 12, 70))
    for (p, i), w in sorted(net.flow.items(), key=lambda x: (net.wave[x[0][0]], x[0])):
        ws.append([p, i, "input", w])
    for f in net.fb:
        ws.append([f["src"], f["dst"], "iteration", f["what"]])
    ws = sheet("DSM", ["Reads from ->"] + net.seq, [24] + [4] * len(net.seq))
    ws.column_dimensions["A"].width = 24
    fb = {(f["dst"], f["src"]) for f in net.fb}
    for i in net.seq:
        ws.append([f"{i} {net.A[i]['name'][:16]}"] + [
            "X" if j in net.pred[i] else ("F" if (i, j) in fb else ("·" if i == j else "")) for j in net.seq])
    ws.append([])
    ws.append(["Row reads from column. X = input (below diagonal), F = iteration loop (above diagonal)."])
    ws = sheet("Change impact", ["If this changes", "Activity", "Re-check activities (all downstream)",
                                 "Deliverables to revise"], (12, 34, 60, 90))
    for i in sorted(net.A, key=lambda i: -len(net.down[i])):
        dd = sorted({no for j in net.down[i] | {i} for no in docs[j]})
        ws.append([i, net.A[i]["name"], ", ".join(sorted(net.down[i], key=net.seq.index)), ", ".join(dd)])
    ws = sheet("To do", ["Order", "ID", "Activity", "Status", "Ready to start?", "Waiting on", "Then re-check",
                         "Note"], (6, 9, 36, 18, 14, 36, 60, 50))
    todo = [i for i in net.seq if net.A[i]["status"] != "done"]
    for k, i in enumerate(todo, 1):
        wait = [p for p in net.pred[i] if net.A[p]["status"] == "open"]
        ws.append([k, i, net.A[i]["name"], STATUS[net.A[i]["status"]], "yes" if not wait else "no", ", ".join(wait),
                   ", ".join(sorted((j for j in net.down[i] if net.A[j]["status"] != "open"), key=net.seq.index)),
                   net.A[i]["note"]])
    ws = sheet("Code check", ["Module", "Activity", "Reads", "Owner", "Result"], (34, 10, 22, 10, 34))
    for r in chk_rows:
        ws.append([r["module"], r["activity"], r["data"], r["producer"], r["result"]])
    for s in chk_issues:
        ws.append(["ISSUE", "", "", "", s])
    for w in wb.worksheets:
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT / "CFU-000-PM-DEP-001_Design-Dependency-Map.xlsx")


# ---------------------------------------------------------------------------------------------- report
def report(net, docs, chk_rows, chk_issues):
    n = {s: sum(a["status"] == s for a in ACTIVITIES) for s in STATUS}
    ndoc = sum(len(v) for v in docs.values())
    res = {}
    for r in chk_rows:
        res[r["result"]] = res.get(r["result"], 0) + 1
    md = [f"""# 1 Purpose
This document maps the engineering activities of the CDU/VDU design: the order they run in, what each one needs
from upstream, and what it affects downstream. Use it to plan the remaining work and to judge the impact of a
change. It covers the {len(ACTIVITIES)} activities that produce this package and the activities still needed
before the design can be used for construction.

How to read it:
- An activity's **needs** are the activities whose results it uses. It can start when they have issued.
- A **wave** is a step in the sequence: wave 1 needs nothing, wave 2 needs only wave 1 results, and so on.
  Activities in the same wave can run in parallel.
- An **iteration loop** is a later result that revises an earlier activity (for example, vendor data revising
  the mechanical calculations). Loops are why FEED takes more than one pass.
- **Downstream** means every activity reachable through the needs links. If an activity changes, all its
  downstream activities must be re-checked.

The network drawing is CFU-000-PM-DEP-002. The workbook has the full sequence, link list, design structure
matrix (DSM), change-impact table, to-do list and code check. An interactive version is in the portal
(dependencies.html).

# 2 Summary
| Item | Value |
|---|---|
| Activities | {len(ACTIVITIES)} ({n['done']} issued, {n['partial']} issued but provisional, {n['open']} not started) |
| Sequence waves | {max(net.wave.values())} |
| Input links | {len(net.flow)} |
| Iteration loops | {len(net.fb)} |
| Deliverables mapped to activities | {ndoc} |
| Data reads traced in the build | {len(chk_rows)} ({', '.join(f'{v} {k}' for k, v in sorted(res.items()))}) |
| Code-check issues | {len(chk_issues)} |

# 3 Design sequence
| Wave | ID | Activity | Discipline | Status | Needs |
|---|---|---|---|---|---|
"""]
    for i in net.seq:
        a = net.A[i]
        md.append(f"| {net.wave[i]} | {i} | {a['name']} | {a['disc']} | {STATUS[a['status']]} | "
                  f"{', '.join(net.pred[i]) or '-'} |\n")
    md.append("""
# 4 Inputs and impacts by activity
For each activity: what it needs (and what flows across each link), what it needs from outside the design team,
and what it feeds. Downstream counts are the activities and deliverables to re-check if it changes.

""")
    for i in net.seq:
        a = net.A[i]
        dd = sum(len(docs[j]) for j in net.down[i])
        md.append(f"**{i} {a['name']}** ({a['disc']}, {STATUS[a['status']].lower()})\n\n")
        if a["after"]:
            md.append("- Needs: " + "; ".join(f"{p} ({w})" for p, w in a["after"]) + "\n")
        if a["ext"]:
            md.append("- External inputs: " + "; ".join(a["ext"]) + "\n")
        if a["writes"]:
            md.append("- Owns data: " + ", ".join(f"data/{f}" for f in a["writes"]) + "\n")
        if docs[i]:
            md.append("- Issues: " + ", ".join(docs[i]) + "\n")
        md.append(f"- Feeds: {', '.join(sorted(net.succ[i], key=net.seq.index)) or 'nothing (end of chain)'}; "
                  f"{len(net.down[i])} downstream activities, {dd} downstream deliverables\n")
        loops_in = [f for f in net.fb if f["dst"] == i]
        if loops_in:
            md.append("- Revised by: " + "; ".join(f"{f['src']} ({f['what']})" for f in loops_in) + "\n")
        if a["note"]:
            md.append(f"- Note: {a['note']}\n")
        md.append("\n")
    md.append("""# 5 Iteration loops
| From | Revises | What changes |
|---|---|---|
""")
    for f in net.fb:
        md.append(f"| {f['src']} | {f['dst']} | {f['what']} |\n")
    md.append("""
Plan each loop as a formal update: issue the upstream activity, collect the later results, then re-issue.
The vendor-data loop (VEN-01) and the HAZOP loops (SAF-01, SAF-03) usually set the FEED schedule.

# 6 Change impact
The activities whose change reaches the most deliverables. Freeze these first.

| ID | Activity | Downstream activities | Downstream deliverables |
|---|---|---|---|
""")
    for i in sorted(net.A, key=lambda i: (-sum(len(docs[j]) for j in net.down[i]), i))[:14]:
        md.append(f"| {i} | {net.A[i]['name']} | {len(net.down[i])} | {sum(len(docs[j]) for j in net.down[i])} |\n")
    md.append("""
# 7 What still needs to be done
In sequence order. "Ready" means nothing it needs is still not started. "Re-opens" lists the issued activities
that must be re-checked when it completes.

| Order | ID | Activity | Status | Ready | Re-opens |
|---|---|---|---|---|---|
""")
    for k, i in enumerate((i for i in net.seq if net.A[i]["status"] != "done"), 1):
        wait = [p for p in net.pred[i] if net.A[p]["status"] == "open"]
        reo = [j for j in sorted(net.down[i], key=net.seq.index) if net.A[j]["status"] != "open"]
        md.append(f"| {k} | {i} | {net.A[i]['name']} | {STATUS[net.A[i]['status']]} | "
                  f"{'yes' if not wait else 'after ' + ', '.join(wait)} | {len(reo)} issued activities |\n")
    md.append(f"""
The provisional items (assay, H&MB, stress screening, SIL determination) sit at the start of long chains:
replacing the shortcut process model (PRC-00, PRC-01) re-opens {len(net.down['PRC-01'])} of the
{len(ACTIVITIES)} activities. Nothing in this package is reviewed or sealed by a licensed engineer; PRJ-04 is the
gate before any use for procurement or construction.

# 8 Check against the code
The build was run with file access recorded (`python -m cfu.wrapup.depmap --trace`, data/dataflow.json).
Each read of a data file was compared with this map:
""")
    for k, v in sorted(res.items()):
        md.append(f"- {k}: {v}\n")
    md.append("\nIterations found in the code (a module reads a file written later in the build, so it uses the "
              "previous revision):\n\n")
    its = sorted({(r["module"], r["data"], r["producer"]) for r in chk_rows
                  if r["result"].startswith(("iteration", "build order"))})
    md += [f"- {m} reads data/{d} from {p}\n" for m, d, p in its] or ["- none\n"]
    md.append("\nIssues:\n\n")
    md += [f"- {s}\n" for s in chk_issues] or ["- none: every traced read is covered by the map\n"]
    md.append("""
# 9 Limits
- The map covers the activities in this FEED package and the main ones still needed. Detailed design adds more
  (structural steel, instrument installation details, procurement, construction planning).
- Links are finish-to-start for clarity. Real schedules overlap activities with holds and early releases.
- The code check covers data files only. Shared constants imported from cfu/basis.py (PRJ-01) reach every module.
""")
    docgen.render("".join(md), OUT / "CFU-000-PM-DEP-001_Design-Dependency-Map", "CFU-000-PM-DEP-001",
                  "Design Dependency Map")


# ---------------------------------------------------------------------------------------------- interactive page
def page(net, G, docs, titles, chk_rows, chk_issues):
    from .portal import GH
    files = {}
    for p in sorted((ROOT / "deliverables").rglob("*")):
        if p.is_file() and p.suffix in (".pdf", ".xlsx"):
            for no in titles:
                if p.name.startswith(no):
                    files.setdefault(no, []).append(p)
    link = {}
    for no, ps in files.items():
        best = sorted(ps, key=lambda p: (p.suffix != ".pdf", "ALL" not in p.name, len(p.name)))[0]
        link[no] = GH + str(best.relative_to(ROOT))
    nodes = []
    for i in net.seq:
        a = net.A[i]
        x, y = G["pos"][i]
        nodes.append(dict(id=i, name=a["name"], disc=a["disc"], status=a["status"], wave=net.wave[i], x=x, y=y,
                          lines=wrap(a["name"]), after=[dict(id=p, what=w) for p, w in a["after"]], ext=a["ext"],
                          writes=a["writes"], note=a["note"],
                          docs=[dict(no=no, t=titles.get(no, ""), href=link.get(no, "")) for no in docs[i]],
                          up=sorted(net.up[i], key=net.seq.index), down=sorted(net.down[i], key=net.seq.index),
                          succ=sorted(net.succ[i], key=net.seq.index)))
    data = dict(nodes=nodes, edges=G["edges"], lanes=G["lanes"], W=G["W"], H=G["H"], waves=G["waves"],
                NW=NW, NH=NH, CW=CW, LANE_W=LANE_W, status=STATUS, loops=net.fb,
                check=dict(n=len(chk_rows), issues=chk_issues,
                           iters=sorted({f"{r['module']} reads {r['data']} ({r['producer']})" for r in chk_rows
                                         if r["result"].startswith(("iteration", "build order"))})))
    tpl = (Path(__file__).with_name("depmap_page.html")).read_text()
    n = {s: sum(a["status"] == s for a in ACTIVITIES) for s in STATUS}
    out = (tpl.replace("/*DATA*/null", json.dumps(data, separators=(",", ":")).replace("</", "<\\/"))
              .replace("{{N_ACT}}", str(len(ACTIVITIES))).replace("{{N_DONE}}", str(n["done"]))
              .replace("{{N_PART}}", str(n["partial"])).replace("{{N_OPEN}}", str(n["open"]))
              .replace("{{N_WAVES}}", str(G["waves"])).replace("{{N_LOOPS}}", str(len(net.fb)))
              .replace("{{N_LINKS}}", str(len(net.flow))).replace("{{REV}}", html.escape(basis.PROJECT["rev"]))
              .replace("{{GH}}", GH))
    PORTAL.mkdir(parents=True, exist_ok=True)
    (PORTAL / "dependencies.html").write_text(out)


# ---------------------------------------------------------------------------------------------- entry
def build():
    net = Net()
    G = geometry(net)
    docs, titles, unmapped = doc_map(net)
    chk_rows, chk_issues = code_check(net)
    chk_issues += [f"{no} is not assigned to an activity" for no in unmapped]
    OUT.mkdir(parents=True, exist_ok=True)
    drawing(net, G)
    workbook(net, docs, titles, chk_rows, chk_issues)
    report(net, docs, chk_rows, chk_issues)
    page(net, G, docs, titles, chk_rows, chk_issues)
    for s in chk_issues:
        print("  DEPMAP:", s)
    return net


if __name__ == "__main__":
    if "--trace" in sys.argv:
        recs = trace()
        print(f"traced {len(recs)} data-file accesses -> data/dataflow.json")
    build()
