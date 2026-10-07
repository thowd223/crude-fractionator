"""Wrap-up deliverables: preliminary HAZOP, cost estimate, document register."""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .. import basis, docgen

ROOT = Path(__file__).resolve().parents[2]
DLV = ROOT / "deliverables"
DOCNO = re.compile(r"(CFU-\d{3}-[A-Z]{2}-[A-Z0-9]{2,4}-[0-9A-Z]{3})")


def register():
    docs = defaultdict(lambda: dict(formats=set(), files=[], title=""))
    for p in sorted(DLV.rglob("*")):
        if not p.is_file() or "figures" in p.parts:
            continue
        m = DOCNO.search(p.name)
        if not m:
            continue
        no = m.group(1)
        d = docs[no]
        d["formats"].add(p.suffix.lstrip("."))
        d["files"].append(str(p.relative_to(ROOT)))
        rest = p.stem[len(no):].strip("_- ")
        if rest and not d["title"]:
            d["title"] = rest.replace("-", " ").replace("_", " ")
        d["folder"] = str(p.parent.relative_to(DLV))
    disc = {"PR": "Process", "ME": "Mechanical", "PI": "Piping", "PL": "Plant layout", "IC": "Instrumentation & control",
            "EL": "Electrical", "PM": "Project management"}
    rows = []
    for no in sorted(docs):
        d = docs[no]
        parts = no.split("-")
        rows.append(dict(no=no, area=parts[1], disc=disc.get(parts[2], parts[2]), type=parts[3],
                         title=d["title"] or "(drawing)", formats=", ".join(sorted(d["formats"])),
                         folder=d["folder"], rev=basis.PROJECT["rev"]))
    out = DLV / "06-wrapup"
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    hdr = ["Document No.", "Area", "Discipline", "Type", "Title", "Rev", "Formats", "Folder"]
    ws.append(hdr)
    for r in rows:
        ws.append([r["no"], r["area"], r["disc"], r["type"], r["title"], r["rev"], r["formats"], r["folder"]])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    for col, w in zip("ABCDEFGH", (24, 6, 22, 6, 50, 5, 18, 36)):
        ws.column_dimensions[col].width = w
    wb.save(out / "CFU-000-PM-REG-001_Document-Register.xlsx")
    md = "# Document register\n\n| Document No. | Discipline | Title | Rev | Formats | Folder |\n|---|---|---|---|---|---|\n"
    md += "".join(f"| {r['no']} | {r['disc']} | {r['title']} | {r['rev']} | {r['formats']} | {r['folder']} |\n"
                  for r in rows)
    docgen.render(md, out / "CFU-000-PM-REG-001_Document-Register", "CFU-000-PM-REG-001", "Document Register")
    return rows


def build():
    from . import depmap, estimate, hazop, portal
    hazop.build()
    estimate.build()
    depmap.build()       # dependency map (reads data/dataflow.json from the last --trace run)
    portal.build()       # also refreshes the register
